# Copyright (c) 2026, CyveTech and contributors
# For license information, please see license.txt

"""Upload Attendance: the template out, the filled file back in.

No Frappe import, like the other *_rules.py modules, so
scripts/verify_upload_attendance.py runs it without a bench.
attendance_upload.py is the part that asks the site and marks the days.

WHY THIS APP HAS ONE

Frappe HR v16 deleted its own Upload Attendance (patch
v16_0/delete_upload_attendance_doctype, "deprecated in favor of the
Employee Attendance Tool"). That tool marks one day at a time; a month kept
on paper goes in faster as one sheet. So the doctype is back, under the
name people look for, as one record per upload.

THE TEMPLATE

One row per employee per day that has no attendance yet (Luuka, 22 Oct
2026, of the register that stood in its place for a day: "you changed the
import file please take it back to this"). A person's rows start on the
period's first day or the day they joined, and stop on the period's last
day, the day they left or today, whichever comes first. A day already
marked, even in draft, is left out, so nobody fills a day twice. A holiday
says so in the Note column.

The rows are written as they are made, never held together, so a template
is as long as one Excel sheet allows (MAX_ROWS): Luuka over one cycle came
to 103,182 rows (21 Oct 2026). openpyxl writes the book around them, and
the rows are written as openpyxl writes them, as text: openpyxl takes some
0.4 ms over a row, so those 103,182 took 42 seconds and two cycles would
outlast the request; as text they take 2, and a full sheet 20.

THE FILE BACK

Only rows with a Status are marked: a row left blank is a day nobody has
marked yet. The file is read as it goes, so one of any length is never held
whole, and every row is checked before it is written, each problem naming
its row. Columns are found by their heading, so a Data Import file
("Employee (employee)") reads too, and other columns are left alone.
"""

import csv
import datetime
import io
import re
import zipfile

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, NamedStyle, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

# Frappe HR's own Attendance statuses, in their order
STATUSES = ("Present", "Absent", "On Leave", "Half Day", "Work From Home")
# Frappe HR marks On Leave without a Leave Type, but then nothing says
# whether the day is paid
NEEDS_LEAVE_TYPE = ("On Leave",)
# the statuses a Leave Type is kept on; Frappe HR clears it on the others
TAKES_LEAVE_TYPE = ("On Leave", "Half Day")
HEADERS = ("Employee", "Employee Name", "Department", "Attendance Date", "Status", "Leave Type", "Shift", "Note")
SHEET, LISTS, HOW = "Attendance", "Lists", "How to fill"
HOLIDAY = "Holiday"
WHY = "Why Not Marked"
# the rows one Excel sheet holds under its heading: the only limit on a template
MAX_ROWS = 1048575
# a template covers a year at most: a period longer than that is a mistake
MAX_DAYS = 366
# how each column read back may be headed, compared without case; "Label
# (fieldname)" and a bare fieldname come to the same thing (heading())
COLUMNS = {
    "employee": ("employee", "employee id"),
    "attendance_date": ("attendance date", "date"),
    "status": ("status",),
    "leave_type": ("leave type",),
    "shift": ("shift",),
}
NEEDED = ("employee", "attendance_date", "status")
# the header is looked for in this many rows from the top
HEADER_ROWS = 10
# day first, as dates are written in Uganda
DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y/%m/%d")
EXCEL_EPOCH = datetime.date(1899, 12, 30)
# the characters XML cannot hold; openpyxl refuses a cell with one
NOT_IN_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
# the text of the one row openpyxl writes in each style, for the rows to take
# its place (_with_rows)
PROBE = "ha-rows-go-here"

NOT_A_SHEET = "This is not an Excel (.xlsx) or .csv file. Save it as an Excel Workbook and attach it again."
NO_HEADER = ("The file has no Employee, Attendance Date and Status columns. "
             "Download the template and fill that in.")
NOTHING_FILLED = "No row has a Status. Put a Status on each day to mark and attach the file again."

HOW_TO_FILL = (
    "Put a Status on each day to mark: Present, Absent, On Leave, Half Day or Work From Home.",
    "A day left blank is not marked.",
    "On Leave needs a Leave Type.",
    "Leave the Shift blank unless the day was worked on a shift that is not the employee's own.",
    "Days already marked are not in this file. A holiday says so in the Note column.",
    "Attach the file to the Upload Attendance and press Import.",
)


# ── The template ──────────────────────────────────────────────────────
def period_days(start, end, today=None, joined=None, left=None):
    """The days of the period a person can be marked on: from its first day
    or the day they joined, to its last day, the day they left or today."""
    first = max(day for day in (_date(start), _date(joined)) if day)
    last = min(day for day in (_date(end), _date(left), _date(today)) if day)
    return [first + datetime.timedelta(days=step) for step in range((last - first).days + 1)]


def template_rows(people, start, end, today, marked, holidays=None):
    """The template's rows, in HEADERS order, one at a time as they are
    made: one per person per day still to be marked.

    people    [{"employee", "employee_name", "department", "date_of_joining",
              "relieving_date", "holiday_list"}], in the order they are listed
    marked    {(employee, date)} already marked, in draft or submitted
    holidays  {holiday list: {dates}}
    """
    holidays = holidays or {}
    for person in people:
        kept = holidays.get(person.get("holiday_list")) or set()
        for day in period_days(start, end, today, person.get("date_of_joining"), person.get("relieving_date")):
            if (person["employee"], day) in marked:
                continue
            yield (person["employee"], person.get("employee_name") or "", person.get("department") or "",
                   day, None, None, None, HOLIDAY if day in kept else None)


def count_rows(people, start, end, today, marked):
    """How many rows the template has, counted before it is made."""
    return sum(1 for person in people
               for day in period_days(start, end, today, person.get("date_of_joining"), person.get("relieving_date"))
               if (person["employee"], day) not in marked)


def build_template(rows, total, leave_types=(), shifts=()):
    """The template as .xlsx bytes: the rows (total of them), a list each
    Status, Leave Type and Shift is picked from, and how to fill it in."""
    book = openpyxl.Workbook(write_only=True)
    _add_styles(book)
    sheet = book.create_sheet(SHEET)
    last = total + 1
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = "A1:%s%d" % (_letter(len(HEADERS)), last)
    for column, width in zip("ABCDEFGH", (14, 30, 28, 16, 18, 22, 18, 12)):
        sheet.column_dimensions[column].width = width
    _pick(sheet, "E", '"%s"' % ",".join(STATUSES), last)
    if leave_types:
        _pick(sheet, "F", "'%s'!$A$2:$A$%d" % (LISTS, len(leave_types) + 1), last)
    if shifts:
        _pick(sheet, "G", "'%s'!$B$2:$B$%d" % (LISTS, len(shifts) + 1), last)
    sheet.append([_styled(sheet, title, "ha_head") for title in HEADERS])
    # one row in each style the rows take, so the book has them all; the
    # rows go in its place
    sheet.append([_styled(sheet, PROBE, "ha_text"), None, None, _styled(sheet, datetime.date(2026, 1, 1), "ha_date"),
                  None, None, None, _styled(sheet, PROBE, "ha_note")])
    lists = book.create_sheet(LISTS)
    lists.sheet_state = "hidden"
    lists.append(["Leave Type", "Shift"])
    for at in range(max(len(leave_types), len(shifts))):
        lists.append([leave_types[at] if at < len(leave_types) else None, shifts[at] if at < len(shifts) else None])
    how = book.create_sheet(HOW)
    how.column_dimensions["A"].width = 100
    for line in HOW_TO_FILL:
        how.append([line])
    out = io.BytesIO()
    book.save(out)
    return _with_rows(out.getvalue(), rows)


def _with_rows(content, rows):
    """The book openpyxl wrote, with the rows in place of its row of styles:
    each written as openpyxl writes one, the Employee as text, the date as
    its day number and a holiday's Note in grey."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as book:
        for info in source.infolist():
            part = source.read(info.filename)
            if PROBE.encode() not in part:
                book.writestr(info.filename, part)
                continue
            xml = part.decode("utf-8")
            start = xml.index('<row r="2"')
            end = xml.index("</row>", start) + len("</row>")
            styles = {}
            for tag in re.findall(r"<c [^>]*>", xml[start:end]):
                style = re.search(r' s="(\d+)"', tag)
                styles[re.search(r' r="([A-Z]+)2"', tag).group(1)] = ' s="%s"' % style.group(1) if style else ""
            with book.open(info.filename, "w") as sheet:
                sheet.write(xml[:start].encode("utf-8"))
                written = []
                for number, row in enumerate(rows, 2):
                    written.append(_row(number, row, styles))
                    if len(written) == 5000:
                        sheet.write("".join(written).encode("utf-8"))
                        written = []
                sheet.write(("".join(written) + xml[end:]).encode("utf-8"))
    return out.getvalue()


def _row(number, row, styles):
    """One template row as openpyxl writes it: a date as its day number,
    anything else as text in the cell, a blank left out."""
    cells = []
    for column, value in zip("ABCDEFGH", row):
        if value is None or value == "":
            continue
        if isinstance(value, datetime.date):
            cells.append('<c r="%s%d"%s t="n"><v>%d</v></c>'
                         % (column, number, styles.get(column, ""), (_date(value) - EXCEL_EPOCH).days))
            continue
        value = NOT_IN_XML.sub("", str(value))
        kept = value.strip()
        # spaces before or after the text are kept, as openpyxl keeps them
        space = ' xml:space="preserve"' if kept and kept != value else ""
        value = value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\r", "&#13;")
        cells.append('<c r="%s%d"%s t="inlineStr"><is><t%s>%s</t></is></c>'
                     % (column, number, styles.get(column, ""), space, value))
    return '<row r="%d">%s</row>' % (number, "".join(cells))


def _add_styles(book):
    # named, so each of a large template's cells points at one style
    # rather than carrying its own (output_sheet.py does the same)
    for name, style in (
        ("ha_head", {"font": Font(bold=True), "fill": PatternFill("solid", fgColor="DDE4EE")}),
        ("ha_text", {"number_format": "@"}),
        ("ha_date", {"number_format": "yyyy-mm-dd"}),
        ("ha_note", {"font": Font(color="808080")}),
    ):
        book.add_named_style(NamedStyle(name=name, **style))


def _styled(sheet, value, style):
    cell = WriteOnlyCell(sheet, value)
    cell.style = style
    return cell


def _pick(sheet, column, formula, last):
    """A list to pick the column's value from, on every row of the sheet: a
    list to pick from, not a fence, so a value typed over it is read back
    and said against its row if it is wrong."""
    rule = DataValidation(type="list", formula1=formula, allow_blank=True)
    rule.showErrorMessage = False
    rule.add("%s2:%s%d" % (column, column, max(last, 2)))
    sheet.data_validations.append(rule)


def _letter(number):
    return "ABCDEFGHIJKLMNOPQRSTUVWXYZ"[number - 1]


# ── The file back ─────────────────────────────────────────────────────
def heading(title):
    """A column heading as it is compared: no case, no "(fieldname)", an
    underscore a space."""
    text = re.sub(r"\s*\([^)]*\)\s*$", "", str(title or "")).strip().lower().replace("_", " ")
    return " ".join(text.split())


def header_columns(header):
    """{what: column number from 0} for the columns read back, or None when
    the row has no Employee, Attendance Date and Status."""
    found = {}
    for at, title in enumerate(header or ()):
        key = heading(title)
        for what, names in COLUMNS.items():
            if key in names and what not in found:
                found[what] = at
    return found if all(what in found for what in NEEDED) else None


def header_of(content, filename=""):
    """The file's heading, as it is: what Import checks before it queues
    anything, and what the days not marked are written under. ValueError
    when the file is not one this reads."""
    for number, values in enumerate(_table(content, filename), 1):
        values = tuple(values)
        if header_columns(values) is not None:
            return values
        if number >= HEADER_ROWS:
            break
    raise ValueError(NO_HEADER)


def days_in(content, filename=""):
    """Every row with a Status, one after the other, read as it goes so a
    file of any length is never held whole: {"row": its number in the file,
    "values": the row as it is, and each column read back}. A row with no
    Status is a day nobody has marked. ValueError when the file is not one
    this reads."""
    columns = None
    for number, values in enumerate(_table(content, filename), 1):
        values = tuple(values)
        if columns is None:
            columns = header_columns(values)
            if columns is None and number >= HEADER_ROWS:
                raise ValueError(NO_HEADER)
            continue
        record = {"row": number, "values": values}
        for what, at in columns.items():
            value = values[at] if at < len(values) else None
            record[what] = value.strip() if isinstance(value, str) else value
        if _text(record.get("status")):
            yield record
    if columns is None:
        raise ValueError(NO_HEADER)


def _table(content, filename):
    if str(filename or "").lower().endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        yield from csv.reader(io.StringIO(text))
        return
    try:
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        # a .xls, a picture, a file cut short: openpyxl says so in its own words
        raise ValueError(NOT_A_SHEET)
    try:
        sheet = book[SHEET] if SHEET in book.sheetnames else book.worksheets[0]
        yield from sheet.iter_rows(values_only=True)
    finally:
        book.close()


def parse_date(value):
    """A date from a cell: a date, a day-first text date, or an Excel day
    number. None when it is none of those."""
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # a date cell read without its format is the number of days since 1899
        return EXCEL_EPOCH + datetime.timedelta(days=int(value)) if 20000 < value < 80000 else None
    text = _text(value).split(" ")[0].split("T")[0]
    for pattern in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


# ── The checks ────────────────────────────────────────────────────────
def check_days(days, employees, names, marked, today, seen=None):
    """What each row comes to, before anything is written.

    days         days_in()'s records
    employees    {name: {"date_of_joining", "relieving_date", "status"}},
                 matched without case, as the database matches them
    names        {"leave_types", "shifts"}: the names the site has
    marked       {(employee, date): (status, the Attendance)}, in draft or
                 submitted
    seen         {(employee, date): row}: the days met so far in the file,
                 kept by the caller across the parts of a large file
    Returns {"to_mark": [...], "already": n, "problems": [(row, why)]}."""
    people = {str(name).lower(): name for name in employees}
    leave_names = {str(name).lower(): name for name in names.get("leave_types") or ()}
    shift_names = {str(name).lower(): name for name in names.get("shifts") or ()}
    statuses = {status.lower(): status for status in STATUSES}
    today = _date(today)
    seen = {} if seen is None else seen
    out = {"to_mark": [], "already": 0, "problems": []}
    for row in days:
        said = _text(row.get("status"))
        who, given = _text(row.get("employee")), row.get("attendance_date")
        name = people.get(who.lower())
        facts = employees.get(name) or {}
        day, status = parse_date(given), statuses.get(said.lower())
        leave, shift = _text(row.get("leave_type")), _text(row.get("shift"))
        joined, left = _date(facts.get("date_of_joining")), _date(facts.get("relieving_date"))
        why = None
        if not who:
            why = "No employee on this row."
        elif not name:
            why = "No employee %s." % who
        elif _text(given) == "":
            why = "No Attendance Date."
        elif day is None:
            why = "%s is not a date. Write it as 2026-10-05." % _text(given)
        elif not status:
            why = "%s is not a Status. Use %s or %s." % (said, ", ".join(STATUSES[:-1]), STATUSES[-1])
        elif day > today:
            why = "%s is after today." % day.isoformat()
        elif joined and day < joined:
            why = "%s joined on %s." % (name, joined.isoformat())
        elif left and day > left:
            why = "%s left on %s." % (name, left.isoformat())
        elif facts.get("status") == "Inactive":
            why = "%s is Inactive, and attendance is not marked for an inactive employee." % name
        elif status in NEEDS_LEAVE_TYPE and not leave:
            why = "%s needs a Leave Type." % status
        elif leave and status in TAKES_LEAVE_TYPE and leave.lower() not in leave_names:
            why = "No leave type %s." % leave
        elif shift and shift.lower() not in shift_names:
            why = "No shift %s." % shift
        elif (name, day) in seen:
            why = "%s on %s is on row %s too." % (name, day.isoformat(), seen[(name, day)])
        if why:
            out["problems"].append((row, why))
            continue
        seen[(name, day)] = row.get("row")
        already = marked.get((name, day))
        if already:
            if already[0] == status:
                out["already"] += 1
            else:
                out["problems"].append((row, "Already marked %s on %s. Cancel that one to mark the day again."
                                        % (already[0], already[1])))
            continue
        out["to_mark"].append({
            "row": row.get("row"), "employee": name, "attendance_date": day, "status": status,
            "leave_type": leave_names.get(leave.lower()) if leave and status in TAKES_LEAVE_TYPE else None,
            "shift": shift_names.get(shift.lower()) if shift else None,
            "source": row,
        })
    return out


# ── The days not marked ───────────────────────────────────────────────
def problems_file(header, problems):
    """The rows not marked as .xlsx bytes, each as it was with why it was
    not: put right, the file is uploaded again as it is."""
    keep = [at for at, title in enumerate(header) if heading(title) != heading(WHY)]
    book = openpyxl.Workbook(write_only=True)
    _add_styles(book)
    sheet = book.create_sheet(SHEET)
    sheet.freeze_panes = "B2"
    sheet.append([_styled(sheet, header[at], "ha_head") for at in keep] + [_styled(sheet, WHY, "ha_head")])
    for row, why in sorted(problems, key=lambda problem: problem[0].get("row") or 0):
        values = list(row.get("values") or ())
        values += [None] * (len(header) - len(values))
        cells = []
        for at in keep:
            value = values[at]
            if isinstance(value, (datetime.date, datetime.datetime)):
                value = _styled(sheet, _date(value), "ha_date")
            cells.append(value)
        sheet.append(cells + [why])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def listed(problems, most):
    """The first problems as lines to show on the upload; the file has them all."""
    ordered = sorted(problems, key=lambda problem: problem[0].get("row") or 0)
    lines = ["Row %s: %s" % (row.get("row"), why) for row, why in ordered[:most]]
    if len(ordered) > most:
        lines.append("And %d more, in the file." % (len(ordered) - most))
    return "\n".join(lines)


def text(value):
    """A cell as text, trimmed; nothing for an empty cell."""
    return _text(value)


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    return datetime.date.fromisoformat(str(value)[:10])
