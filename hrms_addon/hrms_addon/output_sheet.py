# Copyright (c) 2026, CyveTech and contributors
# For license information, please see license.txt

"""Luuka's Per Meter sheet, out and back (openpyxl, no Frappe import, so
scripts/verify_output_pay.py builds and reads it without a bench).

THE LAYOUT, AS LUUKA KEEPS IT (PER METER AUGUST 2026)

  A Emp Names       on each person's first row only
  B Employee Id
  C Attendance Date as text, 26-Jul-2026, every day of the month
  D Work Unit       the size's rate that day
  E Work Done       the metres, typed in
  F TOTAL           =E x D
  G BASIC           the day's TOTALs added up, one cell across the day's rows
  H LOOMS OT        typed in, one cell across the day's rows
  I TOTAL EARN      =G + H, one cell across the day's rows
  J Size            one row a day for each size, in the same order each day

What this writes over Luuka's own: each person's month total adds up the
LOOMS OT and the TOTAL EARN as well as the BASIC (the LOOMS OT is paid
now), and the grand total adds up the month totals by their label, so it
can neither miss a person nor count itself. Only Work Done and LOOMS OT can
be typed in; the rest is protected (no password), so a formula cannot be
typed over by mistake. A hidden sheet names the run it was made for.

READING ONE BACK

Either this sheet or Luuka's own: the columns are found by their headings
(the size column, which has none on Luuka's, is the one after TOTAL EARN),
and a person's rows run from their name in column A to the next name.
Whatever is on it is checked by output_rules.take_sheet.
"""

import datetime
import io
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, NamedStyle, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.cell_range import CellRange

HEADINGS = ("Emp Names", "Employee Id", "Attendance Date", "Work Unit", "Work Done", "TOTAL", "BASIC", "LOOMS OT",
            "TOTAL EARN", "Size")
MONTH_TOTAL = "Month Total"
GRAND_TOTAL = "TOTAL"
TAG_SHEET = "ha"
TAG_MARK = "hrms_addon per meter sheet"
# wide enough for a plant's month in the totals: 999,999,999.99 in bold
WIDTHS = (26, 13, 15, 10, 11, 13, 16, 15, 16, 13)
MONEY = "#,##0.00"
RATE = "0.00"
# as typed: "#,##0.##" would show a whole 98 as "98."
METRES = "General"
DAY_FORMAT = "%d-%b-%Y"

# the headings a column is found by when the sheet is read back
FIND = {
    "name": ("emp names", "employee name", "emp name", "name"),
    "employee_id": ("employee id", "emp id", "id"),
    "date": ("attendance date", "date"),
    "work_unit": ("work unit", "rate"),
    "work_done": ("work done", "metres", "meters"),
    "total_earn": ("total earn",),
    "looms_ot": ("looms ot", "ot"),
    "size": ("size",),
}


def day_text(day):
    return day.strftime(DAY_FORMAT)


def build(data):
    """The month's sheet as .xlsx bytes.

    data: {"title": "AUGUST 2026", "tag": {"run", "branch", "from", "to"},
           "employees": [{"employee", "name", "id"}],
           "dates": [date, ...], "sizes": [{"size", "rates": {date: work unit}}],
           "work": {(employee, date, size): metres},
           "overtime": {(employee, date): amount}}
    """
    # written a row at a time (write-only): a month is 35,000 rows, which
    # the ordinary workbook takes over a minute to build and this seconds
    book = Workbook(write_only=True)
    sheet = book.create_sheet(str(data.get("title") or "PER METER")[:31])
    _add_styles(book)
    for column, width in enumerate(WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.append([_styled(sheet, heading, "ha_head") for heading in HEADINGS])

    work, overtime = data.get("work") or {}, data.get("overtime") or {}
    sizes = data.get("sizes") or []
    dates = data.get("dates") or []
    row = 2
    for person in data.get("employees") or []:
        who = person["employee"]
        first = row
        for day in dates:
            day_first = row
            for size in sizes:
                label = size["size"]
                opening = row == day_first
                sheet.append([
                    _styled(sheet, person.get("name") or who, "ha_name") if row == first else None,
                    person.get("id") or who,
                    day_text(day),
                    _styled(sheet, (size.get("rates") or {}).get(day), "ha_rate"),
                    _styled(sheet, work.get((who, day, label)), "ha_metres"),
                    _styled(sheet, "=E%d*D%d" % (row, row), "ha_money"),
                    _styled(sheet, "=SUM(F%d:F%d)" % (row, row + len(sizes) - 1), "ha_day") if opening else None,
                    _styled(sheet, overtime.get((who, day)), "ha_day_in") if opening else None,
                    _styled(sheet, "=G%d+H%d" % (row, row), "ha_day") if opening else None,
                    label,
                ])
                row += 1
            if row - 1 > day_first:
                for column in ("G", "H", "I"):
                    # the day's three cells span its rows, as on Luuka's sheet
                    sheet.merged_cells.ranges.add(CellRange("%s%d:%s%d" % (column, day_first, column, row - 1)))
        sheet.append([_styled(sheet, MONTH_TOTAL, "ha_total_label")] + [_styled(sheet, None, "ha_total")] * 5 +
                     [_styled(sheet, "=SUM(%s%d:%s%d)" % (column, first, column, row - 1), "ha_total_money")
                      for column in ("G", "H", "I")] + [_styled(sheet, None, "ha_total")])
        row += 1

    end = max(row - 1, 2)
    sheet.append([_styled(sheet, GRAND_TOTAL, "ha_grand_label")] + [_styled(sheet, None, "ha_grand")] * 5 +
                 [_styled(sheet, '=SUMIF($A$2:$A$%d,"%s",%s$2:%s$%d)' % (end, MONTH_TOTAL, column, column, end),
                          "ha_grand_money") for column in ("G", "H", "I")] + [_styled(sheet, None, "ha_grand")])
    sheet.protection.sheet = True
    sheet.protection.formatCells = False
    sheet.protection.formatColumns = False
    sheet.protection.formatRows = False

    tag = book.create_sheet(TAG_SHEET)
    tag.sheet_state = "hidden"
    values = data.get("tag") or {}
    for value in (TAG_MARK, values.get("run"), values.get("branch"), _iso(values.get("from")), _iso(values.get("to"))):
        tag.append([value])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _add_styles(book):
    head, total = PatternFill("solid", fgColor="DDE4EE"), PatternFill("solid", fgColor="F2F2F2")
    bold, unlocked, middle = Font(bold=True), Protection(locked=False), Alignment(vertical="center")
    for name, style in (
        ("ha_head", {"font": bold, "fill": head}),
        ("ha_name", {"font": bold}),
        ("ha_rate", {"number_format": RATE}),
        ("ha_metres", {"number_format": METRES, "protection": unlocked}),
        ("ha_money", {"number_format": MONEY}),
        ("ha_day", {"number_format": MONEY, "alignment": middle}),
        ("ha_day_in", {"number_format": MONEY, "alignment": middle, "protection": unlocked}),
        ("ha_total_label", {"font": bold, "fill": total}),
        ("ha_total", {"fill": total}),
        ("ha_total_money", {"number_format": MONEY, "font": bold, "fill": total}),
        ("ha_grand_label", {"font": bold, "fill": head}),
        ("ha_grand", {"fill": head}),
        ("ha_grand_money", {"number_format": MONEY, "font": bold, "fill": head}),
    ):
        book.add_named_style(NamedStyle(name=name, **style))


def _styled(sheet, value, style):
    cell = WriteOnlyCell(sheet, value)
    cell.style = style
    return cell


def read(content):
    """The lines of a Per Meter sheet: {"sheet", "tag", "lines": [{"row",
    "block", "name", "employee_id", "date", "work_unit", "work_done",
    "size", "looms_ot"}]}. Raises ValueError when it is not an .xlsx
    workbook, or no sheet has the headings."""
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as error:
        raise ValueError("This is not an .xlsx workbook: save it as Excel Workbook (.xlsx) and upload it again.") \
            from error
    try:
        tag = _tag(book)
        for sheet in book.worksheets:
            if sheet.title == TAG_SHEET:
                continue
            found = _read_sheet(sheet)
            if found is not None:
                found["tag"] = tag
                return found
    finally:
        book.close()
    raise ValueError("No Per Meter sheet found: the first rows need the headings Employee Id and Work Done.")


def _read_sheet(sheet):
    rows = sheet.iter_rows(values_only=True)
    columns = None
    header_row = 0
    for number, values in enumerate(rows, 1):
        columns = _columns(values)
        if columns:
            header_row = number
            break
        if number >= 10:
            return None
    if not columns:
        return None
    lines = []
    block, name = None, None
    for number, values in enumerate(rows, header_row + 1):
        def get(field):
            index = columns.get(field)
            return values[index] if index is not None and index < len(values) else None
        raw_name, raw_id, done = get("name"), get("employee_id"), get("work_done")
        if _filled(raw_name) and _filled(raw_id):
            block = (block or 0) + 1
            name = str(raw_name).strip()
        if not _filled(raw_id) and not _filled(done):
            continue  # a total, a heading or an empty row
        lines.append({"row": number, "block": block, "name": name, "employee_id": raw_id, "date": get("date"),
                      "work_unit": get("work_unit"), "work_done": done, "size": get("size"),
                      "looms_ot": get("looms_ot")})
    return {"sheet": sheet.title, "lines": lines}


def _columns(values):
    headings = [str(value).strip().lower() if value is not None else "" for value in values or ()]
    found = {}
    for field, names in FIND.items():
        for index, heading in enumerate(headings):
            if heading in names and index not in found.values():
                found[field] = index
                break
    if "employee_id" not in found or "work_done" not in found:
        return None
    if "size" not in found:
        # Luuka's own sheet leaves the size column unheaded, after TOTAL EARN
        after = found.get("total_earn", max(found.values()))
        found["size"] = after + 1
    return found


def _tag(book):
    if TAG_SHEET not in book.sheetnames:
        return None
    values = [row[0] if row else None for row in book[TAG_SHEET].iter_rows(max_col=1, values_only=True)]
    if not values or values[0] != TAG_MARK:
        return None
    values += [None] * 5
    return {"run": values[1], "branch": values[2], "from": values[3], "to": values[4]}


def _filled(value):
    return value is not None and str(value).strip() != ""


def _iso(value):
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.strftime("%Y-%m-%d")
    return str(value) if value else None
