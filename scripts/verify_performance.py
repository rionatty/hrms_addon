"""Verify performance management, without a bench:

    python scripts/verify_performance.py

Luuka's appraisal round (the flowchart and the test script's cases 1 to 10),
on the shape of the Supervisory Skills Evaluation Form (LPL/HR/18), built on
Frappe HR's Appraisal Cycle and Appraisal so the round keeps its appraisee
list and its Appraisal Overview chart.

  1  the rules: the form's sections, the scale, the bands, the quarters and
     their deadlines, the year's average, what a score suggests
  2  the signatures: Employee, Supervisor, HR Manager, Production Manager,
     General Manager
  3  the forms carry the paper, with the rights each role needs
  4  the glue reads and writes fields that exist, here and upstream
  5  the print-outs say what the paper says
  6  wiring: the doc events, the workflow on migrate, the daily jobs, the
     seed, the patch, the way in

Frappe HR's and ERPNext's own fields are read from FRAPPE_APPS_ROOT
(default ../ERPNext).
"""
import ast
import datetime
import ast
import glob
import importlib.util
import json
import os
import re
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PACKAGE = os.path.join(REPO, "hrms_addon")
APP = os.path.join(PACKAGE, "hrms_addon")
APPS_ROOT = os.environ.get("FRAPPE_APPS_ROOT", os.path.join(os.path.dirname(REPO), "ERPNext"))
fail = []


def read(*parts):
    return open(os.path.join(REPO, *parts), encoding="utf-8").read()


def load(name):
    spec = importlib.util.spec_from_file_location(name, os.path.join(APP, name + ".py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)  # proves it has no Frappe import
    return module


def doctype(name):
    folder = name.lower().replace(" ", "_")
    path = os.path.join(APP, "doctype", folder, folder + ".json")
    return json.load(open(path, encoding="utf-8")) if os.path.exists(path) else {}


def upstream_doctype(name):
    folder = name.lower().replace(" ", "_")
    for app in ("frappe", "erpnext", "hrms"):
        hits = glob.glob(os.path.join(APPS_ROOT, app, app, "**", "doctype", folder, folder + ".json"), recursive=True)
        if hits:
            return json.load(open(hits[0], encoding="utf-8"))
    return None


def fields_of(spec):
    return {f["fieldname"]: f for f in (spec or {}).get("fields", [])}


def upstream_fields(name):
    return {f["fieldname"]: f for f in (upstream_doctype(name) or {}).get("fields", [])}


def custom_fields(dt):
    return {row["fieldname"]: row for row in CUSTOM if row.get("dt") == dt}


def all_fields(name):
    """Every field the site has on a DocType: its own JSON, here or
    upstream, plus this app's Custom Fields."""
    spec = doctype(name) or upstream_doctype(name)
    return dict(fields_of(spec), **custom_fields(name))


def expect(label, got, *needles):
    if not needles:
        if got:
            fail.append("%s: expected no errors, got %s" % (label, got))
        return
    if len(got) != len(needles):
        fail.append("%s: expected %d error(s), got %s" % (label, len(needles), got))
    for needle in needles:
        if not any(needle in message for message in got):
            fail.append("%s: expected an error containing %r, got %s" % (label, needle, got))


def print_format(name):
    folder = name.lower().replace(" ", "_").replace("'", "")
    path = os.path.join(APP, "print_format", folder, folder + ".json")
    return json.load(open(path, encoding="utf-8")) if os.path.exists(path) else None


CUSTOM = json.load(open(os.path.join(PACKAGE, "fixtures", "custom_field.json"), encoding="utf-8"))
SETTERS = json.load(open(os.path.join(PACKAGE, "fixtures", "property_setter.json"), encoding="utf-8"))
R, A, P = load("appraisal_rules"), load("appraisal_approval"), load("pip_rules")
S = load("bsc_rules")
print("loaded appraisal_rules.py, appraisal_approval.py, pip_rules.py and bsc_rules.py without Frappe")

# ── 1. The rules ──────────────────────────────────────────────────────
if len(R.FACTORS) != 12:
    fail.append("Section A of LPL/HR/18 has twelve ratable factors, not %d" % len(R.FACTORS))
if R.FACTORS[0] != "Job performance, work output, quality of work":
    fail.append("Section A must open with the form's first factor")
if "changes brought by the supervisor" not in R.FACTORS[2]:
    fail.append("the third factor is the supervisory form's, not the probation form's")
if R.APPRAISAL_MASTERS != {"Appraisal Factor": ("factor_name", R.FACTORS)}:
    fail.append("the Appraisal Factor list is seeded with exactly Section A")
if len(R.QUESTIONS) != 4:
    fail.append("the form's General section asks four questions, not %d" % len(R.QUESTIONS))
if (R.FACTORS_WEIGHT, R.OBJECTIVES_WEIGHT) != (60, 40):
    fail.append("Section C weights the factors 60 and the objectives 40")
if R.MAX_OBJECTIVES != 8:
    fail.append("the form carries at most eight objectives")
if R.BANDS != ((90, "Excellent"), (75, "Very Good"), (60, "Good"), (50, "Average"), (0, "Below Average")):
    fail.append("the bands are the form's: 90, 75, 60, 50")
if R.NOT_APPLICABLE not in R.RATINGS or R.TOP_RATING != 5:
    fail.append("the scale is 1 to 5 with N/A for a factor that does not fit the job")

if R.section_percent(["5", "5", "N/A", None]) != 100.0:
    fail.append("N/A and blanks are left out of a section's percentage")
if R.section_percent(["N/A", None]) is not None:
    fail.append("a section with nothing rated scores nothing at all")
found = R.scores(["5"] * 12, ["4"] * 8)
if (found["factors"], found["objectives"], found["total"]) != (60.0, 32.0, 92.0):
    fail.append("Section C: all fives and all fours must give 60/60, 32/40, 92%%; got %s" % found)
if R.scores(["3"] * 12, [])["total"] != 60.0:
    fail.append("a section with nothing rated leaves the total to the other alone")
for total, name in ((100, "Excellent"), (90, "Excellent"), (89.9, "Very Good"), (75, "Very Good"),
                    (74, "Good"), (60, "Good"), (59, "Average"), (50, "Average"), (49, "Below Average"), (0, "Below Average")):
    if R.band(total) != name:
        fail.append("%s%% is %s, not %r" % (total, name, R.band(total)))
if R.band(None) is not None:
    fail.append("nothing rated has no band")
if hasattr(R, "annual_average"):
    fail.append("the year to date is bsc_rules.year_to_date, on both forms (Oct 2026): appraisal_rules.annual_average is gone")
if R.PIP_BELOW != 60:
    fail.append("the recommendation puts anyone below 60 on an improvement plan")
if R.recommended(59.9) != R.PIP or R.recommended(60) != R.CLOSE or R.recommended(None) is not None:
    fail.append("below the pass mark suggests a PIP, at or above it suggests closing")
if R.POSITION_CHANGE_FOR.get(R.PROMOTION) != "Promotion" or R.POSITION_CHANGE_FOR.get(R.INCREASE) != "Salary Increment":
    fail.append("a promotion and an increase are carried out by an Employee Position Change")

if R.quarter_window(2026, "Q1") != (datetime.date(2026, 1, 1), datetime.date(2026, 3, 31)):
    fail.append("Q1 runs January to March: %s" % (R.quarter_window(2026, "Q1"),))
if R.quarter_window(2026, "Q4") != (datetime.date(2026, 10, 1), datetime.date(2026, 12, 31)):
    fail.append("Q4 runs October to December: %s" % (R.quarter_window(2026, "Q4"),))
if R.deadlines(2026, "Q1") != (datetime.date(2026, 4, 25), datetime.date(2026, 4, 30)):
    fail.append("Q1 is appraised by the 25th of April, at the latest that month's end: %s" % (R.deadlines(2026, "Q1"),))
if R.deadlines(2026, "Q4") != (datetime.date(2027, 1, 25), datetime.date(2027, 1, 31)):
    fail.append("Q4's deadlines fall in the next year: %s" % (R.deadlines(2026, "Q4"),))
if R.SOFT_DEADLINE_DAY != 25:
    fail.append("the recommendation sets the soft deadline on the 25th")
if R.REMINDER_DAYS != (7, 1, 0):
    fail.append("appraisers are reminded a week before, a day before and on the day")

PLAN = {"year": 2026, "company": "Luuka Plastics Limited",
        "quarters": [{"quarter": q, "from_date": str(R.quarter_window(2026, q)[0]),
                      "to_date": str(R.quarter_window(2026, q)[1]),
                      "soft_deadline": str(R.deadlines(2026, q)[0]), "hard_deadline": str(R.deadlines(2026, q)[1])}
                     for q in R.QUARTERS]}
expect("a whole year planned", R.plan_errors(PLAN))
expect("no year", R.plan_errors(dict(PLAN, year=None)), "year the plan covers")
expect("no quarters", R.plan_errors(dict(PLAN, quarters=[])), "List the quarters")
expect("a quarter listed twice", R.plan_errors(dict(PLAN, quarters=PLAN["quarters"] + [PLAN["quarters"][0]])),
       "is listed twice")
expect("a deadline inside the quarter it appraises",
       R.plan_errors(dict(PLAN, quarters=[dict(PLAN["quarters"][0], soft_deadline=None,
                                               hard_deadline="2026-02-01")])),
       "falls inside the quarter it appraises")
expect("a soft deadline after the hard one",
       R.plan_errors(dict(PLAN, quarters=[dict(PLAN["quarters"][0], soft_deadline="2026-04-29",
                                               hard_deadline="2026-04-28")])),
       "after its hard deadline")

RATED = {"step": "self", "roles": "Run the line", "skills": "Planning",
         "factors": [{"item": f, "employee_rating": "4"} for f in R.FACTORS],
         "objectives": [{"item": "Output", "employee_rating": "4"}]}
expect("a self-assessment complete", R.appraisal_errors(RATED))
expect("a factor not rated",
       R.appraisal_errors(dict(RATED, factors=[dict(row, employee_rating=None) for row in RATED["factors"]])),
       "Rate yourself on every factor")
expect("no objectives listed", R.appraisal_errors(dict(RATED, objectives=[])), "List the objectives")
expect("the General questions unanswered", R.appraisal_errors(dict(RATED, roles="", skills=" ")),
       "Answer the General questions")
SUPERVISED = {"step": "supervisor",
              "factors": [{"item": f, "supervisor_rating": "4"} for f in R.FACTORS],
              "objectives": [{"item": "Output", "supervisor_rating": "3"}]}
expect("the supervisor's rating complete", R.appraisal_errors(SUPERVISED))
expect("the supervisor left one unrated",
       R.appraisal_errors(dict(SUPERVISED, objectives=[{"item": "Output"}])), "Give your rating")
expect("too many objectives",
       R.appraisal_errors(dict(SUPERVISED, objectives=[{"item": str(i), "supervisor_rating": "3"} for i in range(9)])),
       "at most 8 objectives")

due = R.due_quarters([{"quarter": "Q1", "to_date": "2026-03-31"},
                      {"quarter": "Q2", "to_date": "2026-06-30"},
                      {"quarter": "Q3", "to_date": "2026-09-30", "appraisal_cycle": "C3"},
                      {"quarter": "Q4", "to_date": "2026-12-31"}], "2026-10-05")
if [row["quarter"] for row in due] != ["Q1", "Q2"]:
    fail.append("a quarter is due once it has closed and has no cycle yet: %s" % [row["quarter"] for row in due])
if R.due_quarters([{"quarter": "Q1", "to_date": "2026-03-31", "notified_on": "2026-04-01"}], "2026-10-05"):
    fail.append("the HR Officer is told of a quarter once, not every day")
if R.reminders_due("2026-04-30", "2026-04-20", None) != []:
    fail.append("nothing is sent before the first threshold")
if R.reminders_due("2026-04-30", "2026-04-25", None, "2026-04-25") != ["soft", "7"]:
    fail.append("the soft deadline and the week-before reminder fall together: %s"
                % R.reminders_due("2026-04-30", "2026-04-25", None, "2026-04-25"))
if R.reminders_due("2026-04-30", "2026-04-29", "7") != ["1"]:
    fail.append("a day before, once the week's has gone")
if R.reminders_due("2026-04-30", "2026-04-30", "7, 1") != ["0"]:
    fail.append("and on the day")
if R.reminders_due("2026-04-30", "2026-05-01", None) != []:
    fail.append("nothing is sent once the deadline has passed")
if R.reminders_due("2026-04-30", "2026-04-29", None) != ["1"]:
    fail.append("a deadline first seen inside several thresholds gets the nearest")
if "7" not in R.record_reminders(None, ["1"]) or "1" not in R.record_reminders(None, ["1"]):
    fail.append("recording the nearest threshold marks the ones already passed")

for given, wanted in (("5", "5"), (5, "5"), (5.0, "5"), (" n/a ", R.NOT_APPLICABLE), ("7", None), ("", None),
                      (None, None), ("rubbish", None), (4.5, None)):
    if R.rating(given) != wanted:
        fail.append("a rating read from a sheet: %r must be %r, not %r" % (given, wanted, R.rating(given)))
# Appraisal Settings: a Check never saved reads as nothing, not as off
if R.settings_values({}) != {"self_appraisal": 1, "kra_evaluation_method": R.KRA_AUTOMATED}:
    fail.append("unsaved Appraisal Settings are the defaults: employees appraise themselves, KRAs scored "
                "automatically: %s" % R.settings_values({}))
if R.settings_values({"self_appraisal": "0"})["self_appraisal"] != 0 or \
        R.settings_values({"self_appraisal": None})["self_appraisal"] != 1:
    fail.append("the self-appraisal is off only when saved off")
if R.settings_values({"kra_evaluation_method": "Something else"})["kra_evaluation_method"] != R.KRA_AUTOMATED:
    fail.append("a method Frappe HR does not have falls back to the automated one")
if R.SETTINGS_DEFAULTS["kra_evaluation_method"] != "Automated Based on Goal Progress":
    fail.append("the KRA evaluation method defaults to Automated Based on Goal Progress")
if (R.rates_goals_manually(R.KRA_MANUAL), R.rates_goals_manually(R.KRA_AUTOMATED)) != (1, 0):
    fail.append("an appraisal rates its goals by hand only in a cycle rated by hand")
# whether the employee gave a self-appraisal: any rating of their own
for ratings, scores, wanted in (((), (), False), ((None, ""), (0, None, 0.0), False), ((None, "4"), (), True),
                                (("N/A",), (), True), ((), (0, 62.5), True), (("rubbish",), (), False)):
    if R.gave_self_appraisal(ratings, scores) is not wanted:
        fail.append("ratings %r and self scores %r: the employee %s themselves" % (
            ratings, scores, "rated" if wanted else "did not rate"))
if R.gave_self_appraisal(None, None) is not False:
    fail.append("an appraisal with no rows has no self-appraisal")
print("rules: the form's sections and scale, the bands, the quarters and deadlines, the year, the settings")

# ── 2. The signatures ─────────────────────────────────────────────────
if A.DOCTYPE != "Appraisal":
    fail.append("the workflow runs on Frappe HR's Appraisal, so the round keeps its cycle and its chart")
if A.STATUS_FIELD != "custom_appraisal_status":
    fail.append("Frappe HR's Appraisal has no status field; ours must carry the workflow's states")
route = [A.DRAFT, A.PENDING_SELF, A.PENDING_SUPERVISOR, A.PENDING_HRM, A.PENDING_PRODUCTION, A.PENDING_GM,
         A.COMPLETED]
walked, state = [A.DRAFT], A.DRAFT
while True:
    forward = [t for t in A.TRANSITIONS if t["state"] == state
               and t["action"] in (A.SEND_SELF, A.SELF, A.RATE, A.APPROVE)]
    if not forward:
        break
    state = forward[0]["next_state"]
    walked.append(state)
if walked != route:
    fail.append("the form is signed Employee, Supervisor, HR Manager, Production Manager, General Manager: %s" % walked)
SUPERVISORY_SIGNING = {state for state in A.route(A.FORM_SUPERVISORY) if state in A.PENDING_STATES}
if set(A.STAMPS) != SUPERVISORY_SIGNING or set(A.REMARK_FIELDS) != SUPERVISORY_SIGNING:
    fail.append("every comment block on LPL/HR/18 is stamped and signed: %s" % sorted(A.STAMPS))
# the self-appraisal, on or off (Appraisal Settings), each appraisal keeping its own
if A.opening(True) != (A.SEND_SELF, A.PENDING_SELF) or A.opening(False) != (A.SEND_SUPERVISOR, A.PENDING_SUPERVISOR):
    fail.append("a raised appraisal goes to the employee when they appraise themselves, else to the supervisor")
if A.next_states(A.DRAFT, {"HR User"}, A.FORM_SUPERVISORY, True) != [(A.SEND_SELF, A.PENDING_SELF)] or \
        A.next_states(A.DRAFT, {"HR User"}, A.FORM_SUPERVISORY, False) != [(A.SEND_SUPERVISOR, A.PENDING_SUPERVISOR)]:
    fail.append("from Draft HR send it one way or the other, never both: %s / %s"
                % (A.next_states(A.DRAFT, {"HR User"}, A.FORM_SUPERVISORY, True),
                   A.next_states(A.DRAFT, {"HR User"}, A.FORM_SUPERVISORY, False)))
if A.next_states(A.DRAFT, {"Employee"}):
    fail.append("Draft is HR's: the employee acts at their own self-appraisal")
if A.next_states(A.PENDING_SELF, {"Employee"}) != [(A.SELF, A.PENDING_SUPERVISOR)]:
    fail.append("the employee submits their self-appraisal to the supervisor: %s"
                % A.next_states(A.PENDING_SELF, {"Employee"}))
if (A.SELF, A.PENDING_SUPERVISOR) not in A.next_states(A.PENDING_SELF, {"HR User"}):
    fail.append("HR may submit the self-appraisal for someone with no login")
if [row for row in A.STATES if row["state"] == A.DRAFT and row["allow_edit"] == A.APPRAISEE]:
    fail.append("the employee edits their appraisal at Pending Self-Appraisal, not in HR's Draft")
if not [row for row in A.STATES if row["state"] == A.PENDING_SELF and row["allow_edit"] == A.APPRAISEE]:
    fail.append("the employee must be able to fill their self-appraisal in")
for form in A.FORM_TYPES:
    if A.PENDING_SELF in A.route(form, False) or A.route(form, True)[1] != A.PENDING_SELF:
        fail.append("%s: the self-appraisal comes first, and only where the employee appraises themselves" % form)
if A.holds(A.SELF_ON, False, True) is not True or A.holds(A.SELF_OFF, False, True) is not False \
        or A.holds(A.SELF_OFF, True, False) is not True or A.holds(None, True, True) is not True:
    fail.append("the self-appraisal's conditions read the appraisal's own setting")
for condition in {t.get("condition") for t in A.TRANSITIONS if t.get("condition")}:
    try:
        compile(condition, "condition", "eval")
    except SyntaxError:
        fail.append("the workflow's condition %r must be a Python expression Frappe can evaluate" % condition)
    if condition not in (A.IS_BSC, A.NOT_BSC, A.SELF_ON, A.SELF_OFF):
        fail.append("next_states does not know the condition %r" % condition)
if A.SELF_FIELD != "custom_self_appraisal" or A.SELF_FIELD not in A.SELF_ON:
    fail.append("the conditions read the appraisal's own Self-Appraisal")
if A.ROLE_WAITING.get(A.PENDING_SELF) != A.APPRAISEE:
    fail.append("the self-appraisal waits on the employee's own desk")
# the appraisals the supervisor does not have yet follow the setting as it is now
for state, own, wanted in ((A.DRAFT, 1, (1, A.DRAFT)), (A.DRAFT, 0, (0, A.DRAFT)), (None, 0, (0, A.DRAFT)),
                           (A.PENDING_SELF, 1, (1, A.PENDING_SELF)), (A.PENDING_SELF, 0, (0, A.PENDING_SUPERVISOR)),
                           (A.PENDING_SUPERVISOR, 0, None), (A.PENDING_SUPERVISOR, 1, None),
                           (A.PENDING_EMPLOYEE, 0, None), (A.COMPLETED, 1, None)):
    if A.follow_setting(state, own) != wanted:
        fail.append("%s with the self-appraisal %s: %s, not %s" % (state, "on" if own else "off", wanted,
                                                                  A.follow_setting(state, own)))
if set(A.BEFORE_SUPERVISOR) != {A.DRAFT, A.PENDING_SELF}:
    fail.append("only Draft and Pending Self-Appraisal come before the supervisor")
for form in A.FORM_TYPES:
    states = A.route(form, True)
    if any(states.index(state) >= states.index(A.PENDING_SUPERVISOR) for state in A.BEFORE_SUPERVISOR):
        fail.append("%s: Draft and the self-appraisal come before the supervisor" % form)
for own in (0, 1):
    for state in A.BEFORE_SUPERVISOR:
        flag, where = A.follow_setting(state, own)
        if where != state and (A.SELF, where) not in A.next_states(state, {"HR User"}, None, True):
            fail.append("an appraisal sent on when the setting changes goes where the workflow would send it")
        if where == state and flag != own:
            fail.append("an appraisal staying where it is takes the setting")
for state in A.PENDING_STATES:
    if not [t for t in A.TRANSITIONS if t["state"] == state and t["action"] == A.RETURN and t["next_state"] == A.DRAFT]:
        fail.append("%s must be able to return the appraisal to Draft" % state)
if [row for row in A.STATES if row["state"] == A.COMPLETED and row.get("doc_status") != "1"]:
    fail.append("the General Manager's approval submits the appraisal")
if "Production Manager" not in A.NEW_ROLES:
    fail.append("the Production Manager signs the form and appears nowhere else: the role must be created")
for role, ptypes in (("HR User", ("read", "write", "create", "submit")), ("Supervisor", ("read", "write")),
                     ("Production Manager", ("read", "write")), ("General Manager", ("read", "write"))):
    granted = (A.PERMISSIONS.get("Appraisal") or {}).get(role) or ()
    missing = [ptype for ptype in ptypes if ptype not in granted]
    if missing:
        fail.append("Frappe HR gives the Appraisal to the HR Manager alone: %s needs %s on it"
                    % (role, ", ".join(missing)))
    if granted and granted[0] != "read":
        fail.append("PERMISSIONS[Appraisal][%r] must start with read: workflows.py adds the rule with its first right"
                    % role)
granted_roles = {role for grants in A.PERMISSIONS.values() for role in grants}
if not granted_roles - {"HR User", "HR Manager", "Employee"} <= set(A.NEW_ROLES):
    fail.append("NEW_ROLES must name every granted role Frappe HR does not ship: %s"
                % sorted(granted_roles - {"HR User", "HR Manager", "Employee"} - set(A.NEW_ROLES)))

stamped = A.compute_stamps(A.PENDING_SUPERVISOR, A.PENDING_HRM, "sup@luuka", "2026-09-23", {})
if stamped["custom_supervisor_by"] != "sup@luuka" or stamped["custom_supervisor_on"] != "2026-09-23":
    fail.append("the block just passed is signed by whoever passed it: %s" % stamped)
if any(stamped[field] for field in ("custom_hrm_by", "custom_gm_by")):
    fail.append("only the block passed is signed")
if any(A.compute_stamps(A.PENDING_GM, A.DRAFT, "x", "2026-09-23",
                        {"custom_hrm_by": "a", "custom_supervisor_by": "b"}).values()):
    fail.append("a return to Draft clears every signature")
expect("returned without saying why", A.step_errors(A.PENDING_GM, A.DRAFT, {}), "Return Remarks")
expect("returned with a reason", A.step_errors(A.PENDING_GM, A.DRAFT, {"return_remarks": "Re-rate section B"}))
expect("the supervisor passing it on with no comments",
       A.step_errors(A.PENDING_SUPERVISOR, A.PENDING_HRM, {}), "Supervisor's general comments")
if A.next_states(A.PENDING_GM, ["General Manager"]) != [(A.APPROVE, A.COMPLETED), (A.RETURN, A.DRAFT)]:
    fail.append("the General Manager may approve or return: %s" % A.next_states(A.PENDING_GM, ["General Manager"]))
if A.next_states(A.PENDING_GM, ["HR User"]):
    fail.append("nobody but the General Manager acts on an appraisal waiting for them")
print("signatures: Employee to General Manager, returns, stamps, the rights each role needs")

# ── PIP rules ─────────────────────────────────────────────────────────
if P.DEFAULT_MONTHS != 3:
    fail.append("a Performance Improvement Plan runs three months unless HR says otherwise")
if P.end_date("2026-10-01", 3) != datetime.date(2026, 12, 31):
    fail.append("a three-month plan from 1 October runs to 31 December: %s" % P.end_date("2026-10-01", 3))
PIP = {"employee": "HR-EMP-1", "supervisor": "HR-EMP-2", "start_date": "2026-10-01", "end_date": "2026-12-31",
       "objectives": [{"area": "Output", "expected_standard": "95% of target", "measure": "Daily report",
                       "review_date": "2026-11-01"}]}
expect("a plan agreed", P.plan_errors(PIP))
expect("a plan with nothing to improve", P.plan_errors(dict(PIP, objectives=[])), "List what must improve")
expect("a line with no standard",
       P.plan_errors(dict(PIP, objectives=[dict(PIP["objectives"][0], expected_standard="")])), "expected standard")
expect("nobody guiding the employee", P.plan_errors(dict(PIP, supervisor=None)), "supervisor who will guide")
expect("a review date after the plan ends",
       P.plan_errors(dict(PIP, objectives=[dict(PIP["objectives"][0], review_date="2027-02-01")])),
       "falls after the plan ends")
expect("closing with nothing recorded", P.close_errors({}), "at least one review", "how the plan ended",
       "what was agreed")
expect("closing properly", P.close_errors({"reviews": [{"progress": P.MET}], "outcome": P.IMPROVED,
                                           "remarks": "Back to standard."}))
expect("extending with no new date", P.close_errors({"reviews": [{"progress": P.PARTLY}], "outcome": P.EXTENDED,
                                                     "remarks": "More time.", "end_date": "2026-12-31"}),
       "new end date")
if P.suggested_outcome([{"progress": P.MET}, {"progress": P.MET}, {"progress": P.NOT_MET}]) != P.IMPROVED:
    fail.append("two of three met is most of them, and suggests Improved")
if P.suggested_outcome([{"progress": P.NOT_MET}, {"progress": P.MET}]) != P.NOT_IMPROVED:
    fail.append("half met is not most, and does not suggest Improved")
if P.suggested_outcome([{"progress": P.PARTLY}, {"progress": P.PARTLY}]) != P.NOT_IMPROVED:
    fail.append("partly met is not met")
if P.suggested_outcome([]) is not None:
    fail.append("nothing reviewed suggests nothing")
if [row["area"] for row in P.due_reviews([{"area": "A", "review_date": "2026-11-01"},
                                          {"area": "B", "review_date": "2026-12-01"},
                                          {"area": "C", "review_date": "2026-11-01", "reviewed_on": "2026-11-02"}],
                                         "2026-11-15")] != ["A"]:
    fail.append("a review is due once its date has come and nobody has looked at it")
print("improvement plans: the agreement, the reviews, the outcome it suggests")

# ── 3. The forms ──────────────────────────────────────────────────────
PAPER = {
    "Appraisal Plan": ["year", "company", "branch", "quarters", "appraise_all", "employees", "status"],
    "Performance Review": ["appraisal_cycle", "review_date", "employees", "appraised", "average_score",
                           "below_pass", "completion", "shared_with", "shared_on", "management_remarks",
                           "decided_by", "decided_on", "status"],
    "Performance Improvement Plan": ["employee", "supervisor", "appraisal", "appraisal_score", "performance_review",
                                     "start_date", "months", "end_date", "objectives", "reviews", "outcome",
                                     "suggested_outcome", "new_end_date", "outcome_remarks",
                                     "employee_agreed_on", "supervisor_agreed_on", "status"],
}
for name, wanted in PAPER.items():
    spec = doctype(name)
    if not spec:
        fail.append("%s is not there" % name)
        continue
    fields = fields_of(spec)
    for fieldname in wanted:
        if fieldname not in fields:
            fail.append("%s has no %s, which the process asks for" % (name, fieldname))
    if not spec.get("is_submittable"):
        fail.append("%s must be submittable: it is approved, and cancelling undoes what it did" % name)
    roles = {row["role"] for row in spec.get("permissions", [])}
    if not {"HR Manager", "HR User"} <= roles:
        fail.append("%s: the HR Manager and the branch HR Officer must be able to work it (%s)" % (name, sorted(roles)))

quarter = fields_of(doctype("Appraisal Plan Quarter"))
for fieldname in ("quarter", "from_date", "to_date", "soft_deadline", "hard_deadline", "appraisal_cycle",
                  "notified_on", "reminders_sent"):
    if fieldname not in quarter:
        fail.append("Appraisal Plan Quarter has no %s" % fieldname)
if (quarter.get("quarter") or {}).get("options", "").split("\n") != list(R.QUARTERS):
    fail.append("Appraisal Plan Quarter.quarter must offer exactly %s" % (R.QUARTERS,))

review_row = fields_of(doctype("Performance Review Employee"))
for fieldname in ("employee", "appraisal", "total_score", "band", "recommended", "decision", "position_change",
                  "improvement_plan"):
    if fieldname not in review_row:
        fail.append("Performance Review Employee has no %s" % fieldname)
# A choice someone has to make starts blank: Frappe gives a Select with no
# default its first option on every new row, which made every decision a
# Promotion, every rating 1, every point Met and every plan Improved
# before anyone chose (Oct 2026)
if (review_row.get("decision") or {}).get("options", "").split("\n") != [""] + list(R.DECISIONS):
    fail.append("the decision must start blank and offer exactly what management may decide: %s" % (R.DECISIONS,))

for name, wanted in (("Appraisal Factor Rating", ("item", "employee_rating", "supervisor_rating", "supervisor_comment")),
                     ("Appraisal Objective Rating", ("item", "employee_rating", "supervisor_rating", "supervisor_comment"))):
    fields = fields_of(doctype(name))
    for fieldname in wanted:
        if fieldname not in fields:
            fail.append("%s has no %s: the form is rated by the employee and the supervisor" % (name, fieldname))
    for side in ("employee_rating", "supervisor_rating"):
        if fields.get(side, {}).get("options", "").split("\n") != [""] + list(R.RATINGS):
            fail.append("%s.%s must start blank and offer the form's scale %s" % (name, side, (R.RATINGS,)))
        if fields.get(side, {}).get("default"):
            fail.append("%s.%s must not rate a row before anyone has" % (name, side))

pip_row = fields_of(doctype("PIP Objective"))
for fieldname in ("area", "expected_standard", "support", "measure", "review_date", "progress", "reviewed_on"):
    if fieldname not in pip_row:
        fail.append("PIP Objective has no %s" % fieldname)
for name in ("PIP Objective", "PIP Review"):
    if (fields_of(doctype(name)).get("progress") or {}).get("options", "").split("\n") != [""] + list(P.PROGRESS):
        fail.append("%s.progress must start blank and offer exactly %s" % (name, (P.PROGRESS,)))
if (fields_of(doctype("Performance Improvement Plan")).get("outcome") or {}).get("options", "").split("\n") != [""] + list(P.OUTCOMES):
    fail.append("the plan's outcome must start blank and offer exactly %s" % (P.OUTCOMES,))
for name, fieldname in (("Performance Review Employee", "decision"), ("PIP Objective", "progress"),
                        ("PIP Review", "progress"), ("Performance Improvement Plan", "outcome")):
    if (fields_of(doctype(name)).get(fieldname) or {}).get("default"):
        fail.append("%s.%s must not be chosen before anyone has" % (name, fieldname))

# LPL/HR/18 on Frappe HR's Appraisal
appraisal = all_fields("Appraisal")
if not upstream_doctype("Appraisal"):
    fail.append("Frappe HR no longer ships an Appraisal to build the round on")
for fieldname in ("custom_factors", "custom_objectives", "custom_factors_score", "custom_objectives_score",
                  "custom_total_score", "custom_band", "custom_annual_score", "custom_plan", "custom_quarter",
                  "custom_supervisor", "custom_appraisal_status", "custom_return_remarks", "custom_outcome",
                  "custom_performance_review", *A.ALL_STAMP_FIELDS,
                  *[field for field, _who in A.REMARK_FIELDS.values()],
                  *["custom_%s" % field for field, _question in R.QUESTIONS]):
    if fieldname not in appraisal:
        fail.append("the Appraisal has no %s, which LPL/HR/18 asks for" % fieldname)
ours = custom_fields("Appraisal")
if (ours.get("custom_appraisal_status") or {}).get("options", "").split("\n") != list(R.STATUSES):
    fail.append("Appraisal.custom_appraisal_status must follow the workflow's states")
for fieldname in ("custom_factors_score", "custom_objectives_score", "custom_total_score", "custom_band",
                  "custom_annual_score", *A.ALL_STAMP_FIELDS):
    if not (ours.get(fieldname) or {}).get("read_only"):
        fail.append("Appraisal.%s is worked out, not typed: it must be read-only" % fieldname)
if (ours.get("custom_factors") or {}).get("options") != "Appraisal Factor Rating":
    fail.append("Section A must be the Appraisal Factor Rating table")
if (ours.get("custom_objectives") or {}).get("options") != "Appraisal Objective Rating":
    fail.append("Section B must be the Appraisal Objective Rating table")
cycle = custom_fields("Appraisal Cycle")
for fieldname in ("custom_plan", "custom_quarter", "custom_soft_deadline", "custom_hard_deadline",
                  "custom_reminders_sent"):
    if fieldname not in cycle:
        fail.append("the Appraisal Cycle has no %s: the round needs its plan and its deadlines" % fieldname)
print("forms: the plan, the review, the improvement plan, and LPL/HR/18 on Frappe HR's Appraisal")

# ── 4. The glue ───────────────────────────────────────────────────────
glue = read("hrms_addon", "hrms_addon", "appraisals.py")
pips = read("hrms_addon", "hrms_addon", "pips.py")
own = {name: fields_of(doctype(name)) for name in PAPER}


def body_of(source, name):
    start = source.index("\ndef %s(" % name)
    rest = source[start + 1:]
    end = rest.index("\ndef ", 1) if "\ndef " in rest[1:] else len(rest)
    return rest[:end]


for fieldname in sorted(set(re.findall(r'doc\.get\("(custom_\w+)"\)', glue))
                        | set(re.findall(r'doc\.(custom_\w+)\b', glue))):
    if fieldname not in appraisal:
        fail.append("appraisals.py reads or writes Appraisal.%s, which does not exist" % fieldname)
upstream_appraisal = fields_of(upstream_doctype("Appraisal") or {})
for fieldname in ("final_score", "total_score", "self_score", "appraisal_cycle", "employee", "start_date", "end_date"):
    if upstream_appraisal and fieldname not in upstream_appraisal:
        fail.append("appraisals.py sets Appraisal.%s, which Frappe HR no longer has" % fieldname)
upstream_cycle = fields_of(upstream_doctype("Appraisal Cycle") or {})
for fieldname in ("cycle_name", "company", "start_date", "end_date", "status", "branch", "department",
                  "kra_evaluation_method"):
    if upstream_cycle and fieldname not in upstream_cycle:
        fail.append("the cycle is opened with %s, which Frappe HR no longer has" % fieldname)
for needle, why in (
    ('"doctype": "Appraisal Cycle"', "the quarter opens one of Frappe HR's cycles, so the round keeps its chart"),
    ('"doctype": "Appraisal",', "an Appraisal is raised per employee"),
    ("rules.scores(", "Section C comes from the rules"),
    ("doc.final_score = ", "Frappe HR's own score is written too, or the Appraisal Overview chart stays empty"),
    ("approval.compute_stamps(", "the signatures are stamped by the workflow"),
    ("bsc_rules.year_to_date(", "the year to date is the average of the quarters appraised, on either form"),
    ("_year_so_far(doc)", "each appraisal shows the year so far, quarter by quarter"),
    ("rules.due_quarters(", "the HR Officer is told when a quarter closes"),
    ("rules.reminders_due(", "everyone appraising is reminded before the deadlines"),
    ("sheet.build(", "the sheet is Luuka's own form, built by appraisal_sheet.py"),
    ("sheet.read(", "and read back by it"),
    ('frappe.response["type"] = "binary"', "the sheet comes down as a file"),
    ('frappe.response["filecontent"] = content', "carrying the workbook built"),
    ("settings().kra_evaluation_method", "the plan's cycles score KRAs as Appraisal Settings say"),
    ("rules.rates_goals_manually(cycle.get(\"kra_evaluation_method\"))",
     "an appraisal rates its goals by hand only in a cycle rated by hand"),
    ('"custom_self_appraisal": settings().self_appraisal', "an appraisal is raised with the setting as it is"),
    ("own = settings().self_appraisal\n    _action, state = approval.opening(own)",
     "a draft is sent on as the setting says now, not as it said when the draft was made"),
    ("appraisal.set(approval.SELF_FIELD, own)", "and carries that setting from then on"),
    ("wanted = approval.follow_setting(row.get(approval.STATE_FIELD), own)",
     "a saved setting reaches every appraisal the supervisor does not have yet"),
    ('frappe.db.set_value("Appraisal", row.name, approval.SELF_FIELD, flag)\n',
     "a draft that follows the setting is marked modified, so a form opened before is reloaded before it is saved"),
    ('frappe.get_all("Appraisal", filters={"docstatus": 0},', "only open appraisals follow a new setting"),
    ('people.withdraw("Appraisal", name, [frappe.db.get_value("Employee", doc.employee, "user_id")])',
     "the employee's self-appraisal task is withdrawn when it is no longer asked for"),
    ("_tell_next(doc, state)", "and the supervisor told the appraisal is theirs"),
    ('approval.STATE_FIELD: approval.PENDING_SELF})', "the cycle's Self Appraisal Pending counts those waiting on "
                                                     "the employee"),
    ("approval.opening(", "a raised appraisal is sent on to the employee or the supervisor"),
    ('frappe.new_doc("Employee Position Change")', "a promotion or an increase is an Employee Position Change"),
    ('frappe.new_doc("Performance Improvement Plan")', "a PIP decision raises the plan"),
    ("people.hr_officers(", "the branch HR Officer is told"),
):
    if needle not in glue:
        fail.append("appraisals.py: %s (%r not found)" % (why, needle))
for needle, why in (
    ("rules.plan_errors(", "the plan is judged by the rules"),
    ("rules.close_errors(", "and so is closing it"),
    ("rules.suggested_outcome(", "the reviews suggest an outcome"),
    ("rules.due_reviews(", "a review date that has come is chased"),
    ('doc.db_set({"status": rules.CLOSED', "submitting closes the plan"),
):
    if needle not in pips:
        fail.append("pips.py: %s (%r not found)" % (why, needle))
if 'doc.check_permission("submit")' not in body_of(glue, "open_quarter"):
    fail.append("open_quarter raises documents for everyone in scope: it must check the caller may submit the plan")
if "reason = _not_taken(" not in body_of(glue, "upload_sheet"):
    fail.append("upload_sheet writes onto appraisals: each sheet is judged before anything is written")
not_taken = body_of(glue, "_not_taken")
for needle, why in (
    ('frappe.has_permission("Appraisal", "write", name)', "the caller may write the appraisal"),
    ("if doc.docstatus != 0:", "a completed or cancelled appraisal is left alone"),
    ("doc.appraisal_cycle != appraisal_cycle", "a sheet of another cycle is not taken from this one"),
    ("if appraisal and name != appraisal:", "an appraisal's own upload takes only its own sheet"),
    ('values.get("period")', "a scorecard sheet is taken only for the period it was downloaded for"),
    ("SUPERVISOR_STATES + (approval.PENDING_EMPLOYEE,)", "an appraisal past its rating is changed on the system"),
):
    if needle not in not_taken:
        fail.append("the upload checks %s (%r not found in _not_taken)" % (why, needle))
if glue.count("doc.has_permission(\"read\")") < 1 or "doc.has_permission(\"read\")" not in body_of(glue,
                                                                                                "download_sheet"):
    fail.append("download_sheet gives only the appraisals the caller may read")
for name in ("fill_year", "get_appraisals", "download_sheet"):
    if not re.search(r"@frappe\.whitelist\(\)\ndef %s\(" % name, glue):
        fail.append("appraisals.%s must be whitelisted for the form" % name)
for source, module, name in ((glue, "appraisals", "open_quarter"), (glue, "appraisals", "upload_sheet"),
                             (glue, "appraisals", "apply_template"), (glue, "appraisals", "send_drafts"),
                             (glue, "appraisals", "share_with_management"), (pips, "pips", "start")):
    if not re.search(r'@frappe\.whitelist\(methods=\["POST"\]\)\ndef %s\(' % name, source):
        fail.append("%s.%s changes something: a whitelisted POST method" % (module, name))
print("glue: fields that exist here and upstream, the rules followed, the buttons' methods whitelisted")

# ── 5. The print-outs ─────────────────────────────────────────────────
PRINTS = {
    "Supervisory Skills Evaluation Form": ("Appraisal", [
        "LPL/HR/18", "Section A: Ratable Factors", "Section B", "Section C", "Ratable Factors /60",
        "Objectives/KPIs /40", "Total Score", "General comments by the Employee",
        "Human Resources Manager&rsquo;s remarks".replace("&rsquo;", "'"), "Production Manager's remarks",
        "General Manager's remarks", "Rating Scale:"]),
    "Performance Appraisal Report": ("Performance Review", [
        "Performance Appraisal Report", "Average Score", "Below the pass mark", "Decision", "Management's Remarks"]),
    "Performance Improvement Plan Document": ("Performance Improvement Plan", [
        "Performance Improvement Plan", "Area to improve", "Expected standard", "Support from the company",
        "How it is measured", "Review meetings", "Outcome", "Agreement"]),
}
for name, (doc_type, needles) in PRINTS.items():
    spec = print_format(name)
    if not spec:
        fail.append("the %s print format is not there" % name)
        continue
    if spec.get("doc_type") != doc_type:
        fail.append("%s must print a %s, not a %s" % (name, doc_type, spec.get("doc_type")))
    html = spec.get("html") or ""
    for needle in needles:
        if needle not in html:
            fail.append("%s does not say %r, which the paper does" % (name, needle))
    available = all_fields(doc_type)
    for fieldname in sorted(set(re.findall(r"doc\.(\w+)", html))):
        if fieldname in ("name", "creation", "owner", "modified", "doctype", "docstatus", "get"):
            continue
        if fieldname not in available:
            fail.append("%s prints doc.%s, which the %s does not have" % (name, fieldname, doc_type))
    for bad in re.findall(r"\{\{\s*doc\.(\w+)\s*\}\}", html):
        fail.append("%s prints doc.%s unescaped: use v()" % (name, bad))
print("print-outs: LPL/HR/18, the appraisal report and the improvement plan, from fields that exist")

# ── 6. Wiring ─────────────────────────────────────────────────────────
hooks = {}
for node in ast.parse(read("hrms_addon", "hooks.py")).body:
    if isinstance(node, ast.Assign) and isinstance(node.targets[0], ast.Name):
        try:
            hooks[node.targets[0].id] = ast.literal_eval(node.value)
        except ValueError:
            pass
events = hooks.get("doc_events") or {}
for doctype_name, event, function in (("Appraisal", "validate", "appraisals.appraisal_validate"),
                                      ("Appraisal", "on_cancel", "appraisals.appraisal_on_cancel")):
    if (events.get(doctype_name) or {}).get(event) != "hrms_addon.hrms_addon.%s" % function:
        fail.append("doc_events %s %s must be %s" % (doctype_name, event, function))
if "hrms_addon.hrms_addon.appraisals.setup_workflows_on_migrate" not in (hooks.get("after_migrate") or []):
    fail.append("after_migrate must build the appraisal workflow and its roles")
for job in ("hrms_addon.hrms_addon.appraisals.daily", "hrms_addon.hrms_addon.pips.daily"):
    if job not in ((hooks.get("scheduler_events") or {}).get("daily") or []):
        fail.append("the scheduler must run %s" % job)
if "hrms_addon.hrms_addon.pick_lists.seed_appraisal_masters" not in (hooks.get("after_install") or []):
    fail.append("a fresh install seeds Section A as the Appraisal Factor list")
for doctype_name, path in (("Appraisal", "public/js/appraisal.js"), ("Appraisal Cycle", "public/js/appraisal_cycle.js")):
    if (hooks.get("doctype_js") or {}).get(doctype_name) != path or not os.path.exists(os.path.join(PACKAGE, path)):
        fail.append("doctype_js %s must load %s" % (doctype_name, path))
patches = read("hrms_addon", "patches.txt").split("[post_model_sync]")[1]
if "hrms_addon.patches.v1_0.seed_performance" not in patches:
    fail.append("a site that has the app already must get the factors and the roles by patch")
seed = read("hrms_addon", "patches", "v1_0", "seed_performance.py")
if "seed_appraisal_masters()" not in seed or "appraisal_approval.NEW_ROLES" not in seed:
    fail.append("the patch must seed the factor list and exactly the roles the workflow declares")
if "seed_masters(appraisal_rules.APPRAISAL_MASTERS)" not in read("hrms_addon", "hrms_addon", "pick_lists.py"):
    fail.append("pick_lists.seed_appraisal_masters seeds appraisal_rules.APPRAISAL_MASTERS")
for name, module, prefix, methods in (
        ("appraisal_plan", "appraisals", "plan", ("validate", "on_submit", "on_cancel")),
        ("performance_review", "appraisals", "review", ("validate", "on_submit", "on_cancel")),
        ("performance_improvement_plan", "pips", "plan", ("validate", "on_submit", "on_cancel"))):
    controller = open(os.path.join(APP, "doctype", name, name + ".py"), encoding="utf-8").read()
    for method in methods:
        if "    def %s(self):\n        %s.%s_%s(self)" % (method, module, prefix, method) not in controller:
            fail.append("the %s controller must hand %s to %s.%s_%s" % (name, method, module, prefix, method))
navigation = load("navigation_rules")
carded = {link[1] for cards in navigation.CARDS.values() for _card, links in cards for link in links}
sidebarred = {entry[1] for entries in navigation.SIDEBAR.values() for entry in entries}
for name in list(PAPER) + ["Appraisal Factor"]:
    if name not in carded:
        fail.append("%s is on no workspace card: it could only be found by searching for its DocType" % name)
    if name not in sidebarred:
        fail.append("%s is in no sidebar" % name)
if "Performance" not in navigation.CARDS:
    fail.append("Frappe HR ships the Performance page bare: our cards are what fills it")
conn = read("hrms_addon", "hrms_addon", "connections.py")
for name in ("Appraisal", "Performance Improvement Plan"):
    if name not in conn:
        fail.append("%s must show on the Employee's Connections" % name)
for name in PAPER:
    folder = name.lower().replace(" ", "_")
    if not os.path.exists(os.path.join(APP, "doctype", folder, folder + "_dashboard.py")):
        fail.append("%s has no Connections of its own (%s_dashboard.py)" % (name, folder))
print("wiring: the doc events, the workflow on migrate, the daily jobs, the seed, the patch, the way in")


# ── 7. The balanced scorecard, beside the supervisory form ────────────
if S.OBJECTIVES_WEIGHT != 80 or S.COMPETENCIES_WEIGHT != 20:
    fail.append("the scorecard is 80 objectives and 20 competencies")
if S.BANDS != ((90, "Excellent"), (80, "Very Good"), (70, "Good"), (60, "Fair"), (0, "Poor")):
    fail.append("the scorecard's own scale is 90, 80, 70, 60 — not the supervisory form's")
if S.BANDS == R.BANDS:
    fail.append("the two forms band differently; they must not share one scale")
# Luuka, 4 Oct 2026: four quarters, the year to date their average, in
# place of the workbook's three and an annual score out of ten
if S.QUARTERS != ("Q1", "Q2", "Q3", "Q4"):
    fail.append("the scorecard records four quarters: %s" % (S.QUARTERS,))
for gone in ("annual_score", "field_for", "PERIODS", "PERCENT_PERIODS", "SCORE_PERIODS"):
    if hasattr(S, gone):
        fail.append("the annual score out of ten is gone: bsc_rules.%s is still there" % gone)
if set(S.PERSPECTIVES) != {"Financial", "Customer / Stakeholder", "Internal Business Processes", "Learning & Growth"}:
    fail.append("the four balanced scorecard perspectives are the job descriptions' own")
if len(S.COMPETENCIES) != 5 or sum(weight for _n, _i, weight in S.COMPETENCIES) != S.COMPETENCIES_WEIGHT:
    fail.append("the seeded competencies must weigh %d in total" % S.COMPETENCIES_WEIGHT)
if S.FORM_TYPES != (S.FORM_SUPERVISORY, S.FORM_BSC) or A.FORM_TYPES != S.FORM_TYPES:
    fail.append("the rules and the workflow must name the two forms the same way")

# a KPI scores its own weight times the percentage achieved: Luuka's
# workbook divides by a further ten, which they confirmed is not meant
if S.quarter_score(25, 100) != 25.0 or S.quarter_score(12.5, 80) != 10.0:
    fail.append("a KPI scores its weight times the percentage achieved: %s, %s"
                % (S.quarter_score(25, 100), S.quarter_score(12.5, 80)))
if (S.percent_field("Q2"), S.score_field("Q3"), S.comments_field("Q4")) != ("q2_percent", "q3_score", "q4_comments"):
    fail.append("each quarter has its percentage, its weighted score and its comments")
if [S.quarter_of(month) for month in (1, 3, 4, 6, 7, 9, 10, 12)] != ["Q1", "Q1", "Q2", "Q2", "Q3", "Q3", "Q4", "Q4"]:
    fail.append("a month falls in its calendar quarter")
CARD_KPIS = [{"perspective": "Financial", "kpi": "Savings", "weight": 12.5},
             {"perspective": "Financial", "kpi": "Variance", "weight": 12.5},
             {"perspective": "Customer / Stakeholder", "kpi": "Lead times", "weight": 15},
             {"perspective": "Internal Business Processes", "kpi": "Quotations", "weight": 20},
             {"perspective": "Internal Business Processes", "kpi": "Approvals", "weight": 10},
             {"perspective": "Learning & Growth", "kpi": "CIPS", "weight": 10}]
WHOLE_CARD = [dict(row, q1_percent=100, q2_percent=50) for row in CARD_KPIS]
if S.section_a(WHOLE_CARD, "Q1") != 80.0:
    fail.append("a quarter fully achieved scores the whole 80: %s" % S.section_a(WHOLE_CARD, "Q1"))
if S.section_a(WHOLE_CARD, "Q2") != 40.0:
    fail.append("half achieved scores half: %s" % S.section_a(WHOLE_CARD, "Q2"))
if S.section_a(WHOLE_CARD, "Q3") is not None:
    fail.append("a quarter with nothing recorded scores nothing at all")
if S.section_a([], "Q1") is not None or S.quarter_score(25, None) is not None:
    fail.append("nothing recorded is nothing, never a zero that drags the score down")
SHARED = [dict(row, q2_percent=50) for row in S.spread_weights([{"perspective": "Financial", "weight": 25},
                                                                 {"perspective": "Financial"},
                                                                 {"perspective": "Financial"}])]
if [S.quarter_score(row["weight"], 50) for row in SHARED] != [4.17, 4.17, 4.17] \
        or S.section_a(SHARED, "Q2") != 12.5 or S.perspective_summary(SHARED)[0]["q2_score"] != 12.5:
    fail.append("totals add the KPIs' exact scores and round once: 8.33, 8.33 and 8.34 half achieved make 12.5, "
                "not 12.51: %s" % S.section_a(SHARED, "Q2"))
summary = S.perspective_summary(WHOLE_CARD)
if [(row["perspective"], row["weight"], row["q1_score"], row["q2_score"], row["q3_score"], row["year_to_date"])
        for row in summary] != [("Financial", 25.0, 25.0, 12.5, None, 18.75),
                                ("Customer / Stakeholder", 15.0, 15.0, 7.5, None, 11.25),
                                ("Internal Business Processes", 30.0, 30.0, 15.0, None, 22.5),
                                ("Learning & Growth", 10.0, 10.0, 5.0, None, 7.5)]:
    fail.append("the perspectives below the KPIs sum them up: their weight, each quarter, their year to date: %s"
                % summary)
if S.perspective_summary([{"perspective": None, "weight": 5}]) != [] or S.perspective_summary([]) != []:
    fail.append("a KPI under no perspective sums into none")
if (S.year_to_date([80, None, 70]), S.year_to_date([0, 90]), S.year_to_date([]), S.year_to_date([None])) \
        != (75.0, 45.0, None, None):
    fail.append("the year to date averages the quarters appraised, a quarter that scored nothing included: %s"
                % ((S.year_to_date([80, None, 70]), S.year_to_date([0, 90])),))
# Frappe keeps a number left blank as 0 (Float and Percent are NOT NULL in
# v16): a column that is 0 throughout was never filled in, and once any KPI
# has a figure, a 0 beside it is a real 0
ZEROS = [dict(row, q2_percent=0) for row in CARD_KPIS]
SOME = [dict(row, q2_percent=(80 if index == 0 else 0)) for index, row in enumerate(CARD_KPIS)]
if (S.recorded(ZEROS, "q2_percent"), S.recorded(SOME, "q2_percent"), S.recorded([], "q2_percent")) \
        != (False, True, False):
    fail.append("a column is filled in once any KPI has a figure in it")
if S.section_a(ZEROS, "Q2") is not None or S.section_a(SOME, "Q2") != 10.0:
    fail.append("a quarter left at 0 throughout is not scored; one with a figure counts its 0s as 0%%: %s, %s"
                % (S.section_a(ZEROS, "Q2"), S.section_a(SOME, "Q2")))
if [row["q2_score"] for row in S.perspective_summary(SOME)] != [10.0, 0.0, 0.0, 0.0] \
        or {row["q2_score"] for row in S.perspective_summary(ZEROS)} != {None}:
    fail.append("a perspective scores 0 in a quarter filled in, nothing in one never filled in: %s"
                % [row["q2_score"] for row in S.perspective_summary(SOME)])
expect("a quarter left at 0 throughout is not scored: every KPI named",
       S.appraisal_errors({"step": "appraiser", "quarter": "Q2", "kpis": ZEROS, "competencies": []}),
       "Record Q2 percentage achieved for every KPI: Savings (Financial)")
expect("one filled in takes its 0s as 0%",
       S.appraisal_errors({"step": "appraiser", "quarter": "Q2", "kpis": SOME, "competencies": []}))
if S.section_b([{"weight": 4, "score": 0}, {"weight": 16, "score": 0}]) is not None \
        or S.section_b([{"weight": 4, "score": 0}, {"weight": 16, "score": 5}]) != 8.0:
    fail.append("Section B left at 0 throughout is not scored; scored, its 0s are 0")
expect("competencies left at 0 throughout are not scored",
       S.appraisal_errors({"step": "appraiser", "quarter": "Q2", "kpis": SOME,
                           "competencies": [{"competency": "One", "score": 0}]}), "Score every competency")
WHOLE_B = [{"weight": weight, "score": 10} for _n, _i, weight in S.COMPETENCIES]
if S.section_b(WHOLE_B) != 20.0:
    fail.append("every competency at ten scores the whole 20: %s" % S.section_b(WHOLE_B))
if S.overall(80.0, 20.0) != 100.0 or S.band(S.overall(80.0, 20.0)) != "Excellent":
    fail.append("80 and 20 make 100, which is Excellent")
for total, name in ((95, "Excellent"), (90, "Excellent"), (85, "Very Good"), (80, "Very Good"), (75, "Good"),
                    (70, "Good"), (65, "Fair"), (60, "Fair"), (59.9, "Poor"), (0, "Poor")):
    if S.band(total) != name:
        fail.append("the scorecard rates %s as %s, not %r" % (total, name, S.band(total)))
if S.band(None) is not None:
    fail.append("nothing scored has no rating")
if set(S.BAND_MEANING) != {name for _floor, name in S.BANDS}:
    fail.append("every band must carry the words the form prints beside it")

CARD = {"designation": "Procurement Manager", "kpis": CARD_KPIS,
        "competencies": [{"competency": n, "weight": w} for n, _i, w in S.COMPETENCIES]}
expect("a whole scorecard", S.template_errors(CARD))
expect("no role", S.template_errors(dict(CARD, designation=None)), "Name the role")
expect("KPIs whose weights do not total 80",
       S.template_errors(dict(CARD, kpis=[{"perspective": "Financial", "kpi": "Savings", "weight": 90}])),
       "must total 80")
expect("a KPI with no weight, the total still 80",
       S.template_errors(dict(CARD, kpis=CARD_KPIS + [{"perspective": "Financial", "kpi": "Unweighed"}])),
       "Give every KPI its weight: Unweighed (Financial)")
expect("a KPI under no perspective",
       S.template_errors(dict(CARD, kpis=CARD_KPIS[:-1] + [{"kpi": "Stray", "weight": 10}])),
       "Put every KPI under one of the perspectives: Stray (no perspective)")
expect("a negative weight",
       S.template_errors(dict(CARD, kpis=CARD_KPIS + [{"perspective": "Financial", "kpi": "Minus", "weight": -5},
                                                      {"perspective": "Financial", "kpi": "Plus", "weight": 5}])),
       "cannot be negative: Minus (Financial)")
expect("competencies that do not total 20",
       S.template_errors(dict(CARD, competencies=[{"competency": "One", "weight": 5}])), "must total 20")
expect("no KPIs", S.template_errors(dict(CARD, kpis=[])), "List the KPIs")
many = [{"perspective": "Financial", "kpi": "K%d" % number} for number in range(8)]
expect("eight KPIs with no weight: five named, the rest counted",
       S.template_errors(dict(CARD, kpis=many)), "K0 (Financial), K1 (Financial), K2 (Financial), K3 (Financial), "
       "K4 (Financial) and 3 more", "not 0")

SCORED = {"step": "appraiser", "quarter": "Q1",
          "kpis": [dict(row, q1_percent=90) for row in CARD_KPIS],
          "competencies": [{"competency": n, "score": 8} for n, _i, _w in S.COMPETENCIES]}
expect("a quarter scored", S.appraisal_errors(SCORED))
expect("not the appraiser's step", S.appraisal_errors(dict(SCORED, step=None)))
expect("a KPI left blank",
       S.appraisal_errors(dict(SCORED, kpis=[{"perspective": "Financial", "kpi": "Savings", "weight": 20}])),
       "Record Q1 percentage achieved for every KPI: Savings (Financial)")
expect("a percentage over a hundred",
       S.appraisal_errors(dict(SCORED, kpis=[{"perspective": "Financial", "kpi": "Savings", "weight": 20,
                                              "q1_percent": 140}])),
       "from 0 to 100")
expect("another quarter's figure does not stand for this one's",
       S.appraisal_errors(dict(SCORED, kpis=[{"perspective": "Financial", "kpi": "Savings", "weight": 20,
                                              "q2_percent": 90}])),
       "Record Q1 percentage achieved")
expect("no quarter", S.appraisal_errors(dict(SCORED, quarter=None)), "Say which quarter")
expect("a competency unscored",
       S.appraisal_errors(dict(SCORED, competencies=[{"competency": "One"}])), "Score every competency")
expect("no scorecard at all", S.appraisal_errors(dict(SCORED, kpis=[])), "no KPIs")

# Luuka's own workbook, read as openpyxl hands it over
SHEET = [
    ["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
    ["Role / Position:", "", "Procurement Manager", "", "", "", "", "Department:", "", "", "Procurement"],
    ["Employee Name:", "", "", "", "", "", "", "Grade:", "", "", "G15"],
    ["", "", "", "", "", "", "", "Review Period:", "", "", "January - December 2026"],
    ["SECTION A  —  OBJECTIVES & KPIs"],
    ["#", "BSC Perspective", "KPI / Objective", "Timing", "Weight"],
    ["1", "Financial", "Zero stock-outs", "Monthly", 25],
    ["2", S.CONTINUATION, "Savings of 5%", "Monthly", None],
    ["3", "Customer / Stakeholder", "Lead times met", "Monthly", 15],
    ["4", "Internal Process", "Three quotations", "Weekly", 30],
    ["5", "Learning & Growth", "Report by the 5th", "Ongoing", 10],
    ["WEIGHT CHECK & QUARTERLY TOTALS", "", "", "", 80],
    ["SECTION B  —  COMPETENCIES"],
    ["Competency", "", "", "Behavioural Indicators", "", "", "", "", "", "Weight"],
    ["Job Knowledge & Technical Excellence", "", "", "Deep expertise", "", "", "", "", "", 4],
    ["Compliance & Governance", "", "", "Champions policy", "", "", "", "", "", 3],
    ["Commitment & Results Delivery", "", "", "High drive", "", "", "", "", "", 4],
    ["Strategic Planning & Decision-Making", "", "", "Clear priorities", "", "", "", "", "", 4],
    ["Inclusive Leadership & Team Development", "", "", "Coaches reports", "", "", "", "", "", 5],
    ["COMPETENCY WEIGHT CHECK & SECTION B SCORE", "", "", "", "", "", "", "", "", 20],
]
from_sheet = S.parse_sheet(SHEET)
if from_sheet["role"] != "Procurement Manager" or from_sheet["grade"] != "G15":
    fail.append("the sheet's header must be read: %s"
                % {k: from_sheet[k] for k in ("role", "grade", "department")})
if [(row["perspective"], row["weight"]) for row in from_sheet["perspectives"]] != [
        ("Financial", 25.0), ("Customer / Stakeholder", 15.0),
        ("Internal Business Processes", 30.0), ("Learning & Growth", 10.0)]:
    fail.append("the weights the sheet writes once per perspective are read, and Internal Process is named as the "
                "JDs name it: %s" % from_sheet["perspectives"])
if len(from_sheet["kpis"]) != 5:
    fail.append("every KPI is read, including the ones marked with the continuation arrow: %s" % len(from_sheet["kpis"]))
if from_sheet["kpis"][1]["perspective"] != "Financial":
    fail.append("a KPI under the continuation arrow belongs to the perspective above it: %s" % from_sheet["kpis"][1])
if len(from_sheet["competencies"]) != 5 or from_sheet["competencies"][0]["weight"] != 4.0:
    fail.append("Section B is read with its indicators and weights: %s" % from_sheet["competencies"])
if [row["weight"] for row in from_sheet["kpis"]] != [25.0, None, 15.0, 30.0, 10.0]:
    fail.append("each KPI keeps the weight written beside it, the perspective's on its first KPI: %s"
                % [row["weight"] for row in from_sheet["kpis"]])
shared = S.spread_weights(from_sheet["kpis"])
if [row["weight"] for row in shared] != [12.5, 12.5, 15.0, 30.0, 10.0]:
    fail.append("a perspective the sheet weighs once has its weight shared between its KPIs: %s"
                % [row["weight"] for row in shared])
expect("Luuka's own sheet, its weights shared out", S.template_errors(
    {"designation": from_sheet["role"], "kpis": shared, "competencies": from_sheet["competencies"]}))
weighed = S.parse_sheet(SHEET[:7] + [["2", S.CONTINUATION, "Savings of 5%", "Monthly", 10]] + SHEET[8:])
if [row["weight"] for row in weighed["kpis"]][:2] != [25.0, 10.0] \
        or [row["weight"] for row in S.spread_weights(weighed["kpis"])][:2] != [25.0, 10.0]:
    fail.append("a KPI the sheet weighs of its own keeps its weight, and its perspective is not shared out: %s"
                % [row["weight"] for row in weighed["kpis"]])
if S.parse_sheet([]) != {"role": None, "department": None, "grade": None, "review_period": None,
                         "form_reference": None, "revision": None,
                         "perspectives": [], "kpis": [], "competencies": []}:
    fail.append("an empty sheet reads as nothing")
if S.normalise_perspective(S.CONTINUATION) is not None or S.normalise_perspective("") is not None:
    fail.append("the continuation arrow is not a perspective")
footer = "Luuka Plastics Limited  |  PMS BSC Appraisal Form FY 2026  |  Procurement  |  PROC/002  |  CONFIDENTIAL  |  Rev 01"
if S.footer_parts(footer) != ("PROC/002", "Rev 01") or S.footer_parts("Nothing here") != (None, None):
    fail.append("the footer names the form and its revision: %s" % (S.footer_parts(footer),))
if S.parse_sheet(SHEET + [[footer]])["form_reference"] != "PROC/002":
    fail.append("the sheet's form reference is read from its footer")

# ── the template laid out as the form is, every KPI weighed ───────────
typed = [{"perspective": "Financial", "kpi": "Savings", "weight": 15},
         {"perspective": "Customer / Stakeholder", "kpi": "Lead times", "weight": 15},
         {"perspective": "Financial", "kpi": "Variance", "weight": 10},
         {"perspective": "Internal Business Processes", "kpi": "Quotations", "weight": 20},
         {"perspective": "Internal Business Processes", "kpi": "Approvals", "weight": 10},
         {"perspective": "Learning & Growth", "kpi": "CIPS", "weight": 10}]
arranged, perspectives = S.arrange_kpis(typed)
if [row["kpi"] for row in arranged] != ["Savings", "Variance", "Lead times", "Quotations", "Approvals", "CIPS"]:
    fail.append("a perspective's KPIs sit together, in the order the perspectives first appear: %s"
                % [row["kpi"] for row in arranged])
if [row["weight"] for row in arranged] != [15, 10, 15, 20, 10, 10]:
    fail.append("every KPI keeps its own weight: %s" % [row["weight"] for row in arranged])
if perspectives != [{"perspective": "Financial", "weight": 25.0},
                    {"perspective": "Customer / Stakeholder", "weight": 15.0},
                    {"perspective": "Internal Business Processes", "weight": 30.0},
                    {"perspective": "Learning & Growth", "weight": 10.0}]:
    fail.append("a perspective weighs what its KPIs weigh: %s" % perspectives)
if S.arrange_kpis(arranged)[0] != arranged:
    fail.append("laying out a laid-out template changes nothing")
if S.arrange_kpis([]) != ([], []):
    fail.append("no KPIs, no perspectives")
given = [{"perspective": "A", "weight": 25}, {"perspective": "A"}, {"perspective": "A"},
         {"perspective": "B", "weight": 0.29}, {"perspective": "B", "weight": None},
         {"perspective": "C", "weight": 7}, {"perspective": "D", "weight": 4}, {"perspective": "D", "weight": 6}]
if [row.get("weight") for row in S.spread_weights(given)] != [8.33, 8.33, 8.34, 0.14, 0.15, 7, 4, 6]:
    fail.append("a perspective's one weight is shared evenly to the hundredth, the last taking what rounding "
                "leaves; one KPI, or KPIs weighed one by one, are left: %s"
                % [row.get("weight") for row in S.spread_weights(given)])
if given[1].get("weight") is not None:
    fail.append("sharing out the weights leaves the rows it was given as they were")
expect("a supervisory template with its factors", S.supervisory_template_errors(
    {"factors": [{"factor": "Attendance"}], "objectives": [{"objective": "Output"}]}))
expect("one with none", S.supervisory_template_errors({"factors": []}), "ratable factors")
expect("one with a factor twice", S.supervisory_template_errors(
    {"factors": [{"factor": "Attendance"}, {"factor": "Attendance"}]}), "listed twice")
expect("one with nine objectives", S.supervisory_template_errors(
    {"factors": [{"factor": "A"}], "objectives": [{"objective": "O%d" % n} for n in range(9)]}), "at most 8")

# ── the employee's own scorecard ──────────────────────────────────────
OWN = {"step": "self", "quarter": "Q1",
       "kpis": [{"perspective": "Financial", "kpi": "Savings", "weight": 50, "self_percent": 80},
                {"perspective": "Customer / Stakeholder", "kpi": "Lead times", "weight": 30, "self_percent": 50}],
       "competencies": [{"competency": "One", "weight": 20, "self_score": 7}]}
expect("a self-appraisal scored throughout", S.appraisal_errors(OWN))
expect("a self-appraisal with a KPI left", S.appraisal_errors(dict(OWN, kpis=[
    {"perspective": "Financial", "kpi": "Savings", "weight": 50}])), "your own Q1 percentage achieved")
expect("a self-appraisal with a competency left", S.appraisal_errors(dict(OWN, competencies=[
    {"competency": "One", "weight": 20}])), "Score yourself on every competency")
expect("a self-appraisal scoring a competency eleven", S.appraisal_errors(dict(OWN, competencies=[
    {"competency": "One", "weight": 20, "self_score": 11}])), "out of ten")
expect("the appraiser's own step does not read the employee's figures",
       S.appraisal_errors(dict(OWN, step="appraiser")), "Record Q1 percentage achieved", "Score every competency")
own = S.self_scores(OWN["kpis"], OWN["competencies"], "Q1")
if own != {"section_a": 55.0, "section_b": 14.0, "overall": 69.0}:
    fail.append("the employee's own scores are worked out as the appraiser's are: %s" % own)
if S.self_scores([], [], "Q4") != {"section_a": None, "section_b": None, "overall": None}:
    fail.append("nothing rated, no score of their own")
print("the scorecard: 80 and 20, every KPI weighed and scored, the perspectives summing them up, four quarters and "
      "their average, its own bands, Luuka's workbook read and shared out, the self-appraisal")

# ── 8. Two forms, two chains, one workflow ────────────────────────────
if set(A.ROUTES) != set(A.FORM_TYPES):
    fail.append("each form must declare the states it passes through")
for form, wanted in ((A.FORM_SUPERVISORY, (A.DRAFT, A.PENDING_SELF, A.PENDING_SUPERVISOR, A.PENDING_HRM,
                                           A.PENDING_PRODUCTION, A.PENDING_GM, A.COMPLETED)),
                     (A.FORM_BSC, (A.DRAFT, A.PENDING_SELF, A.PENDING_SUPERVISOR, A.PENDING_EMPLOYEE, A.PENDING_HOD,
                                   A.PENDING_HRM, A.PENDING_ED, A.COMPLETED))):
    for own in (True, False):
        expected = wanted if own else tuple(state for state in wanted if state != A.PENDING_SELF)
        if A.route(form, own) != expected:
            fail.append("%s is signed %s, not %s" % (form, expected, A.route(form, own)))
        walked, state = [A.DRAFT], A.DRAFT
        guard = 0
        while guard < 12:
            guard += 1
            forward = [step for action, step in A.next_states(state, {"HR User", "Employee", "Supervisor",
                                                                      "Head of Department", "HR Manager",
                                                                      "Production Manager", "General Manager",
                                                                      "Executive Director"}, form, own)
                       if step != A.DRAFT and step not in walked]
            if not forward:
                break
            state = forward[0]
            walked.append(state)
            if state == A.COMPLETED:
                break
        if tuple(walked) != expected:
            fail.append("walking %s (self-appraisal %s) by its own transitions gives %s, not %s"
                        % (form, "on" if own else "off", walked, expected))
# every state either form can reach has a way on and a way back, whoever holds it
for form in A.FORM_TYPES:
    for own in (True, False):
        for state in A.route(form, own):
            if state in (A.DRAFT, A.COMPLETED):
                continue
            holders = {row["allow_edit"] for row in A.STATES if row["state"] == state}
            moves = [pair for role in holders for pair in A.next_states(state, {role}, form, own)]
            if not [pair for pair in moves if pair[1] != A.DRAFT]:
                fail.append("%s (%s): whoever edits it cannot pass it on" % (form, state))
            if not [pair for pair in moves if pair == (A.RETURN, A.DRAFT)]:
                fail.append("%s (%s): whoever holds it cannot return it" % (form, state))
if A.next_states(A.PENDING_HRM, {"HR Manager"}, A.FORM_BSC) == A.next_states(A.PENDING_HRM, {"HR Manager"},
                                                                            A.FORM_SUPERVISORY):
    fail.append("after the HR Manager the two forms part: one to Production, one to the Executive Director")
for state in A.PENDING_STATES:
    if not [t for t in A.TRANSITIONS if t["state"] == state and t["action"] == A.RETURN and t["next_state"] == A.DRAFT]:
        fail.append("%s must be able to return the appraisal to Draft" % state)
for form in A.FORM_TYPES:
    stamps, remarks = A.stamps_for(form), A.remarks_for(form)
    # every pending state signs the form, except the scorecard's
    # self-appraisal: its employee block is signed over the appraiser's scores
    signing = [state for state in A.route(form)
               if state in A.PENDING_STATES and not (state == A.PENDING_SELF and form == A.FORM_BSC)]
    if set(stamps) != set(signing):
        fail.append("%s: every state it passes through signs the form (%s vs %s)"
                    % (form, sorted(stamps), sorted(signing)))
    if set(remarks) != set(stamps):
        fail.append("%s: everyone who signs also comments" % form)
if A.stamps_for(A.FORM_BSC).get(A.PENDING_EMPLOYEE) != ("custom_employee_signed_by", "custom_employee_signed_on"):
    fail.append("on the scorecard the employee signs after the appraiser has scored, not in Draft")
if A.stamps_for(A.FORM_SUPERVISORY).get(A.PENDING_SELF) != ("custom_employee_signed_by", "custom_employee_signed_on"):
    fail.append("on the supervisory form the employee signs their own self-appraisal")
if A.DRAFT in A.stamps_for(A.FORM_SUPERVISORY) or A.DRAFT in A.stamps_for(A.FORM_BSC):
    fail.append("HR sending the appraisal on from Draft signs nobody's block")
if A.compute_stamps(A.PENDING_SELF, A.PENDING_SUPERVISOR, "emp@luuka", "2026-09-24", {},
                    A.FORM_BSC).get("custom_employee_signed_by"):
    fail.append("the scorecard's self-appraisal is not the employee's signature over the appraiser's scores")
if A.compute_stamps(A.PENDING_SELF, A.PENDING_SUPERVISOR, "emp@luuka", "2026-09-24", {},
                    A.FORM_SUPERVISORY).get("custom_employee_signed_by") != "emp@luuka":
    fail.append("the supervisory form's self-appraisal is signed by whoever submitted it")
if A.compute_stamps(A.DRAFT, A.PENDING_SUPERVISOR, "hro@luuka", "2026-09-24", {},
                    A.FORM_SUPERVISORY).get("custom_employee_signed_by"):
    fail.append("with no self-appraisal, HR sending it to the supervisor does not sign as the employee")
if set(A.ROLE_WAITING) != set(A.PENDING_STATES):
    fail.append("every pending state must know whose desk it is on")
if A.ROLE_WAITING[A.PENDING_EMPLOYEE] != A.APPRAISEE or A.ROLE_WAITING[A.PENDING_ED] != A.ED:
    fail.append("the scorecard waits on the employee, then the Head of Department, then the Executive Director")
expect("the appraiser passing the scorecard on with no comments",
       A.step_errors(A.PENDING_SUPERVISOR, A.PENDING_EMPLOYEE, {"form_type": A.FORM_BSC}), "Appraiser's general comments")
expect("the appraiser having commented",
       A.step_errors(A.PENDING_SUPERVISOR, A.PENDING_EMPLOYEE,
                     {"form_type": A.FORM_BSC, "custom_supervisor_remarks": "Strong year."}))
for role, ptypes in (("Executive Director", ("read", "write")), ("Head of Department", ("read", "write"))):
    granted = (A.PERMISSIONS.get("Appraisal") or {}).get(role) or ()
    if [ptype for ptype in ptypes if ptype not in granted]:
        fail.append("%s signs the scorecard and needs %s on the Appraisal" % (role, ptypes))
print("two forms: each route walked, each signed by its own people, the junctions conditional")

# ── 9. The scorecard's DocTypes and glue ──────────────────────────────
# Frappe HR ships an Appraisal Template. Luuka's scorecard is carried on
# it, the way both appraisal forms are carried on their Appraisal, so
# there must be no second template DocType standing beside it.
if doctype("BSC Appraisal Template"):
    fail.append("the scorecard belongs on Frappe HR's Appraisal Template, not a DocType beside it")
if not upstream_doctype("Appraisal Template"):
    fail.append("Frappe HR's Appraisal Template is not where it was: the scorecard is built on it")
theirs = custom_fields("Appraisal Template")
for fieldname in ("custom_designation", "custom_review_year", "custom_department", "custom_grade",
                  "custom_review_period", "custom_company", "custom_is_active", "custom_perspectives",
                  "custom_kpis", "custom_competencies", "custom_objectives_weight",
                  "custom_competencies_weight", "custom_source_file", "custom_source_sheet",
                  "custom_import_remarks"):
    if fieldname not in theirs:
        fail.append("Appraisal Template has no %s, which the scorecard asks for" % fieldname)
for fieldname, options in (("custom_perspectives", "BSC Template Perspective"),
                           ("custom_kpis", "BSC Template KPI"),
                           ("custom_competencies", "BSC Template Competency")):
    if (theirs.get(fieldname) or {}).get("options") != options:
        fail.append("Appraisal Template.%s must be a table of %s" % (fieldname, options))
for fieldname in ("custom_objectives_weight", "custom_competencies_weight", "custom_source_file",
                  "custom_source_sheet", "custom_import_remarks"):
    if not (theirs.get(fieldname) or {}).get("read_only"):
        fail.append("Appraisal Template.%s is worked out, not typed" % fieldname)
if (theirs.get("custom_designation") or {}).get("options") != "Designation":
    fail.append("a scorecard is one role's: Appraisal Template.custom_designation links a Designation")
# their own KRA table and rating criteria are neither of Luuka's forms
setter_names = {row["name"] for row in SETTERS}
for fieldname in ("goals", "rating_criteria"):
    if "Appraisal Template-%s-hidden" % fieldname not in setter_names:
        fail.append("Appraisal Template.%s is not part of either Luuka form and must be put away" % fieldname)
# Hiding a table does not stop it being mandatory, and it does not stop a
# blank row in it refusing the save. Frappe HR ships goals as reqd, so the
# flag has to be cleared as well, or a scorecard template cannot be saved
# at all.
if (upstream_fields("Appraisal Template").get("goals") or {}).get("reqd"):
    if "Appraisal Template-goals-reqd" not in setter_names:
        fail.append("Frappe HR's goals table is mandatory: a scorecard template will not save "
                    "until a Property Setter clears reqd on it")
    else:
        cleared = [row for row in SETTERS if row["name"] == "Appraisal Template-goals-reqd"]
        if str(cleared[0].get("value")) != "0":
            fail.append("Appraisal Template-goals-reqd must clear the flag, not set it")
# and a blank row in either hidden table is dropped, on the form and on the
# server, so neither can refuse a save over a table nobody can see
form = read("hrms_addon", "public", "js", "appraisal_template.js")
if "before_save(frm)" not in form or '"goals"' not in form:
    fail.append("the template form must drop a blank row from the hidden KRA table before "
                "Frappe's mandatory check runs")
glue = read("hrms_addon", "hrms_addon", "bsc.py")
if "_drop_blank_upstream_rows" not in glue:
    fail.append("and the server must do it too, for an import or an API call")
# the appraisal names its template through their own link, not one of ours
if "Appraisal-appraisal_template-hidden" in setter_names:
    fail.append("Appraisal.appraisal_template IS Luuka's template now: it must not be hidden")
if "Appraisal-appraisal_template-description" not in setter_names:
    fail.append("Appraisal.appraisal_template should say which template it means")
if "custom_bsc_template" in custom_fields("Appraisal"):
    fail.append("the Appraisal must name its template once, through Frappe HR's own appraisal_template")
if (hooks.get("doc_events", {}).get("Appraisal Template", {}).get("validate")
        != "hrms_addon.hrms_addon.bsc.template_validate"):
    fail.append("the scorecard's weights are checked when their template is saved")
if "hrms_addon.patches.v1_0.scorecard_onto_appraisal_template" not in read("hrms_addon", "patches.txt"):
    fail.append("a site that already imported scorecards must have them carried across")
QUARTER_COLUMNS = tuple("%s_%s" % (quarter.lower(), kind) for quarter in S.QUARTERS
                        for kind in ("percent", "score", "comments"))
for name, wanted in (
    ("BSC Template Perspective", ("perspective", "weight")),
    ("BSC Template KPI", ("perspective", "kpi", "timing", "weight")),
    ("BSC Template Competency", ("competency", "indicators", "weight")),
    ("BSC Appraisal KPI", ("perspective", "kpi", "timing", "weight", "self_percent", "score") + QUARTER_COLUMNS),
    ("BSC Appraisal Perspective", ("perspective", "weight", "q1_score", "q2_score", "q3_score", "q4_score",
                                   "year_to_date")),
    ("Appraisal Quarter Result", ("quarter", "appraisal", "section_a", "section_b", "total", "band", "status")),
    ("BSC Appraisal Competency", ("competency", "indicators", "weight", "score", "weighted_score")),
    ("BSC Competency", ("competency_name", "indicators", "default_weight")),
    ("Appraisal Assignment", ("task", "assignment_given", "expected_outcome", "employee_comments",
                              "supervisor_comments")),
    ("Development Action", ("action", "duration", "by_when", "by_whom", "estimated_cost")),
):
    fields = fields_of(doctype(name))
    if not fields:
        fail.append("%s is not there" % name)
        continue
    for fieldname in wanted:
        if fieldname not in fields:
            fail.append("%s has no %s, which the scorecard asks for" % (name, fieldname))
worked_out = [("BSC Appraisal Competency", "weighted_score"), ("BSC Appraisal Competency", "weight"),
              ("BSC Template Perspective", "weight")]
worked_out += [("BSC Appraisal Perspective", fieldname) for fieldname in fields_of(doctype("BSC Appraisal Perspective"))]
worked_out += [("Appraisal Quarter Result", fieldname) for fieldname in fields_of(doctype("Appraisal Quarter Result"))]
worked_out += [("BSC Appraisal KPI", fieldname) for fieldname in ("perspective", "kpi", "timing", "weight", "score")
               + tuple(S.score_field(quarter) for quarter in S.QUARTERS)]
for name, fieldname in worked_out:
    if not (fields_of(doctype(name)).get(fieldname) or {}).get("read_only"):
        fail.append("%s.%s is worked out or carried from the template, not typed" % (name, fieldname))
# each quarter's percentage and comments, and the employee's own figure, are
# typed: the form opens only the quarter appraised (appraisal.js), the server
# carries the earlier ones in (appraisals._carry_earlier_quarters)
for fieldname in ("self_percent",) + tuple(field for quarter in S.QUARTERS
                                           for field in (S.percent_field(quarter), S.comments_field(quarter))):
    if (fields_of(doctype("BSC Appraisal KPI")).get(fieldname) or {}).get("read_only"):
        fail.append("BSC Appraisal KPI.%s is typed for the quarter appraised: it cannot be read-only" % fieldname)
if (fields_of(doctype("BSC Appraisal KPI")).get("timing") or {}).get("options", "").split("\n")[0] != "":
    fail.append("a KPI's timing starts blank: Frappe pre-fills a Select's first option")
ours = custom_fields("Appraisal")
for fieldname in ("custom_form_type", "custom_quarter", "custom_bsc_perspectives",
                  "custom_bsc_kpis", "custom_bsc_competencies", "custom_assignments",
                  "custom_bsc_section_a_score", "custom_bsc_section_b_score", "custom_bsc_overall",
                  "custom_bsc_band", "custom_bsc_band_meaning", "custom_hod_by", "custom_hod_on",
                  "custom_hod_remarks", "custom_ed_by", "custom_ed_on", "custom_ed_remarks",
                  "custom_continue", "custom_stop", "custom_start", "custom_development_actions",
                  "custom_quarter_results", "custom_annual_score", "custom_year_band", "custom_on_pip",
                  "custom_improvement_plan"):
    if fieldname not in ours:
        fail.append("the Appraisal has no %s, which the scorecard asks for" % fieldname)
if (ours.get("custom_form_type") or {}).get("options", "").split("\n") != list(S.FORM_TYPES):
    fail.append("Appraisal.custom_form_type must offer exactly the two forms")
if "custom_period" in ours:
    fail.append("the scorecard's own period gave way to the appraisal's quarter: custom_period must go")
quarter = ours.get("custom_quarter") or {}
if quarter.get("options", "").split("\n") != [""] + list(S.QUARTERS) or quarter.get("read_only") \
        or "custom_plan" not in (quarter.get("read_only_depends_on") or "") \
        or "Pending Self-Appraisal" not in (quarter.get("read_only_depends_on") or ""):
    fail.append("the quarter is the plan's, or HR's on an appraisal made by hand until it is rated: %s" % quarter)
if (ours.get("custom_quarter_results") or {}).get("options") != "Appraisal Quarter Result":
    fail.append("Appraisal.custom_quarter_results is a table of Appraisal Quarter Result")
for fieldname in ("custom_bsc_section_a_score", "custom_bsc_section_b_score", "custom_bsc_overall",
                  "custom_bsc_band", "custom_hod_by", "custom_ed_by", "custom_bsc_perspectives",
                  "custom_quarter_results", "custom_annual_score", "custom_year_band", "custom_on_pip",
                  "custom_improvement_plan"):
    if not (ours.get(fieldname) or {}).get("read_only"):
        fail.append("Appraisal.%s is worked out, not typed" % fieldname)
for fieldname in ("custom_annual_score", "custom_quarter_results", "custom_results_section"):
    if (ours.get(fieldname) or {}).get("depends_on"):
        fail.append("Appraisal.%s shows the year so far on both forms" % fieldname)
if not (ours.get("custom_on_pip") or {}).get("in_list_view") \
        or not (ours.get("custom_on_pip") or {}).get("in_standard_filter"):
    fail.append("the Appraisal list shows and filters who is on an improvement plan")
# each form's own sections are shown only for that form
for fieldname in ("custom_factors", "custom_objectives", "custom_total_score", "custom_production_remarks"):
    if "Balanced Scorecard" not in (ours.get(fieldname) or {}).get("depends_on", ""):
        fail.append("Appraisal.%s belongs to the supervisory form and must be hidden on a scorecard" % fieldname)
for fieldname in ("custom_bsc_perspectives", "custom_bsc_competencies", "custom_bsc_overall"):
    if "Balanced Scorecard" not in (ours.get(fieldname) or {}).get("depends_on", ""):
        fail.append("Appraisal.%s belongs to the scorecard and must be hidden on the supervisory form" % fieldname)

glue_bsc = read("hrms_addon", "hrms_addon", "bsc.py")
# it writes the appraisal and the template, both of them Frappe HR's own
both = dict(all_fields("Appraisal"), **all_fields("Appraisal Template"))
for fieldname in sorted(set(re.findall(r'doc\.get\("(custom_\w+)"\)', glue_bsc))
                        | set(re.findall(r"doc\.(custom_\w+)\b", glue_bsc))
                        | set(re.findall(r'card\.(custom_\w+)\b', glue_bsc))):
    if fieldname not in both:
        fail.append("bsc.py reads or writes %s, which is on neither the Appraisal nor its Template" % fieldname)
if 'TEMPLATE = "Appraisal Template"' not in glue_bsc:
    fail.append("bsc.py must name Frappe HR's own Appraisal Template as the scorecard's home")
for needle, why in (
    ("rules.template_errors(", "a scorecard is judged by the rules"),
    ("rules.quarter_score(", "each KPI's quarter is scored by the rules"),
    ("rules.perspective_summary(", "and the perspectives sum them up"),
    ("rules.spread_weights(", "a workbook's perspective weight is shared out between its KPIs"),
    ("rules.section_a(", "Section A comes from the rules"),
    ("rules.section_b(", "and Section B"),
    ("rules.band(", "and the band"),
    ("rules.parse_sheet(", "Luuka's workbook is read by the rules"),
    ("load_workbook(", "the importer opens the workbook"),
    ('frappe.has_permission(TEMPLATE, "create")', "the importer makes documents: it checks first"),
    ("doc.custom_is_active = 1 if (activate and not problems) else 0",
     "a sheet whose weights do not add up is imported but left inactive"),
    ("doc.template_title = _title(designation, year, sheet)",
     "their template is named after the role, because that is its autoname"),
):
    if needle not in glue_bsc:
        fail.append("bsc.py: %s (%r not found)" % (why, needle))
if "return appraisals.apply_template(appraisal, template)" not in glue_bsc:
    fail.append("bsc.get_scorecard, the button's older name, takes the template the one way apply_template does")
glue_appraisals = read("hrms_addon", "hrms_addon", "appraisals.py")
if 'doc.check_permission("write")' not in body_of(glue_appraisals, "apply_template"):
    fail.append("apply_template writes onto an appraisal: it must check the caller may write it")
for name in ("get_scorecard", "import_workbook"):
    if not re.search(r'@frappe\.whitelist\(methods=\["POST"\]\)\ndef %s\(' % name, glue_bsc):
        fail.append("bsc.%s changes something: a whitelisted POST method" % name)
for needle, why in (
    ("bsc.template_for(", "an active scorecard made for the role is the last place the template is looked for"),
    ('frappe.db.get_value("Designation", designation, "appraisal_template")',
     "the Job Title's own template comes before any other"),
    ('"Appraisee", {"parent": cycle, "parenttype": "Appraisal Cycle"',
     "and the one the cycle names for the employee before that"),
    ("doc.appraisal_template = template_for(",
     "the employee's template attaches itself, on Frappe HR's own link"),
    ("doc.custom_form_type = form_of(template)", "the template says which form the employee is on"),
    ("bsc.fill(", "a scorecard appraisal is filled from the template"),
    ("bsc.score(", "and scored by the scorecard's own rules"),
    ("bsc_rules.appraisal_errors(", "and judged by them"),
    ("_carry_scores(", "the scorecard's overall becomes the appraisal's score, so one report reads both forms"),
    ("form_type)", "the signatures follow the form's own chain"),
):
    if needle not in glue_appraisals:
        fail.append("appraisals.py: %s (%r not found)" % (why, needle))
if "hrms_addon.hrms_addon.pick_lists.seed_bsc_masters" not in (hooks.get("after_install") or []):
    fail.append("a fresh install seeds the scorecard's competencies")
if "seed_bsc_masters()" not in read("hrms_addon", "patches", "v1_0", "seed_performance.py"):
    fail.append("and so does the patch, on a site that has the app already")
if "Appraisal Template" not in carded:
    fail.append("Appraisal Template is on no workspace card")
if "Appraisal Template" in sidebarred:
    fail.append("Frappe HR already lists Appraisal Template under Setup: leave their entry where it is")
for name in ("BSC Competency",):
    if name not in carded:
        fail.append("%s is on no workspace card" % name)
    if name not in sidebarred:
        fail.append("%s is in no sidebar" % name)
print("the scorecard's forms, its glue, the importer guarded, the seed and the way in")

# ── 10. The template each appraisal is filled from, and the self-appraisal ──
status = custom_fields("Appraisal").get("custom_appraisal_status") or {}
missing = sorted({row["status"] for row in A.STATES} - set((status.get("options") or "").split("\n")))
if missing:
    fail.append("the workflow writes %s into the Appraisal Status, which does not offer them: the save would be "
                "refused" % missing)
if set(R.STATUSES) != set((status.get("options") or "").split("\n")):
    fail.append("appraisal_rules.STATUSES and the Appraisal Status options are one list")
ours = custom_fields("Appraisal")
name_field = ours.get("custom_supervisor_name") or {}
if name_field.get("fetch_from") != "custom_supervisor.employee_name" or not name_field.get("read_only") \
        or name_field.get("insert_after") != "custom_supervisor":
    fail.append("the supervisor's name shows beside the supervisor, fetched and read-only: %s" % name_field)
if "custom_supervisor_name = frappe.db.get_value(\"Employee\", doc.custom_supervisor, \"employee_name\")" \
        not in glue_appraisals:
    fail.append("the supervisor's name is filled on the server too, wherever the appraisal is made")
if "if (boss !== (frm.doc.custom_supervisor || \"\")) frm.set_value(\"custom_supervisor\", boss);" \
        not in read("hrms_addon", "public", "js", "appraisal.js"):
    fail.append("picking the employee fills in their supervisor, even over one already there")
for fieldname in ("custom_self_appraisal", "custom_bsc_self_score"):
    if not (ours.get(fieldname) or {}).get("read_only"):
        fail.append("Appraisal.%s is set by the system: read-only" % fieldname)
def read_upstream(*parts):
    return open(os.path.join(APPS_ROOT, *parts), encoding="utf-8").read()


def appraisals_self_ratings():
    """appraisals.SELF_RATINGS, read without importing Frappe."""
    for node in ast.parse(glue_appraisals).body:
        if isinstance(node, ast.Assign) and getattr(node.targets[0], "id", None) == "SELF_RATINGS":
            return ast.literal_eval(node.value)
    return ()


validate_body = body_of(glue_appraisals, "appraisal_validate")
if "if doc.is_new() or (doc.get(approval.STATE_FIELD) or approval.DRAFT) == approval.DRAFT:\n" \
        "        # one HR has not sent on yet follows Appraisal Settings as they are now\n" \
        "        doc.custom_self_appraisal = settings().self_appraisal" not in validate_body:
    fail.append("an appraisal made by hand, and any still in Draft, follows the setting as it is now")
if ((hooks.get("doc_events") or {}).get("Appraisal Settings") or {}).get("on_update") \
        != "hrms_addon.hrms_addon.appraisals.settings_on_update":
    fail.append("saving Appraisal Settings reaches the appraisals already raised")
if (hooks.get("override_whitelisted_methods") or {}).get(
        "hrms.hr.doctype.appraisal_cycle.appraisal_cycle.get_appraisal_cycle_summary") \
        != "hrms_addon.hrms_addon.appraisals.get_appraisal_cycle_summary":
    fail.append("the Appraisal Cycle's Self Appraisal Pending counts the appraisals waiting on the employee")
tab_setter = [row for row in json.load(open(os.path.join(PACKAGE, "fixtures", "property_setter.json"),
                                            encoding="utf-8"))
              if row["name"] == "Appraisal-self_appraisal_tab-depends_on"]
if not tab_setter or (tab_setter[0]["doc_type"], tab_setter[0]["field_name"], tab_setter[0]["property"],
                      tab_setter[0]["value"]) != ("Appraisal", "self_appraisal_tab", "depends_on",
                                                  "eval:doc.custom_self_appraisal"):
    fail.append("Frappe HR's Self Appraisal tab shows only on an appraisal with a self-appraisal")
if (upstream_fields("Appraisal").get("self_appraisal_tab") or {}).get("fieldtype") != "Tab Break":
    fail.append("Frappe HR's Appraisal has no Self Appraisal tab to hide")
if '"Appraisal-self_appraisal_tab-depends_on",' not in read("hrms_addon", "hooks.py"):
    fail.append("hooks.py's fixture filter names the Self Appraisal tab's setter")
table_of = {row.get("options"): row.get("fieldname") for row in ours.values() if row.get("fieldtype") == "Table"}
for child, fieldname in appraisals_self_ratings():
    if child not in table_of or fieldname not in fields_of(doctype(child)):
        fail.append("the employee's own ratings are %s.%s, a table on the Appraisal" % (child, fieldname))
if {table_of.get(child) for child, _field in appraisals_self_ratings()} != \
        {"custom_factors", "custom_objectives", "custom_bsc_kpis", "custom_bsc_competencies"}:
    fail.append("every table the employee rates in is read for their self-appraisal")
people_glue = read("hrms_addon", "hrms_addon", "people.py")
if "_remove(doctype, name, user, ignore_permissions=True)" not in body_of(people_glue, "withdraw"):
    fail.append("a task withdrawn is cancelled through Frappe's own assignment, so the document's list follows")
for path, needle, why in (
    (("frappe", "frappe", "desk", "form", "assign_to.py"), "def _remove(doctype, name, assign_to, ignore_permissions=False):",
     "Frappe withdraws an assignment through _remove"),
    (("frappe", "frappe", "public", "js", "frappe", "form", "layout.js"), "const fields = this.fields_list.concat(this.tabs);",
     "a tab's depends_on is followed, so the Self Appraisal tab can hide"),
    (("hrms", "hrms", "hr", "doctype", "appraisal_cycle", "appraisal_cycle.py"),
     "def get_appraisal_cycle_summary(cycle_name: str) -> dict:", "Frappe HR's cycle summary, which ours wraps"),
    (("hrms", "hrms", "hr", "doctype", "appraisal_cycle", "appraisal_cycle.py"),
     'summary["self_appraisal_pending"] = frappe.db.count(', "and the count ours puts right"),
    (("frappe", "frappe", "model", "workflow.py"), "doc = frappe.get_doc(frappe.parse_json(doc))\n\tdoc.load_from_db()",
     "a workflow action reads the appraisal as stored, so a form opened before the setting changed cannot send "
     "it on the old way"),
):
    if needle not in read_upstream(*path):
        fail.append("%s: %s (%r not found)" % ("/".join(path), why, needle))
follow_patch = read("hrms_addon", "patches", "v1_0", "self_appraisal_follows_settings.py")
for needle, why in (
    ("appraisals.follow_settings()", "the appraisals already raised follow the setting as it is"),
    ("gave = appraisals.gave_self_appraisal(further)", "one further on keeps its self-appraisal where it was given"),
    ('frappe.db.set_value("Appraisal", name, approval.SELF_FIELD, 0, update_modified=False)',
     "and loses it where it was not"),
):
    if needle not in follow_patch:
        fail.append("the self-appraisal patch: %s (%r not found)" % (why, needle))
listed_patches = read("hrms_addon", "patches.txt").split()
if "hrms_addon.patches.v1_0.self_appraisal_follows_settings" not in listed_patches or \
        listed_patches.index("hrms_addon.patches.v1_0.self_appraisal_follows_settings") \
        < listed_patches.index("hrms_addon.patches.v1_0.appraisal_templates_and_self_appraisal"):
    fail.append("the self-appraisal patch runs after the one that turned it on for everyone")
for child, fieldname in (("BSC Appraisal KPI", "self_percent"), ("BSC Appraisal Competency", "self_score"),
                         ("BSC Appraisal KPI", "q1_comments"), ("BSC Template KPI", "weight"),
                         ("Appraisal Template Factor", "factor"), ("Appraisal Template Objective", "objective")):
    if fieldname not in fields_of(doctype(child)):
        fail.append("%s has no %s" % (child, fieldname))
for child in ("BSC Appraisal Perspective", "BSC Appraisal Competency", "BSC Appraisal KPI", "BSC Template KPI",
              "Employee Tool", "Onboarding Tool"):
    listed = sum(f.get("columns") or 0 for f in doctype(child)["fields"] if f.get("in_list_view"))
    if listed > 10:
        fail.append("%s's grid asks for %d columns: Frappe shows ten, and drops the rest" % (child, listed))
settings_spec = doctype("Appraisal Settings")
settings_fields = fields_of(settings_spec)
if not settings_spec.get("issingle") or (settings_fields.get("self_appraisal") or {}).get("default") != "1" \
        or (settings_fields.get("kra_evaluation_method") or {}).get("default") != R.KRA_AUTOMATED \
        or tuple((settings_fields.get("kra_evaluation_method") or {}).get("options", "").split("\n")) != R.KRA_METHODS:
    fail.append("Appraisal Settings: a single, self-appraisal on and KRAs automated unless changed")
if set(settings_fields) - {"self_section", "cycle_section"} != set(R.SETTINGS_DEFAULTS):
    fail.append("Appraisal Settings holds what settings_values reads, and nothing else")
if "frappe.db.get_singles_dict(SETTINGS)" not in body_of(glue_appraisals, "settings"):
    fail.append("the settings are read as stored: a Check never saved would read 0, turning the self-appraisal off")
if "Appraisal Settings" not in carded or "Appraisal Settings" not in sidebarred:
    fail.append("Appraisal Settings has a way in, beside the templates")

template = custom_fields("Appraisal Template")
form = template.get("custom_form_type") or {}
if tuple((form.get("options") or "").split("\n")) != (S.FORM_BSC, S.FORM_SUPERVISORY) or form.get("default") != S.FORM_BSC:
    fail.append("a template says which form it carries, the scorecard unless it says otherwise: %s" % form)
perspectives = template.get("custom_perspectives") or {}
if perspectives.get("hidden") or not perspectives.get("read_only") or perspectives.get("insert_after") != "custom_kpis":
    fail.append("the perspectives follow from the KPIs: shown below them, never typed into: %s" % perspectives)
if (template.get("custom_kpis") or {}).get("insert_after") != "custom_section_a":
    fail.append("Section A opens with the KPIs, the perspectives below them")
if (template.get("custom_kpis") or {}).get("label") != "Objectives & KPIs":
    fail.append("Section A is one table, as the workbook has it")
for fieldname, options in (("custom_factors", "Appraisal Template Factor"),
                           ("custom_objectives", "Appraisal Template Objective")):
    if (template.get(fieldname) or {}).get("options") != options:
        fail.append("a supervisory template carries its %s" % fieldname)
for fieldname in ("custom_factors_section", "custom_objectives_section"):
    if S.FORM_SUPERVISORY not in ((template.get(fieldname) or {}).get("depends_on") or ""):
        fail.append("%s shows only on a supervisory template" % fieldname)
for fieldname in ("custom_section_a", "custom_section_b"):
    if S.FORM_SUPERVISORY not in ((template.get(fieldname) or {}).get("depends_on") or ""):
        fail.append("%s shows only on a scorecard" % fieldname)
chain, seen = [], set()
by_after = {}
for fieldname, row in template.items():
    by_after.setdefault(row.get("insert_after"), []).append(fieldname)
twice = {after: names for after, names in by_after.items() if len(names) > 1}
if twice:
    fail.append("two template fields follow the same field, so their order is left to chance: %s" % twice)
SETTERS = json.load(open(os.path.join(PACKAGE, "fixtures", "property_setter.json"), encoding="utf-8"))
if not [row for row in SETTERS if row["doc_type"] == "Designation" and row.get("field_name") == "appraisal_template"
        and row["property"] == "reqd" and str(row["value"]) == "1"]:
    fail.append("every Job Title must name its Appraisal Template: Designation.appraisal_template is mandatory")
if "frm.set_query(\"appraisal_template\", () => ({ filters: { custom_is_active: 1 } }));" \
        not in read("hrms_addon", "public", "js", "designation.js"):
    fail.append("a Job Title picks from the templates made ready to be used")
if "designation.flags.ignore_mandatory = True" not in body_of(glue_bsc, "_designation"):
    fail.append("a Job Title the workbook makes is saved before its template exists, so its template is named after")
if "_link_designation(designation, doc.name, year)" not in glue_bsc:
    fail.append("the workbook names each Job Title's template once the template is saved")
if "hrms_addon.hrms_addon.bsc.seed_supervisory_template" not in (hooks.get("after_install") or []):
    fail.append("a fresh install makes the supervisory form's template")
if "doc.custom_form_type == rules.FORM_SUPERVISORY" not in body_of(glue_bsc, "template_validate") \
        or "_arrange_kpis(doc)" not in body_of(glue_bsc, "template_validate"):
    fail.append("the template is checked as the form it carries, its Section A laid out like the workbook")
patch = read("hrms_addon", "patches", "v1_0", "appraisal_templates_and_self_appraisal.py")
for needle, why in (
    ('sync_fixtures("hrms_addon")', "the fields exist before the patch writes them"),
    ("frappe.db.get_singles_dict(SETTINGS)", "settings someone already saved are kept"),
    ("_lay_out(name)", "each template is laid out like the workbook"),
    ("bsc.seed_supervisory_template()", "the supervisory form gets its template"),
    ('"appraisal_template": ["is", "not set"]', "only a Job Title with no template is given one"),
    ("set custom_self_appraisal = 1", "the appraisals under way keep the self-appraisal they had"),
    ('"workflow_state": approval.PENDING_SELF', "a plan's draft waits on the employee's self-appraisal"),
    ("continue  # goals rated by hand stay rated by hand", "a cycle with goals rated by hand is left as it is"),
):
    if needle not in patch:
        fail.append("the patch: %s (%r not found)" % (why, needle))
if "hrms_addon.patches.v1_0.appraisal_templates_and_self_appraisal" not in read("hrms_addon", "patches.txt"):
    fail.append("the patch is listed in patches.txt")
print("templates: the form each carries, laid out like the workbook, named by every Job Title; the self-appraisal "
      "as Appraisal Settings say now, Frappe HR's tab and count with it; every state a status")

# ── 11. The sheet, Luuka's own form, out and back ─────────────────────
SH = load("appraisal_sheet")
if tuple(floor for floor, *_rest in SH.BSC_BANDS) != tuple(floor for floor, _name in S.BANDS) or \
        tuple(name for _floor, name, *_rest in SH.BSC_BANDS) != tuple(name for _floor, name in S.BANDS):
    fail.append("the sheet rates on the scorecard's own bands")
if tuple(SH.SUPERVISORY_BANDS) != tuple(R.BANDS) or tuple(SH.RATINGS) != tuple(R.RATINGS):
    fail.append("the supervisory sheet rates on LPL/HR/18's bands and scale")
if SH.MAX_OBJECTIVES != R.MAX_OBJECTIVES or (SH.FORM_BSC, SH.FORM_SUPERVISORY) != (S.FORM_BSC, S.FORM_SUPERVISORY):
    fail.append("the sheet and the rules name the forms and the limits alike")
if [key for key, _question in SH.QUESTIONS] != [key for key, _question in R.QUESTIONS]:
    fail.append("the sheet asks LPL/HR/18's General questions")
BSC_DATA = {
    "name": "HR-APR-2026-00012", "form_type": S.FORM_BSC, "period": "Q2", "self_appraisal": 1,
    "company": "Luuka Plastics Limited", "year": 2026, "employee_name": "Ferdinand: Musembi / Senior Procurement",
    "kpis": [{"perspective": "Financial", "kpi": "Savings", "timing": "Monthly", "weight": 30, "q1_percent": 90,
              "q1_comments": "On track."},
             {"perspective": "Financial", "kpi": "Variance\ntracked", "timing": "Monthly", "weight": 20},
             {"perspective": "Customer / Stakeholder", "kpi": "Stock-outs", "timing": "Weekly", "weight": 30}],
    "results": {"Q1": 81.5},
    "competencies": [{"competency": "One", "weight": 12}, {"competency": "Two", "weight": 8}],
    "assignments": [], "remarks": {}, "names": {}, "plan": {}, "actions": [],
}
LPL_DATA = {
    "name": "HR-APR-2026-00013", "form_type": S.FORM_SUPERVISORY, "self_appraisal": 0, "company": "Luuka",
    "employee_name": "John Okello",
    "factors": [{"item": "Attendance and time management", "employee_rating": "4"}],
    "objectives": [{"item": "Daily output"}], "answers": {}, "remarks": {},
}
content = SH.build([BSC_DATA, LPL_DATA, dict(BSC_DATA, name="HR-APR-2026-00014")], logo=b"not a picture")
import io  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

book = load_workbook(io.BytesIO(content))
titles = book.sheetnames
if len(titles) != 3 or len(set(title.lower() for title in titles)) != 3 or \
        any(len(title) > 31 or set(title) & set("[]:*?/\\") for title in titles):
    fail.append("one sheet per appraisal, titled as Excel allows, each once: %s" % titles)
card, lpl = book.worksheets[0], book.worksheets[1]
if (card["R1"].value, card["R2"].value, card["R4"].value, card["R5"].value) != (SH.MARK, "HR-APR-2026-00012", "Q2",
                                                                              SH.LAYOUT) \
        or not card.column_dimensions["R"].hidden:
    fail.append("each sheet names its appraisal, its quarter and its layout in a hidden column")
if SH.QUARTERS != S.QUARTERS or [SH.PERIOD_COLUMNS[quarter] for quarter in SH.QUARTERS] != [
        ("F", "G", "H"), ("I", "J", "K"), ("L", "M", "N"), ("O", "P", "Q")]:
    fail.append("the sheet carries the scorecard's four quarters, three columns each, A to Q")
if sorted(SH.BSC_WIDTHS) != [chr(code) for code in range(ord("A"), ord("Q") + 1)] or sorted(SH.WIDTHS)[-1] != "P":
    fail.append("the scorecard's sheet runs A to Q, the supervisory form's still A to P")
if not card.protection.sheet or card.protection.formatColumns:
    fail.append("the sheet is locked where the system filled it in, and its columns can still be widened")
if card["K8"].value != "HR-APR-2026-00012":
    fail.append("the HR Ref on the sheet is the appraisal")
rows = {card["C%d" % row].value: row for row in range(12, 20) if card["R%d" % row].value == "kpi"}
first = rows.get("Savings")
if not first or card["J%d" % first].protection.locked or not card["G%d" % first].protection.locked:
    fail.append("only the quarter being appraised is open: Q2's % open, Q1's locked")
if first and card["G%d" % first].value != 0.9:
    fail.append("an earlier quarter shows what was recorded for it (Q1 90%% as 0.9): %s" % card["G%d" % first].value)
if first and card["K%d" % first].value != '=IF(J%d="","",E%d*J%d)' % (first, first, first):
    fail.append("a quarter scores weight times percent, as the system does, without the workbook's tenth: %s"
                % card["K%d" % first].value)
if first and card["F%d" % first].value != "On track.":
    fail.append("an earlier quarter's comments show against its KPIs")
if first and (card["I%d" % first].protection.locked is not False or not card["F%d" % first].protection.locked):
    fail.append("the quarter appraised takes a comment against each KPI; an earlier quarter's are locked")
if first and card["B%d" % (first + 1)].value != S.CONTINUATION:
    fail.append("a KPI under the same perspective carries the workbook's arrow")
if first and ((card["E%d" % first].value, card["E%d" % (first + 1)].value) != (30, 20)
              or "E%d:E%d" % (first, first + 1) in [str(span) for span in card.merged_cells.ranges]):
    fail.append("every KPI carries its own weight on its own row")
if first and (card["P11"].value != "Q4 %\nAchieved" or not card["P%d" % first].protection.locked
              or card["Q%d" % first].value != '=IF(P%d="","",E%d*P%d)' % (first, first, first)):
    fail.append("the fourth quarter is on the sheet, locked until it is appraised")
total_row = next(row for row in range(12, 30) if card["A%d" % row].value == "WEIGHT CHECK & QUARTERLY TOTALS")
summary_row = next((row for row in range(total_row, total_row + 8) if card["A%d" % row].value == "Financial"), None)
if not summary_row or card["E%d" % summary_row].value != "=SUMIF($S$12:$S$14,$A%d,$E$12:$E$14)" % summary_row \
        or "SUMIF($S$12:$S$14,$A%d,$H$12:$H$14)" % summary_row not in str(card["F%d" % summary_row].value):
    fail.append("the perspectives below the KPIs sum them up by the key each KPI row carries: %s"
                % (summary_row and card["E%d" % summary_row].value))
scored_row = next((row for row in range(12, card.max_row + 1) if card["A%d" % row].value == "Overall score"), None)
overall_row = next((row for row in range(12, card.max_row + 1)
                    if str(card["A%d" % row].value or "").startswith("OVERALL SCORE")), None)
if not scored_row or card["F%d" % scored_row].value != 81.5 \
        or card["I%d" % scored_row].value != '=IFERROR(Q%d,"")' % overall_row \
        or card["L%d" % scored_row].value is not None or card["O%d" % scored_row].value is not None:
    fail.append("the year so far: Q1 as its own appraisal recorded it, Q2 from this sheet, the later ones blank")
if not scored_row or card["Q%d" % (scored_row + 1)].value != '=IFERROR(AVERAGE(F%d:Q%d),"")' % (scored_row, scored_row):
    fail.append("the year to date is the average of the quarters appraised")
checks = [(check.type, check.formula1, check.formula2) for check in card.data_validations.dataValidation
          if first and "J%d" % first in str(check.sqref).split() and "J%d" % (first + 1) in str(check.sqref).split()]
if checks != [("decimal", "0", "1")]:
    fail.append("the percentage achieved is checked as it is typed, 0%% to 100%%: %s" % checks)
if lpl["G12"].protection.locked is False:
    fail.append("with the self-appraisal off, the employee's column is not filled in")
if lpl["I12"].protection.locked is not False:
    fail.append("the supervisor's rating is open")
if lpl["G12"].value != 4 or lpl["B12"].value != "Attendance and time management":
    fail.append("the supervisory sheet carries the factors and what is rated: %s" % [lpl["B12"].value, lpl["G12"].value])
# the appraiser fills it in, renames the sheet and uploads it
card.title = "Renamed"
card["J%d" % first] = 0.85
card["J%d" % rows["Stock-outs"]] = 1.4
card["I%d" % (first + 1)] = "Variance down"
blank = [row for row in range(1, card.max_row + 1) if card["R%d" % row].value == "assignment"]
card["B%d" % blank[0]] = "Stocktake"
card["C%d" % blank[0]] = "Count the stores"
for row in range(1, card.max_row + 1):
    if card["R%d" % row].value == "competency":
        card["L%d" % row] = {"One": 8, "Two": 11}[card["S%d" % row].value]
    if card["R%d" % row].value == "remark" and card["S%d" % row].value == "supervisor":
        card["C%d" % row] = "Strong quarter."
lpl["I12"] = 5
lpl["K12"] = "Always early"
objective_row = next(row for row in range(1, lpl.max_row + 1) if lpl["R%d" % row].value == "objective")
lpl["I%d" % objective_row] = "seven"
out = io.BytesIO()
book.save(out)
found = SH.read(out.getvalue())
got = found.get("HR-APR-2026-00012") or {}
if got.get("sheet") != "Renamed" or got.get("period") != "Q2":
    fail.append("a renamed sheet is still matched to its appraisal: %s" % sorted(found))
if got.get("kpis") != {("Financial", "Savings"): {"percent": 85.0, "comments": None},
                       ("Financial", "Variance tracked"): {"percent": None, "comments": "Variance down"}}:
    fail.append("each KPI's percentage comes back as a percentage, its comment with it, only for the quarter "
                "appraised, and one over 100 is not taken: %s" % got.get("kpis"))
if not any("Stock-outs (Customer / Stakeholder)" in text for text in got.get("problems") or []):
    fail.append("a percentage over 100 is reported against its KPI, not guessed at")
if got.get("outdated"):
    fail.append("a sheet laid out now is read")
if got.get("competencies") != {"One": 8.0} or not any("Two" in text for text in got.get("problems") or []):
    fail.append("a competency over ten is reported, the rest taken: %s" % got.get("competencies"))
if got.get("remarks") != {"supervisor": "Strong quarter."}:
    fail.append("the appraiser's comments come back: %s" % got.get("remarks"))
if [row["task"] for row in got.get("assignments") or []] != ["Stocktake"]:
    fail.append("an assignment typed on the sheet comes back: %s" % got.get("assignments"))
supervisory = found.get("HR-APR-2026-00013") or {}
factor = (supervisory.get("factors") or [{}])[0]
if (factor.get("supervisor_rating"), factor.get("supervisor_comment")) != ("5", "Always early"):
    fail.append("a supervisory rating comes back on the scale's own terms: %s" % factor)
if not any("seven" in text for text in supervisory.get("problems") or []):
    fail.append("a rating off the scale is reported")
if "HR-APR-2026-00014" not in found:
    fail.append("every sheet of the workbook is read, not only the first")
older = book.worksheets[2]
older["R5"] = None
out = io.BytesIO()
book.save(out)
stale = SH.read(out.getvalue()).get("HR-APR-2026-00014") or {}
if not stale.get("outdated") or stale.get("kpis"):
    fail.append("a sheet laid out before every KPI carried its own weight is not read: %s" % stale)
bare = io.BytesIO()
from openpyxl import Workbook  # noqa: E402

Workbook().save(bare)
if SH.read(bare.getvalue()) != {}:
    fail.append("a workbook this app did not make says nothing")
foreign = Workbook()
foreign.active["R2"] = "HR-APR-2026-00012"
foreign.active["R3"] = S.FORM_BSC
bare = io.BytesIO()
foreign.save(bare)
if SH.read(bare.getvalue()) != {}:
    fail.append("a workbook without the sheet's own mark is not read, whatever its cells hold")
for needle, why in (
    ("_earlier_quarters(doc)", "the earlier quarters come from the employee's earlier appraisals"),
    ("EMPLOYEE_STATES = (approval.DRAFT, approval.PENDING_SELF)", "the employee's part is taken until they submit it"),
    ("SUPERVISOR_STATES = (approval.DRAFT, approval.PENDING_SELF, approval.PENDING_SUPERVISOR)",
     "the supervisor's until they pass it on"),
    ("_logo(docs[0].company)", "the sheet carries the company's logo"),
):
    if needle not in glue_appraisals:
        fail.append("appraisals.py: %s (%r not found)" % (why, needle))
print("the sheet: Luuka's own form per appraisal, locked but for the period appraised, scored by the system's "
      "rules, read back whatever it is renamed to, every sheet of it")

# A cycle's sheet that would come out empty is explained on the form, not
# as the bare 417 page a refused download opens (Oct 2026)
for facts, wanted in (
        ({"cycle": "Q4", "appraisals": 3, "open": 2, "theirs": 1, "readable": 1}, None),
        ({"cycle": "Q4", "appraisals": 0, "open": 0, "theirs": 0, "readable": 0}, "Q4 has no appraisals yet"),
        ({"cycle": "Q4", "appraisals": 3, "open": 0, "theirs": 0, "readable": 0}, "Every appraisal in Q4 is submitted"),
        ({"cycle": "Q4", "appraisals": 3, "open": 2, "theirs": 0, "readable": 0, "supervisor": "Sarah"},
         "None of the open appraisals in Q4 is Sarah's to rate"),
        ({"cycle": "Q4", "appraisals": 3, "open": 2, "theirs": 2, "readable": 0}, "You may not open the appraisals in Q4")):
    got = R.no_sheet_reason(facts)
    if (got is None) != (wanted is None) or (wanted and wanted not in got):
        fail.append("no_sheet_reason(%r): got %r, want %r" % (facts, got, wanted))
cycle_js = read("hrms_addon", "public", "js", "appraisal_cycle.js")
for needle, why in (
        ('.xcall("hrms_addon.hrms_addon.appraisals.sheet_count", args)', "the cycle's button asks before downloading"),
        ("if (values.supervisor) args.supervisor = values.supervisor;", "an empty supervisor is left out of the call"),
        ("if (!found.count) {", "an empty sheet is explained, not downloaded"),
        ("frappe.utils.escape_html(found.reason", "the reason is shown as text")):
    if needle not in cycle_js:
        fail.append("appraisal_cycle.js: %s (%r not found)" % (why, needle))
if cycle_js.index("sheet_count") > cycle_js.index("download_sheet?appraisal_cycle="):
    fail.append("appraisal_cycle.js must ask sheet_count before it opens the download")
for needle, why in (
        ("def sheet_count(appraisal_cycle: str, supervisor: str | None = None) -> dict:",
         "the count the cycle's button asks for"),
        ("names = [appraisal] if appraisal else _sheet_names(appraisal_cycle, supervisor)",
         "the download takes the same appraisals as the count"),
        ("or (not row.custom_supervisor and row.employee in theirs)",
         "an appraisal naming no supervisor yet is its employee's Reports To's"),
        ("else _(_no_sheet(appraisal_cycle, supervisor, names, docs)))", "a refused download says why"),
        ('readable = [name for name in names if frappe.has_permission("Appraisal", "read", name)]',
         "the count is of what the user may open, as the download's is")):
    if needle not in glue_appraisals:
        fail.append("appraisals.py: %s (%r not found)" % (why, needle))
print("the cycle's sheet: asked for before it downloads, and an empty one explained")

# ── 12. The scorecard template drawn as the workbook's form ───────────
# The Appraisal Template shows a Balanced Scorecard as the LPL PMS BSC
# Appraisal Form; it must draw what the offline sheet draws, so the two
# cannot drift: the same columns, widths, palette, perspective colours,
# scale and signatories.
template_js = read("hrms_addon", "public", "js", "appraisal_template.js")
sheet_source = read("hrms_addon", "hrms_addon", "appraisal_sheet.py")
on_template = custom_fields("Appraisal Template")
for fieldname, fieldtype, after in (("custom_form_section", "Section Break", "template_title"),
                                    ("custom_form_view", "HTML", "custom_form_section"),
                                    ("custom_role_section", "Section Break", "custom_form_view")):
    row = on_template.get(fieldname) or {}
    if (row.get("fieldtype"), row.get("insert_after")) != (fieldtype, after):
        fail.append("Appraisal Template.%s must be a %s after %s, so the form comes first" % (fieldname, fieldtype, after))
if (on_template.get("custom_form_section") or {}).get("depends_on") != 'eval:doc.custom_form_type == "Balanced Scorecard"':
    fail.append("the workbook's form is the scorecard's: shown on a Balanced Scorecard template only")


def js_list(name):
    found = re.search(r"const %s = (\[[^\n]*\]|\[.*?\n\]);" % name, template_js, re.S)
    return found.group(1) if found else ""


widths = [int(value) for value in re.findall(r"\d+", js_list("HA_SHEET_WIDTHS"))]
if widths != [SH.BSC_WIDTHS[column] for column in "ABCDEFGHIJKLMNOPQ"]:
    fail.append("the form's columns must be the scorecard sheet's, A to Q, in its widths: %s" % widths)
# the template drawn with the sheet's own heads, the four quarters each
# with comments, a percentage and a weighted score
sheet_row = [card["%s11" % column].value for column in "ABCDEFGHIJKLMNOPQ"]
drawn_first = re.findall(r'\["([^"]*)", "head', js_list("HA_SHEET_HEADS"))[:5]
drawn_quarters = [head.replace("${quarter}", quarter) for quarter in S.QUARTERS
                  for head in re.findall(r"\[`([^`]*)`, \"head", js_list("HA_SHEET_HEADS"))]
if [head.replace("\\n", "\n") for head in drawn_first + drawn_quarters] != sheet_row:
    fail.append("the template's form heads its columns as the sheet does: %s against %s"
                % (drawn_first + drawn_quarters, sheet_row))
drawn_heads = re.findall(r'\["([^"]*)", "head', js_list("HA_SHEET_HEADS"))
sheet_heads = re.findall(r'\("[A-P]", "([^"]*)", TEAL', sheet_source.split("def _section_a(")[1].split("for column, text, fill in heads")[0])
if not sheet_heads or drawn_heads != sheet_heads:
    fail.append("Section A's headings must be the sheet's: %s against %s" % (drawn_heads, sheet_heads))
palette = dict(re.findall(r'^\t(\w+): "#(\w{6})",$', template_js.split("const HA_SHEET = {")[1].split("};")[0], re.M))
if palette != {"navy": SH.NAVY, "teal": SH.TEAL, "tealDark": SH.TEAL_DARK, "gold": SH.GOLD, "label": SH.LABEL,
               "note": SH.NOTE, "soft": SH.SOFT, "cream": SH.CREAM, "green": SH.GREEN_SOFT, "grey": SH.GREY_SOFT}:
    fail.append("the form's palette must be the sheet's: %s" % palette)
colours = re.findall(r'\["(\w+)", "#(\w{6})", "#(\w{6})"\]', js_list("HA_PERSPECTIVE_COLOURS"))
for perspective, (light, dark) in SH.PERSPECTIVE_COLOURS.items():
    match = [(fill, text) for word, fill, text in colours if word in perspective.lower()]
    if match != [(light, dark)]:
        fail.append("the form must colour %s as the sheet does (%s on %s)" % (perspective, dark, light))
if re.findall(r'"([^"]+)"', js_list("HA_SIGNATORIES")) != [label for _key, label in SH.BSC_SIGNATORIES]:
    fail.append("Part D must be signed by the sheet's signatories")
scale = re.findall(r'\["(\w[\w ]*)", "[^"]*", "#(\w{6})", "([^"]*)"\]',
                   template_js.split('"Balanced Scorecard": [')[1].split("],\n\t[")[0])
if scale != [(name, colour, meaning) for _floor, name, colour, meaning in SH.BSC_BANDS]:
    fail.append("the form's scale must be the sheet's: %s" % scale)
for needle, why in (
        ('frm.trigger("show_scale");\n\t\tfrm.trigger("show_form");', "drawn when the template opens and as its form type changes"),
        ("show_form(frm) {\n\t\tha_bsc_form(frm);", "the form is drawn by ha_bsc_form"),
        ('const field = frm.get_field("custom_form_view");', "into the form view field"),
        ("if (frm.doc.custom_form_type !== HA_TEMPLATE_BSC) {\n\t\tfield.$wrapper.empty();",
         "a supervisory template shows no scorecard"),
        ("frappe.utils.escape_html(", "what is typed is shown as text"),
        ("${td(esc(kpi.kpi))}", "a KPI is escaped"),
        ("td(esc(kpi.perspective)", "a perspective is escaped"),
        ("${td(esc(row.competency)", "a competency is escaped"),
        ("esc(row.indicators)", "its indicators are escaped"),
        ("/^(\\/|https?:\\/\\/)/.test(logo)", "only a logo the site serves is shown"),
        ("custom_kpis_remove(frm) {\n\t\tha_template_perspectives(frm);\n\t\tfrm.trigger(\"show_weights\");\n"
         "\t\tfrm.trigger(\"show_form\");", "a KPI taken off redraws the form and its perspectives"),
        ("weight(frm) {\n\t\tha_template_perspectives(frm);", "a KPI's weight sums into its perspective as it is typed"),
        ("custom_competencies_remove(frm) {\n\t\tfrm.trigger(\"show_weights\");\n\t\tfrm.trigger(\"show_form\");",
         "a competency taken off redraws the form")):
    if needle not in template_js:
        fail.append("appraisal_template.js: %s (%r not found)" % (why, needle))
if template_js.count('frm.trigger("show_scale");\n\t\tfrm.trigger("show_form");') != 2:
    fail.append("appraisal_template.js: the form is drawn both when the template opens and when its form type changes")
if re.search(r"(?<![\w-])(eval|new Function)\(|\.innerHTML\s*=", template_js):
    fail.append("appraisal_template.js must not evaluate or write raw HTML")
print("the template: a scorecard drawn as Luuka's BSC Appraisal Form, the offline sheet's columns, palette, scale "
      "and signatories")

# ── What Frappe chose by itself, cleared where it is still open ───────
def rows_rated(employee, supervisor):
    return [{"name": "r%d" % index, "employee_rating": mine, "supervisor_rating": theirs}
            for index, (mine, theirs) in enumerate(zip(employee, supervisor))]


ONES, BLANKS = ["1", "1", "1"], ["", "", ""]
for label, rows, supervisor_had_it, employee_had_it, wanted in (
        ("neither rater has had it: both columns' 1s", rows_rated(ONES, ONES), False, False,
         {"supervisor_rating": ["r0", "r1", "r2"], "employee_rating": ["r0", "r1", "r2"]}),
        ("a rater who has had it and left every rating 1", rows_rated(ONES, ONES), True, True,
         {"supervisor_rating": ["r0", "r1", "r2"], "employee_rating": ["r0", "r1", "r2"]}),
        ("a rater who changed some keeps their 1s", rows_rated(["3", "1", "1"], ["1", "4", "1"]), True, True, {}),
        ("before the supervisor, their 1s go even beside other values",
         rows_rated(["3", "1", "1"], ["1", "4", "1"]), False, True, {"supervisor_rating": ["r0", "r2"]}),
        ("a blank among 1s is still all 1", rows_rated(["1", "", "1"], BLANKS), True, True,
         {"employee_rating": ["r0", "r2"]}),
        ("N/A is a choice", rows_rated(["N/A", "1", "1"], BLANKS), True, True, {}),
        ("nothing rated, nothing to clear", rows_rated(BLANKS, BLANKS), False, False, {})):
    got = R.prefilled_ratings(rows, supervisor_had_it, employee_had_it)
    if got != wanted:
        fail.append("prefilled_ratings, %s: got %r, want %r" % (label, got, wanted))
if R.PREFILLED_RATING != R.RATINGS[0] or R.PREFILLED_RATING != "1":
    fail.append("the rating Frappe filled in was the scale's first, 1")
if R.prefilled_decisions([{"name": "a", "decision": R.PROMOTION}, {"name": "b", "decision": R.CLOSE},
                          {"name": "c", "decision": ""}]) != ["a"]:
    fail.append("prefilled_decisions: every Promotion on an open review, nothing else")
if P.prefilled(P.IMPROVED, [{"name": "a", "progress": P.MET}, {"name": "b", "progress": P.MET, "reviewed_on": "2026-09-20"},
                            {"name": "c", "progress": P.NOT_MET}]) != (True, ["a"]) \
        or P.prefilled(P.EXTENDED, []) != (False, []):
    fail.append("pip_rules.prefilled: the outcome Improved, and Met where nobody reviewed the point")
patch = read("hrms_addon", "patches", "v1_0", "clear_prefilled_choices.py")
if "hrms_addon.patches.v1_0.clear_prefilled_choices" not in read("hrms_addon", "patches.txt").split("[post_model_sync]")[1]:
    fail.append("clear_prefilled_choices must run after the doctypes are migrated (patches.txt, post_model_sync)")
for needle, why in (
        ("appraisal_rules.prefilled_ratings(", "the appraisals are cleared by the tested rule"),
        ("supervisor_had_it=state not in approval.BEFORE_SUPERVISOR",
         "the supervisor has had an appraisal once it is past the employee"),
        ("employee_had_it=bool(appraisal.get(approval.SELF_FIELD)) and state != approval.DRAFT",
         "the employee has had it once out of Draft, and only where they appraise themselves"),
        ("if appraisal.docstatus == 1:\n            filed.append(appraisal.name)\n            continue",
         "a submitted appraisal is named, never changed"),
        ("appraisals._score(doc)", "an appraisal cleared is scored again"),
        ("{field: flt(doc.get(field)) for field in SCORE_FIELDS}", "a score not given is written as 0, never NULL"),
        ("appraisals.gave_self_appraisal(unrated)", "one left with no rating of the employee's was never self-appraised"),
        ("appraisal_rules.prefilled_decisions(rows)", "the reviews are cleared by the tested rule"),
        ("if review.docstatus == 1:", "a filed review is named, never changed"),
        ('add_comment("Info"', "the review says whose decision was cleared"),
        ("pip_rules.prefilled(plan.outcome, objectives)", "the plans are cleared by the tested rule"),
        ('filters={"docstatus": 0}, fields=["name", "outcome"]', "only open plans are touched")):
    if needle not in patch:
        fail.append("clear_prefilled_choices.py: %s (%r not found)" % (why, needle))
print("what Frappe chose by itself: the choices start blank, and the patch clears only what is open and unchosen")

# ── 13. Luuka, 4 Oct 2026: KPIs weighed, the quarters, the remarks, the PIP ──
def frappe_order(doctype_json, custom):
    """Frappe v16's own placing of custom fields (meta.sort_fields and
    _update_field_order_based_on_insert_after): a field after another's
    whole chain when two share an anchor, a break anchored on a standard
    field moved to the end of that field's section."""
    kinds = {field["fieldname"]: field["fieldtype"] for field in doctype_json["fields"]}
    kinds.update({field["fieldname"]: field["fieldtype"] for field in custom})
    order = [name for name in (doctype_json.get("field_order") or []) if name in kinds] \
        or [field["fieldname"] for field in doctype_json["fields"]]
    insertion = {}
    for field in custom:
        target = field.get("insert_after")
        if field["fieldtype"] in ("Section Break", "Column Break") and target in order:
            original = target
            for current in order[order.index(target) + 1:]:
                if kinds[current] == "Section Break" or kinds[current] == kinds[original]:
                    break
                target = current
        insertion.setdefault(target, []).append(field["fieldname"])
    retry = True
    while retry:
        retry = False
        for anchor in list(insertion):
            if anchor in order:
                at = order.index(anchor)
                for name in insertion.pop(anchor):
                    at += 1
                    order.insert(at, name)
                retry = True
    for names in insertion.values():
        order.extend(names)
    return order


appraisal_custom = [row for row in CUSTOM if row["dt"] == "Appraisal"]
# fields sharing an anchor land in the order the database hands them over
# (every fixture's idx is 0): one field per anchor keeps the form as drawn
anchors = {}
for row in appraisal_custom:
    anchors.setdefault(row.get("insert_after"), []).append(row["fieldname"])
shared = {anchor: names for anchor, names in anchors.items() if len(names) > 1}
if shared:
    fail.append("Appraisal custom fields share an anchor, so their order is the database's: %s" % shared)
placed = frappe_order(upstream_doctype("Appraisal"), appraisal_custom)
at = {name: index for index, name in enumerate(placed)}
remark_sections = ("custom_employee_section", "custom_supervisor_section", "custom_hod_section", "custom_hrm_section",
                   "custom_production_section", "custom_gm_section", "custom_ed_section")
content = ("custom_challenges", "custom_bsc_self_score", "custom_year_band", "custom_development_actions",
           "custom_quarter_results", "custom_bsc_competencies", "custom_factors", "custom_objectives")
if min(at[name] for name in remark_sections) < max(at[name] for name in content):
    fail.append("the remarks sit below the appraisal, after both forms, the year so far and the plan: %s"
                % [name for name in placed if name.startswith("custom_")][:80])
if [name for name in placed if name in remark_sections] != list(remark_sections):
    fail.append("the remarks follow the signing order, the HOD before the HR Manager: %s"
                % [name for name in placed if name in remark_sections])
if not at["custom_bsc_kpis"] < at["custom_bsc_perspectives"] < at["custom_assignments_section"]:
    fail.append("Section A shows the KPIs, then the perspectives summing them up")
if not at["custom_round_section"] < at["custom_form_type"] < at["custom_quarter"] < at["custom_on_pip"] \
        < at["custom_section_a"]:
    fail.append("the form, the quarter and the improvement plan sit at the top, with the appraisal's details")
if not at["custom_band"] < at["custom_general_section"] < at["custom_bsc_section_a"]:
    fail.append("the supervisory form's General questions follow its own scores, not the year so far")
gm = next(row for row in appraisal_custom if row["fieldname"] == "custom_gm_section")
if "Balanced Scorecard" not in (gm.get("depends_on") or ""):
    fail.append("the General Manager signs the supervisory form only: their remarks are not on the scorecard")

# the remarks are their signatory's, written when the appraisal is with them
for form, step, field in ((A.FORM_BSC, A.PENDING_SUPERVISOR, "custom_supervisor_remarks"),
                          (A.FORM_BSC, A.PENDING_EMPLOYEE, "custom_employee_remarks"),
                          (A.FORM_BSC, A.PENDING_HOD, "custom_hod_remarks"),
                          (A.FORM_BSC, A.PENDING_ED, "custom_ed_remarks"),
                          (A.FORM_SUPERVISORY, A.PENDING_SELF, "custom_employee_remarks"),
                          (A.FORM_SUPERVISORY, A.PENDING_GM, "custom_gm_remarks")):
    if A.remark_steps(form).get(field) != step:
        fail.append("%s's %s are written at %s" % (form, field, step))
    expect("%s writes their own remarks at %s" % (field, step), A.remark_errors(form, step, [field]))
expect("the HR Manager cannot write the appraiser's remarks",
       A.remark_errors(A.FORM_BSC, A.PENDING_HRM, ["custom_supervisor_remarks"]), "The Appraiser's remarks")
expect("nobody writes remarks in Draft", A.remark_errors(A.FORM_SUPERVISORY, A.DRAFT, ["custom_employee_remarks"]),
       "The Employee's remarks")
expect("a remark the form has no place for", A.remark_errors(A.FORM_BSC, A.PENDING_HRM, ["custom_gm_remarks"]),
       "no place for those remarks")
expect("nothing changed, nothing to say", A.remark_errors(A.FORM_BSC, A.PENDING_HRM, []))
if set(A.remark_steps(A.FORM_BSC)) | set(A.remark_steps(A.FORM_SUPERVISORY)) != set(A.ALL_REMARK_FIELDS):
    fail.append("every remark field belongs to one signatory's step")
appraisal_js = read("hrms_addon", "public", "js", "appraisal.js")
js_remarks = re.findall(r'"(custom_\w+_remarks)"', appraisal_js.split("const HA_REMARK_FIELDS")[1].split("];")[0])
if set(js_remarks) != set(A.ALL_REMARK_FIELDS):
    fail.append("the form opens and closes every signatory's remarks: %s" % js_remarks)
for needle, why in (
    ('frm.set_df_property(fieldname, "read_only", steps[fieldname] === state ? 0 : 1);',
     "a signatory's remarks open only at their own step"),
    ("frm.doc.__onload && frm.doc.__onload.remark_steps", "the steps come from the server, the rules' own"),
    ('grid.update_docfield_property(fieldname, "read_only", closed);', "only the quarter appraised is open"),
    ('quarter.toLowerCase() + "_percent", quarter.toLowerCase() + "_comments"', "its percentage and its comments"),
    ("flt((flt(row.weight) * flt(percent)) / 100, 2)", "a KPI's weighted score shows as it is typed"),
    ('frm.set_intro(', "an employee on an improvement plan is said at the top"),
    ('"red"', "in red"),
    ('["custom_bsc_kpis", "self_percent"]', "the employee's own percentage is a column of the KPIs"),
):
    if needle not in appraisal_js:
        fail.append("appraisal.js: %s (%r not found)" % (why, needle))
if "custom_period" in appraisal_js or "custom_period" in glue_appraisals:
    fail.append("the scorecard's period is gone: the form and the glue read the quarter")

# the glue: the quarter, the earlier quarters carried in, the year so far
validate_now = body_of(glue_appraisals, "appraisal_validate")
for needle in ("_settle_quarter(doc)", "_mark_pip(doc)", "_carry_earlier_quarters(doc)", "_year_so_far(doc)",
               "_check_remarks(doc)"):
    if needle not in validate_now:
        fail.append("appraisal_validate must call %s" % needle)
if validate_now.index("_carry_earlier_quarters(doc)") > validate_now.index("bsc.score(doc)") \
        or validate_now.index("_year_so_far(doc)") < validate_now.index("bsc.score(doc)") \
        or validate_now.index("_check_remarks(doc)") > validate_now.index("_check_step(doc)"):
    fail.append("the earlier quarters are carried in before scoring, the year so far after it, the remarks checked "
                "before the step")
for name, needles in (
    ("_settle_quarter", ('doc.custom_quarter = before.custom_quarter', "approval.BEFORE_SUPERVISOR",
                         'bsc_rules.quarter_of(getdate(start).month)', 'cycle.custom_quarter')),
    ("_year_appraisals", ('filters["custom_plan"] = doc.custom_plan', '"between", ["%s-01-01" % year',
                          'order_by="docstatus asc, modified asc"', '"name": ["!=", doc.name or ""]')),
    ("_carry_earlier_quarters", ("found = (earlier.get(each) or {}).get(key) or {}", "if each == quarter:", "row.set(bsc_rules.percent_field(each)",
                                 "row.set(bsc_rules.comments_field(each)")),
    ("_year_so_far", ("bsc_rules.QUARTERS[:bsc_rules.QUARTERS.index(quarter) + 1]",
                      "doc.custom_annual_score = bsc_rules.year_to_date(", "doc.custom_year_band =")),
    ("_check_remarks", ('doc.flags.get("from_sheet")', "approval.remark_errors(")),
    ("appraisal_on_change", ("row.docstatus != 0", "other.db_update()", 'other.update_child_table(table)')),
    ("appraisal_onload", ('doc.set_onload("remark_steps", approval.remark_steps(',)),
    ("_apply_sheet", ("doc.flags.from_sheet = True",)),
    ("_not_taken", ('values.get("outdated")',)),
):
    found = body_of(glue_appraisals, name)
    for needle in needles:
        if needle not in found:
            fail.append("appraisals.%s: %r not found" % (name, needle))
for event, function in (("onload", "appraisal_onload"), ("on_change", "appraisal_on_change")):
    if (events.get("Appraisal") or {}).get(event) != "hrms_addon.hrms_addon.appraisals.%s" % function:
        fail.append("doc_events Appraisal %s must be appraisals.%s" % (event, function))
if ("BSC Appraisal KPI", "self_percent") not in appraisals_self_ratings():
    fail.append("a self-appraisal on the scorecard is the employee's own percentage against each KPI")

# the improvement plan, in red wherever its employee is listed
pip_glue = read("hrms_addon", "hrms_addon", "pips.py")
if P.OPEN != (P.DRAFT, P.AGREED, P.IN_PROGRESS):
    fail.append("an employee is on a plan from the day it is raised until it is closed or cancelled")
for name, needles in (
    ("open_plan", ('"docstatus": 0, "status": ["in", list(rules.OPEN)]', 'order_by="creation desc", limit=1')),
    ("mark_appraisals", ('"docstatus": ["!=", 2]', 'update_modified=False', '"custom_improvement_plan": plan')),
    ("open_plans", ('frappe.has_permission("Performance Improvement Plan", "read")', "[:500]",
                    "row.name if readable else 1")),
    ("plan_on_change", ("mark_appraisals(",)),
):
    found = body_of(pip_glue, name)
    for needle in needles:
        if needle not in found:
            fail.append("pips.%s: %r not found" % (name, needle))
if not re.search(r"@frappe\.whitelist\(\)\ndef open_plans\(employees: list \| str \| None = None\) -> dict:", pip_glue):
    fail.append("pips.open_plans is whitelisted, reads only, and takes the list as Frappe sends it")
controller = read("hrms_addon", "hrms_addon", "doctype", "performance_improvement_plan",
                  "performance_improvement_plan.py")
for needle in ("def on_change(self):\n        pips.plan_on_change(self)",
               "def after_delete(self):\n        pips.plan_after_delete(self)"):
    if needle not in controller:
        fail.append("a plan raised, closed, cancelled or deleted marks its employee's appraisals")
pip_js = read("hrms_addon", "public", "js", "hrms_addon_pip.js")
if "/assets/hrms_addon/js/hrms_addon_pip.js" not in (hooks.get("app_include_js") or []):
    fail.append("the red marking is loaded on every desk page")
for needle in ('"hrms_addon.hrms_addon.pips.open_plans"', 'grid.update_docfield_property(field || "employee_name", '
               '"formatter"', 'indicator-pill red', "frappe.utils.escape_html("):
    if needle not in pip_js:
        fail.append("hrms_addon_pip.js: %r not found" % needle)
for path, needle in (
        (("public", "js", "appraisal_cycle.js"), 'frappe.after_ajax(() => hrms_addon.pip.mark_rows(frm, "appraisees"));'),
        (("hrms_addon", "doctype", "appraisal_plan", "appraisal_plan.js"),
         'hrms_addon.pip.mark_rows(frm, "employees", "employee");'),
        (("hrms_addon", "doctype", "performance_review", "performance_review.js"),
         'hrms_addon.pip.mark_rows(frm, "employees");'),
        (("hrms_addon", "doctype", "performance_review", "performance_review.js"),
         "Number(value) < HA_REVIEW_PASS_MARK")):
    if needle not in read("hrms_addon", *path):
        fail.append("%s: %r not found" % (path[-1], needle))
if (hooks.get("doctype_list_js") or {}).get("Appraisal") != "public/js/appraisal_list.js":
    fail.append("the Appraisal list has its own settings, for the improvement plan and the pass mark")
appraisal_list = read("hrms_addon", "public", "js", "appraisal_list.js")
for needle in ("custom_on_pip(value, df, doc)", "Object.assign(frappe.listview_settings[\"Appraisal\"] || {}",
               "Number(value) < HA_PASS_MARK"):
    if needle not in appraisal_list:
        fail.append("appraisal_list.js: %r not found" % needle)
for name in ("HA_PASS_MARK", "HA_REVIEW_PASS_MARK"):
    source = appraisal_list if name == "HA_PASS_MARK" else read(
        "hrms_addon", "hrms_addon", "doctype", "performance_review", "performance_review.js")
    if "const %s = %d;" % (name, R.PIP_BELOW) not in source:
        fail.append("%s is the pass mark the rules recommend a plan below (%d)" % (name, R.PIP_BELOW))

# the patch that moves a site onto all of it
patch_text = read("hrms_addon", "patches", "v1_0", "scorecard_by_kpi.py")
listed = read("hrms_addon", "patches.txt").split()
if "hrms_addon.patches.v1_0.scorecard_by_kpi" not in listed:
    fail.append("scorecard_by_kpi runs on migrate")
for needle, why in (
    ('sync_fixtures("hrms_addon")', "the fields it writes exist first"),
    ("bsc_rules.spread_weights(", "a template's perspective weight is shared out between its KPIs"),
    ("round(flt(recorded[\"annual_score\"]) * 10, 2)", "an annual score out of ten becomes the fourth quarter's percentage"),
    ("doc.custom_quarter = bsc_rules.QUARTERS[-1]", "the old Annual is the fourth quarter"),
    ("frappe.db.has_column(", "the old columns are read only where they are"),
    ("if doc.docstatus == 0:", "a submitted appraisal keeps the totals it was submitted with"),
    ("appraisals._year_so_far(doc)", "every appraisal gets its year so far"),
    ("pips.mark_appraisals(", "and says who is on a plan"),
    ('frappe.delete_doc("Custom Field", "Appraisal-custom_period"', "the old period goes"),
    ("row.docstatus = doc.docstatus", "the rows written carry their record's state"),
    ("doc.update_child_table(table)", "and are written without the record's checks"),
):
    if needle not in patch_text:
        fail.append("scorecard_by_kpi.py: %s (%r not found)" % (why, needle))
print("Oct 2026: one field per anchor and the remarks at the bottom, each opened only at its signatory's step; the "
      "quarter, the earlier quarters carried in, the year so far; the improvement plan in red; the patch")

print()
if fail:
    print("FAILURES:")
    for f in fail:
        print("  -", f)
    sys.exit(1)
print("ALL PERFORMANCE CHECKS PASSED")
