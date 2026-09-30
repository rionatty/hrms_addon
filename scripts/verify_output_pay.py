"""Verify output pay — Per Meter, Per Piece and the hourly casuals — without
a bench:

    python scripts/verify_output_pay.py

The minutes of 16 and 20 July 2026 (Human Resource, Reward and
Compensation), §4.6, and Luuka's own sheet, PER METER AUGUST 2026. Each
size of material has its own rate, the Work Unit, and the machine it was
made on does not matter (Luuka, 30 Sep 2026). Each person's Work Done on
each size each day is paid at that day's Work Unit, the day's LOOMS OT is
typed in beside it and paid with it, and the month, 26th to 25th, is added
on top of the basic salary. Per Piece is the same with piece categories.
The hourly casuals are paid hours x the standard rate, their overtime
priced by the overtime module.

  1  the sizes: each its own Work Unit, from a date
  2  a day: Work Done x Work Unit, the Looms OT, every wrong line refused
  3  one place for a person's day: a supervisor's reports or the sheet
  4  the month: days folded per person per size, then per person
  5  the sheet checked line by line, as Luuka's own file needs
  6  the sheet built with its formulas and read back: the same pay both ways
  7  the hourly casuals: the standard hours paid, the overtime not twice
  8  the paper: the DocTypes carry what the sheet carries, and no machines
  9  the glue: priced on the day, paid as Additional Salary, never twice
 10  the wiring: events, the patches, the field, and the way in
"""
import datetime
import importlib.util
import io
import json
import os
import re
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
APP = os.path.join(REPO, "hrms_addon", "hrms_addon")
D = datetime.date
fail = []


def read(*parts):
    return open(os.path.join(REPO, *parts), encoding="utf-8").read()


def load(name):
    spec = importlib.util.spec_from_file_location(name, os.path.join(APP, name + ".py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)  # proves it has no Frappe import
    return module


def expect(label, errors, needle=None):
    if needle is None:
        if errors:
            fail.append("%s: expected nothing wrong, got %s" % (label, errors))
    elif not any(needle in error for error in errors):
        fail.append("%s: expected %r among %s" % (label, needle, errors))


O = load("output_rules")
S = load("output_sheet")
A = load("attendance_rules")
print("loaded output_rules.py and output_sheet.py without Frappe")

# ── 1. The sizes ──────────────────────────────────────────────────────
if list(O.SHEET_SIZES) != [("45CM-60CM", 5.2), ("74CM", 6.5), ("60CM", 6.9), ("61CM-79CM", 7.8),
                           ("ABOVE 80CM", 8.7), ("ABOVE 39CM", 8.0)]:
    fail.append("the six sizes and their Work Units are the August 2026 sheet's, in its order: %s"
                % (O.SHEET_SIZES,))
if str(A.cycle_window(2026, 8)[0]) != O.SHEET_FROM:
    fail.append("the August sheet's Work Units start with its payroll month, on 26 July")
if O.size_key(" 45 cm - 60cm ") != "45CM-60CM":
    fail.append("a size is the same size however it is spaced or cased")
UNITS_74 = [{"work_unit": 7.0, "valid_from": "2026-09-01"}, {"work_unit": 6.5, "valid_from": "2026-07-26"}]
if O.work_unit_on(UNITS_74, "2026-08-30") != 6.5:
    fail.append("August's work at 74CM is paid August's Work Unit")
if O.work_unit_on(UNITS_74, "2026-09-01") != 7.0:
    fail.append("a Work Unit that starts on 1 September prices that day on, whatever order the rows are in")
if O.work_unit_on(UNITS_74, "2026-07-25") is not None:
    fail.append("work from before any Work Unit began has none, not nought")
if O.line_total(850, 5.2) != 4420.0 or O.line_total("1,000", 6.5) != 6500.0:
    fail.append("a line's TOTAL is Work Done x Work Unit, as the sheet's F column has it")
expect("a sound size", O.size_errors(O.PER_METER, "74CM", UNITS_74))
expect("a sound piece category",
       O.size_errors(O.PER_PIECE, "CAT A", [{"work_unit": 150, "valid_from": "2026-01-01"}]))
for label, section, size, rows, needle in (
    ("a size with no name", O.PER_METER, "", UNITS_74, "Name the size"),
    ("a piece category with no name", O.PER_PIECE, " ", UNITS_74, "Name the piece category"),
    ("a size with no Work Unit", O.PER_METER, "74CM", [], "Give the size its Work Unit"),
    ("a Work Unit of nothing", O.PER_METER, "74CM", [{"work_unit": 0, "valid_from": "2026-07-26"}],
     "more than nothing"),
    ("a Work Unit with no start", O.PER_METER, "74CM", [{"work_unit": 6.5}], "say the day it starts"),
    ("two Work Units from the same day", O.PER_METER, "74CM",
     [{"work_unit": 6.5, "valid_from": "2026-07-26"}, {"work_unit": 7, "valid_from": "2026-07-26"}],
     "starts the same day"),
    ("a size for the hourly casuals", O.HOURLY, "74CM", UNITS_74, "Per Meter or the Per Piece"),
):
    expect(label, O.size_errors(section, size, rows), needle)
print("sizes: each its own Work Unit from a date, the August sheet's six to start with")

# ── 2. A day ──────────────────────────────────────────────────────────
IN_USE = dict({label: O.PER_METER for label, _unit in O.SHEET_SIZES}, **{"CAT A": O.PER_PIECE})
PEOPLE = {"E1": O.PER_METER, "E2": O.PER_METER, "P1": O.PER_PIECE, "M1": "Monthly"}
ok = {"employee": "E1", "size": "74CM", "output": 100, "rate": 6.5}
expect("a sound day", O.report_errors(O.PER_METER, [ok, dict(ok, size="60CM", rate=6.9)],
                                      [{"employee": "E1", "amount": 2000}], PEOPLE, IN_USE))
expect("a sound Per Piece day", O.report_errors(
    O.PER_PIECE, [{"employee": "P1", "size": "CAT A", "output": 40, "rate": 150}], [], PEOPLE, IN_USE))
for label, change, needle in (
    ("nobody on the line", {"employee": None}, "say who made which size"),
    ("no size on the line", {"size": ""}, "say who made which size"),
    ("less than nothing", {"output": -1}, "less than nothing"),
    ("a size not in use", {"size": "90CM"}, "not a size in use"),
    ("a piece category on a Per Meter report", {"size": "CAT A"}, "Per Piece piece category, not Per Meter"),
    ("no Work Unit that day", {"rate": None}, "has no Work Unit on this day"),
    ("somebody paid monthly", {"employee": "M1"}, "paid Monthly"),
):
    expect(label, O.report_errors(O.PER_METER, [dict(ok, **change)], [], PEOPLE, IN_USE), needle)
expect("a Per Piece report asks for the piece category",
       O.report_errors(O.PER_PIECE, [{"employee": "P1", "size": None, "output": 1}], [], PEOPLE, IN_USE),
       "which piece category")
expect("the same person on the same size twice",
       O.report_errors(O.PER_METER, [ok, dict(ok)], [], PEOPLE, IN_USE), "already on this report")
for label, section, overtime, needle in (
    ("Looms OT on a Per Piece report", O.PER_PIECE, [{"employee": "P1", "amount": 100}],
     "Looms OT belongs on a Per Meter report"),
    ("Looms OT for nobody", O.PER_METER, [{"employee": None, "amount": 100}], "say whose it is"),
    ("Looms OT of less than nothing", O.PER_METER, [{"employee": "E1", "amount": -5}], "less than nothing"),
    ("Looms OT for somebody paid monthly", O.PER_METER, [{"employee": "M1", "amount": 100}],
     "paid Monthly, not Per Meter"),
    ("Looms OT twice for one person", O.PER_METER,
     [{"employee": "E1", "amount": 100}, {"employee": "E1", "amount": 50}], "already has Looms OT"),
):
    expect(label, O.report_errors(section, [], overtime, PEOPLE, IN_USE), needle)
numbered = O.overtime_errors(O.PER_METER, [{"employee": "E1", "amount": -1, "idx": 4}], PEOPLE, where="run")
if not numbered or not numbered[0].startswith("Looms OT row 4:"):
    fail.append("a run checks only the Looms OT typed in, and names the row as its table does: %s" % numbered)
print("a day: Work Done x Work Unit on each size, the Looms OT beside it, every wrong line refused")

# ── 3. One place for a person's day ───────────────────────────────────
OTHER = [{"report": "R1", "source": O.ENTERED, "shift": "Day", "employee": "E1", "size": "74CM"},
         {"report": "R1", "source": O.ENTERED, "shift": "Day", "employee": "E1", "size": None}]
mine = [{"employee": "E1", "size": "74CM"}]
expect("the night shift on the same size", O.day_clashes(mine, [], O.ENTERED, "Night", OTHER))
expect("another size the same shift",
       O.day_clashes([{"employee": "E1", "size": "60CM"}], [], O.ENTERED, "Day", OTHER))
expect("somebody else", O.day_clashes([{"employee": "E2", "size": "74CM"}], [], O.ENTERED, "Day", OTHER))
expect("the same size the same shift", O.day_clashes(mine, [], O.ENTERED, "Day", OTHER),
       "74CM is already on R1 for this shift")
expect("Looms OT twice the same shift", O.day_clashes([], [{"employee": "E1"}], O.ENTERED, "Day", OTHER),
       "Looms OT is already on R1")
clash = O.day_clashes(mine + [{"employee": "E1", "size": "60CM"}], [{"employee": "E1"}], O.UPLOADED, "Day", OTHER)
if len(clash) != 1 or "from the daily reports" not in clash[0]:
    fail.append("a day on a supervisor's report is not taken from the sheet as well, and that is said "
                "once: %s" % clash)
expect("a supervisor's report on a day the sheet has",
       O.day_clashes(mine, [], O.ENTERED, "Night", [dict(OTHER[0], source=O.UPLOADED)]), "from the uploaded sheet")
print("one place for a person's day: a supervisor's reports or the sheet, and a size once a shift")

# ── 4. The month ──────────────────────────────────────────────────────
DAYS = [{"employee": "E1", "size": "74CM", "output": 100, "amount": 650},
        {"employee": "E1", "size": "74CM", "output": 100, "amount": 700},  # after 74CM went to 7
        {"employee": "E1", "size": "60CM", "output": 50, "amount": 345},
        {"employee": "E2", "size": "74CM", "output": 10, "amount": 65}]
merged = O.merge_lines(DAYS)
first = [row for row in merged if row["employee"] == "E1" and row["size"] == "74CM"]
if not first or first[0]["shifts"] != 2 or first[0]["amount"] != 1350 or first[0]["rate"] != 6.75:
    fail.append("a person's days on one size fold into one line, each day at its own Work Unit: %s" % first)
if len(merged) != 3:
    fail.append("a second size is a line of its own: %s" % merged)
extra = O.overtime_by_employee([{"employee": "E1", "amount": 2000}, {"employee": "E1", "amount": 1500},
                                {"employee": "E3", "amount": 1000}, {"employee": None, "amount": 99}])
if extra != {"E1": 3500.0, "E3": 1000.0}:
    fail.append("each person's Looms OT is their days' added up: %s" % extra)
people = {row["employee"]: row for row in O.per_employee(merged, extra)}
seat = people.get("E1", {})
if (seat.get("output_amount"), seat.get("looms_ot"), seat.get("amount")) != (1695, 3500, 5195):
    fail.append("a person's month is every day's TOTAL EARN, the metres and the Looms OT together: %s" % seat)
if people.get("E3", {}).get("amount") != 1000:
    fail.append("Looms OT is paid in a month with no metres too: %s" % people.get("E3"))
if (people.get("E2", {}).get("amount"), people.get("E2", {}).get("looms_ot")) != (65, 0):
    fail.append("and metres with no Looms OT are paid as they are: %s" % people.get("E2"))
if A.cycle_window(2026, 9) != (D(2026, 8, 26), D(2026, 9, 25)):
    fail.append("the payroll month is the 26th to the 25th")
print("the month: days folded per person per size, then each person's metres and Looms OT")

# ── 5. The sheet checked line by line ─────────────────────────────────
if O.employee_key(" lpl12261. ") != "LPL12261":
    fail.append("an Employee Id is matched without the dots and spaces the sheet adds")
IDS = {"LPL001": "E1", "LPL002": "E2", "LPL003": "E3", "LPL004": "E4"}
CATEGORIES = {"E1": O.PER_METER, "E2": O.PER_METER, "E3": "Monthly", "E4": O.PER_METER}
NAMES = {"E1": "ALICE", "E2": "BOB", "E3": "CAROL", "E4": "EVE"}
KEYED = {O.size_key(label): label for label, _unit in O.SHEET_SIZES}
RATES = {label: [{"work_unit": unit, "valid_from": O.SHEET_FROM}] for label, unit in O.SHEET_SIZES}
RATES["ABOVE 39CM"] = [{"work_unit": 8.0, "valid_from": "2026-08-01"}]  # none yet on 30 July
PERIOD = ("2026-07-26", "2026-08-25")


def line(row, block, name, emp_id, date, unit, done, size, looms_ot=None):
    return {"row": row, "block": block, "name": name, "employee_id": emp_id, "date": date,
            "work_unit": unit, "work_done": done, "size": size, "looms_ot": looms_ot}


SHEET = [
    line(2, 1, "ALICE", "LPL001", "26-Jul-2026", 5.2, 100, "45CM-60CM", 2000),  # taken, with Looms OT
    line(3, 1, "ALICE", "LPL001", "26-Jul-2026", 6.5, 50, "74CM"),  # taken
    line(4, 1, "ALICE", "LPL001", "26-Jul-2026", 6.5, 20, "74CM"),  # the same day and size again
    line(5, 1, "ALICE", "LPL002", "27-Jul-2026", 5.2, 30, "45CM-60CM"),  # BOB's ID in ALICE's rows
    line(6, 1, "ALICE", "LPL002", "28-Jul-2026", 5.2, None, "45CM-60CM"),  # and again, with no work
    line(7, 2, "BOB", "LPL002.", "26-Jul-2026", 5.2, 40, " 45 cm-60cm "),  # a dot; the size spaced
    line(8, 2, "BOB", "LPL002", "26-Jul-2026", 5.3, 10, "60CM"),  # the sheet's Work Unit is not ours
    line(9, 2, "BOB", "LPL002", "26-Jul-2026", 7.8, 15, None),  # no size
    line(10, 2, "BOB", "LPL002", "26-Jul-2026", 7.8, 15, "90CM"),  # a size not in use
    line(11, 2, "BOB", "LPL002", "30-Jul-2026", 8.0, 12, "ABOVE 39CM"),  # no Work Unit that day
    line(12, 3, "CAROL", "LPL003", "26-Jul-2026", 5.2, 10, "45CM-60CM"),  # paid monthly
    line(13, 4, "DAN", "LPL999", "26-Jul-2026", 5.2, 10, "45CM-60CM"),  # nobody here
    line(14, 5, "EVE", "LPL004", "26-Aug-2026", 5.2, 10, "45CM-60CM"),  # next month
    line(15, 5, "EVE", "LPL004", "25-Aug-2026", 5.2, "ten", "45CM-60CM"),  # not a number
    line(16, 5, "EVE", "LPL004", "25-Aug-2026", 5.2, 0, "45CM-60CM", "abc"),  # Looms OT not an amount
    line(17, 5, "EVE", "LPL004", "24-Aug-2026", 5.2, "1,000", "45CM-60CM"),  # a thousand, comma and all
    line(18, 6, "ALICE", "LPL001", "27-Jul-2026", 5.2, 10, "45CM-60CM", 500),  # ALICE heads a second set
    line(19, 6, "ALICE", "LPL001", "27-Jul-2026", 6.5, None, "74CM", 700),  # her day's Looms OT twice
    line(20, 6, "ALICE", None, "27-Jul-2026", 6.9, 5, "60CM"),  # no ID
]
taken = O.take_sheet(SHEET, IDS, CATEGORIES, KEYED, RATES, PERIOD, NAMES)
want_work = {("E1", D(2026, 7, 26), "45CM-60CM"): (100.0, 5.2, 2), ("E1", D(2026, 7, 26), "74CM"): (50.0, 6.5, 3),
             ("E2", D(2026, 7, 26), "45CM-60CM"): (40.0, 5.2, 7), ("E2", D(2026, 7, 26), "60CM"): (10.0, 6.9, 8),
             ("E4", D(2026, 8, 24), "45CM-60CM"): (1000.0, 5.2, 17),
             ("E1", D(2026, 7, 27), "45CM-60CM"): (10.0, 5.2, 18)}
if taken["work"] != want_work:
    fail.append("the sheet's work is taken where it is sound, each line at the system's Work Unit for its "
                "day:\n    got  %s\n    want %s" % (taken["work"], want_work))
if taken["overtime"] != {("E1", D(2026, 7, 26)): (2000.0, 2), ("E1", D(2026, 7, 27)): (500.0, 18)}:
    fail.append("the sheet's Looms OT is taken once a person and day: %s" % taken["overtime"])
for needle in (
    "Row 4: ALICE (E1)'s 74CM for 26-Jul-2026 is on the sheet twice (rows 3 and 4)",
    "Row 5: the ID LPL002 is BOB (E2)'s, in ALICE's rows. Left out: correct the ID",
    "Row 9: no size, so its 15 metres are left out.",
    "Row 10: 90CM is not a size in use, so its 15 metres are left out.",
    "Row 11: ABOVE 39CM has no Work Unit on 30-Jul-2026.",
    "CAROL (E3) is paid Monthly, not Per Meter",
    "LPL999 is not an employee here",
    "Row 14: 26-Aug-2026 is not a day of this month (26-Jul-2026 to 25-Aug-2026).",
    "Row 15: Work Done 'ten' is not a number of metres.",
    "Row 16: Looms OT 'abc' is not an amount.",
    "Row 19: ALICE (E1)'s Looms OT for 27-Jul-2026 is on the sheet twice (rows 18 and 19)",
    "Row 20: no Employee Id, so its work is left out.",
):
    expect("the sheet's problems", taken["problems"], needle)
if len(taken["problems"]) != 12:
    fail.append("each line left out is said once, and nothing else: %s" % taken["problems"])
for needle in (
    "Rows of ALICE carry the ID LPL002, which is BOB (E2)'s (from row 6).",
    "BOB (E2) is written 'LPL002.' on the sheet, read as LPL002.",
    "The sheet's Work Unit for 60CM is 5.3; the system's on those days is 6.9, which is what is paid.",
    "ALICE (E1) heads more than one set of rows (from rows 2 and 18)",
):
    expect("the sheet's notes", taken["notes"], needle)
if len(taken["notes"]) != 4:
    fail.append("what was seen and changed nothing is said once: %s" % taken["notes"])
days = O.sheet_days(taken["work"], taken["overtime"])
if list(days) != [D(2026, 7, 26), D(2026, 7, 27), D(2026, 8, 24)]:
    fail.append("the sheet makes one report a day it has work for, in order: %s" % list(days))
first_day = days.get(D(2026, 7, 26), {})
if [(row["employee"], row["size"], row["rate"]) for row in first_day.get("lines", [])] != [
        ("E1", "45CM-60CM", 5.2), ("E1", "74CM", 6.5), ("E2", "45CM-60CM", 5.2), ("E2", "60CM", 6.9)]:
    fail.append("a day's lines keep the sheet's order: %s" % first_day.get("lines"))
if first_day.get("overtime") != [{"employee": "E1", "amount": 2000.0}]:
    fail.append("and its Looms OT goes on the same day's report: %s" % first_day.get("overtime"))
kept = {"work": dict(taken["work"]), "overtime": dict(taken["overtime"])}
changed = O.keep_reported_days(kept, {
    ("E1", D(2026, 7, 26)): {"45CM-60CM": 100, "74CM": 50},  # as the downloaded sheet showed it
    ("E2", D(2026, 7, 26)): {"45CM-60CM": 40, "60CM": 12},  # the sheet says 10
    ("E1", D(2026, 7, 27)): {"45CM-60CM": 10, None: 7},  # and a line from before the sizes
    ("E3", D(2026, 7, 26)): {"74CM": 30},  # not on the sheet at all
}, {("E1", D(2026, 7, 26)): 2000, ("E1", D(2026, 7, 27)): 500}, set(KEYED.values()))
if changed != [("E2", D(2026, 7, 26))]:
    fail.append("a day on a submitted report comes back from the downloaded sheet unchanged and needs no word; "
                "only a day the sheet changed is said: %s" % changed)
if set(kept["work"]) != {("E4", D(2026, 8, 24), "45CM-60CM")} or kept["overtime"]:
    fail.append("and every day a submitted report has stays the report's, changed or not: %s" % kept)
kept = {"work": dict(taken["work"]), "overtime": dict(taken["overtime"])}
if O.keep_reported_days(kept, {("E1", D(2026, 7, 26)): {"45CM-60CM": 100, "74CM": 50}},
                        {("E1", D(2026, 7, 26)): 1500}, set(KEYED.values())) != [("E1", D(2026, 7, 26))]:
    fail.append("a day whose Looms OT the sheet changed is said too")
print("the sheet checked: wrong IDs, doubled days, missing sizes and other months left out and said")

# ── 6. The sheet built and read back ──────────────────────────────────
from openpyxl import Workbook, load_workbook  # noqa: E402 — the sheet's own library, as on the bench

DATES = [D(2026, 7, 26), D(2026, 7, 27), D(2026, 7, 28)]
RATES_BUILT = {label: {day: unit for day in DATES} for label, unit in O.SHEET_SIZES}
RATES_BUILT["74CM"][D(2026, 7, 28)] = 7.0  # 74CM's Work Unit goes up on the 28th
WORK = {("E1", D(2026, 7, 26), "45CM-60CM"): 100, ("E1", D(2026, 7, 26), "74CM"): 50.5,
        ("E1", D(2026, 7, 28), "74CM"): 20, ("E2", D(2026, 7, 27), "ABOVE 80CM"): 300}
EXTRA = {("E1", D(2026, 7, 26)): 2000, ("E2", D(2026, 7, 27)): 1500}
built = S.build({
    "title": "AUGUST 2026", "tag": {"run": "LPL-OPR-2026-00001", "branch": "Namanve",
                                    "from": D(2026, 7, 26), "to": D(2026, 8, 25)},
    "employees": [{"employee": "E1", "name": "ALICE", "id": "LPL001"},
                  {"employee": "E2", "name": "BOB", "id": "LPL002"}],
    "dates": DATES, "sizes": [{"size": label, "rates": RATES_BUILT[label]} for label, _unit in O.SHEET_SIZES],
    "work": WORK, "overtime": EXTRA,
})
book = load_workbook(io.BytesIO(built))
sheet = book.worksheets[0]
if sheet.title != "AUGUST 2026" or [cell.value for cell in sheet[1]] != list(S.HEADINGS):
    fail.append("the sheet has Luuka's headings, with the size column named: %s" % [c.value for c in sheet[1]])
for ref, want in (("A2", "ALICE"), ("B2", "LPL001"), ("C2", "26-Jul-2026"), ("D2", 5.2), ("E2", 100),
                  ("F2", "=E2*D2"), ("G2", "=SUM(F2:F7)"), ("H2", 2000), ("I2", "=G2+H2"), ("J2", "45CM-60CM"),
                  ("A3", None), ("G3", None), ("H3", None), ("J3", "74CM"), ("D15", 7.0), ("J15", "74CM"),
                  ("A20", S.MONTH_TOTAL), ("G20", "=SUM(G2:G19)"), ("H20", "=SUM(H2:H19)"),
                  ("I20", "=SUM(I2:I19)"), ("A21", "BOB"), ("A40", S.GRAND_TOTAL),
                  ("I40", '=SUMIF($A$2:$A$39,"Month Total",I$2:I$39)')):
    if sheet[ref].value != want:
        fail.append("the sheet's %s is %r, not %r" % (ref, sheet[ref].value, want))
merges = {str(cell_range) for cell_range in sheet.merged_cells.ranges}
if not {"G2:G7", "H2:H7", "I2:I7", "G14:G19"} <= merges:
    fail.append("a day's BASIC, LOOMS OT and TOTAL EARN span its rows, as on Luuka's sheet")
if not sheet.protection.sheet or sheet["E2"].protection.locked or sheet["H2"].protection.locked \
        or not sheet["D2"].protection.locked or not sheet["F2"].protection.locked or not sheet["I20"].protection.locked:
    fail.append("only Work Done and LOOMS OT can be typed in; the Work Units and the formulas are kept")
if book[S.TAG_SHEET].sheet_state != "hidden":
    fail.append("the sheet names the run it was made for, out of the way")

cells = {(cell.column_letter, cell.row): cell.value for row in sheet.iter_rows() for cell in row
         if cell.value is not None}
known = {}


def value(column, row):
    """A cell worked out as Excel would, for the formulas the sheet has."""
    if (column, row) not in known:
        raw = cells.get((column, row))
        known[(column, row)] = formula(raw) if isinstance(raw, str) and raw.startswith("=") else \
            float(raw) if isinstance(raw, (int, float)) else 0.0
    return known[(column, row)]


def formula(text):
    found = re.fullmatch(r"=([A-Z])(\d+)\*([A-Z])(\d+)", text)
    if found:
        return value(found[1], int(found[2])) * value(found[3], int(found[4]))
    found = re.fullmatch(r"=([A-Z])(\d+)\+([A-Z])(\d+)", text)
    if found:
        return value(found[1], int(found[2])) + value(found[3], int(found[4]))
    found = re.fullmatch(r"=SUM\(([A-Z])(\d+):\1(\d+)\)", text)
    if found:
        return sum(value(found[1], row) for row in range(int(found[2]), int(found[3]) + 1))
    found = re.fullmatch(r'=SUMIF\(\$A\$(\d+):\$A\$(\d+),"([^"]+)",([A-Z])\$\1:\4\$\2\)', text)
    if found:
        return sum(value(found[4], row) for row in range(int(found[1]), int(found[2]) + 1)
                   if cells.get(("A", row)) == found[3])
    fail.append("the sheet has a formula this check cannot work out: %s" % text)
    return 0.0


for row, want in ((20, (988.25, 2000, 2988.25)), (39, (2610, 1500, 4110)), (40, (3598.25, 3500, 7098.25))):
    got = tuple(round(value(column, row), 2) for column in "GHI")
    if got != want:
        fail.append("row %d's BASIC, LOOMS OT and TOTAL EARN work out to %s, not %s: each month total adds "
                    "the LOOMS OT in, and the grand total every month total" % (row, got, want))

back = S.read(built)
if back["tag"] != {"run": "LPL-OPR-2026-00001", "branch": "Namanve", "from": "2026-07-26", "to": "2026-08-25"}:
    fail.append("the sheet read back names its run, plant and month: %s" % back["tag"])
if len(back["lines"]) != 2 * 3 * 6 or back["lines"][0] != {
        "row": 2, "block": 1, "name": "ALICE", "employee_id": "LPL001", "date": "26-Jul-2026", "work_unit": 5.2,
        "work_done": 100, "size": "45CM-60CM", "looms_ot": 2000}:
    fail.append("every row of every day comes back, the first as it was written: %s"
                % (back["lines"][:1],))
rates_back = {label: [{"work_unit": unit, "valid_from": O.SHEET_FROM}] for label, unit in O.SHEET_SIZES}
rates_back["74CM"].append({"work_unit": 7.0, "valid_from": "2026-07-28"})
again = O.take_sheet(back["lines"], {"LPL001": "E1", "LPL002": "E2"}, {"E1": O.PER_METER, "E2": O.PER_METER},
                     KEYED, rates_back, PERIOD)
if {key: done for key, (done, _unit, _row) in again["work"].items()} != {key: float(v) for key, v in WORK.items()} \
        or {key: amount for key, (amount, _row) in again["overtime"].items()} != EXTRA \
        or again["problems"] or again["notes"]:
    fail.append("what the sheet was given comes back as it was, with nothing to say: %s" % again)
paid = {row["employee"]: row["amount"] for row in O.per_employee(
    O.merge_lines([{"employee": key[0], "size": key[2], "output": done, "amount": O.line_total(done, unit)}
                   for key, (done, unit, _row) in again["work"].items()]),
    O.overtime_by_employee([{"employee": key[0], "amount": amount}
                            for key, (amount, _row) in again["overtime"].items()]))}
if paid != {"E1": round(value("I", 20), 2), "E2": round(value("I", 39), 2)}:
    fail.append("the system pays each person the month total on the sheet: %s against %s and %s"
                % (paid, value("I", 20), value("I", 39)))

luuka = Workbook()
own = luuka.active
own.title = "PER METER"
own.append(["PER METER AUGUST 2026"])
own.append([])
own.append(["Emp Names", "Employee Id", "Attendance Date", "Work Unit", "Work Done", "TOTAL", "BASIC",
            "LOOMS OT", "TOTAL EARN", None])
own.append(["ALICE", "LPL001", "26-Jul-2026", 5.2, 100, "=E4*D4", "=SUM(F4:F5)", 2000, "=G4+H4", "45CM-60CM"])
own.append([None, "LPL001", "26-Jul-2026", 6.5, None, "=E5*D5", None, None, None, "74CM"])
own.append(["Month Total", None, None, None, None, None, "=SUM(G4:G5)", None, None, None])
own.append(["BOB", "LPL002.", "26-Jul-2026", 8.7, 300, "=E7*D7", "=SUM(F7:F7)", None, "=G7+H7", "ABOVE 80CM"])
out = io.BytesIO()
luuka.save(out)
theirs = S.read(out.getvalue())
if theirs["tag"] is not None or [(row["row"], row["block"], row["name"], row["employee_id"], row["size"])
                                 for row in theirs["lines"]] != [
        (4, 1, "ALICE", "LPL001", "45CM-60CM"), (5, 1, "ALICE", "LPL001", "74CM"), (7, 2, "BOB", "LPL002.", "ABOVE 80CM")]:
    fail.append("Luuka's own sheet is read too: headings a few rows down, the size column unheaded, a person's "
                "rows from their name to the next: %s" % theirs)
try:
    S.read(b"not a workbook")
    fail.append("a file that is not a workbook is refused")
except ValueError as error:
    if ".xlsx" not in str(error):
        fail.append("a file that is not a workbook is refused, saying so: %s" % error)
blank = io.BytesIO()
Workbook().save(blank)
try:
    S.read(blank.getvalue())
    fail.append("a workbook with no Per Meter sheet is refused")
except ValueError as error:
    if "Employee Id and Work Done" not in str(error):
        fail.append("a workbook with no Per Meter sheet is refused, saying what it needs: %s" % error)
print("the sheet: Luuka's layout with its formulas, the same pay on it as in the system, and read back")

# ── 7. The hourly casuals ─────────────────────────────────────────────
month = O.hourly_pay([12, 10, 8, 0], 2500)
if month != {"days": 3, "hours": 28.0, "overtime": 2.0, "amount": 70000.0}:
    fail.append("an hourly casual is paid their standard hours; a twelve-hour day pays ten here and its "
                "two hours of overtime are the overtime module's: %s" % month)
if O.STANDARD_HOURS != A.STANDARD_HOURS:
    fail.append("the standard hours are the attendance register's: a twelve-hour shift, two of them overtime")
expect("hours but no rate", O.run_errors(O.HOURLY, [{"employee": "H1", "hours": 20, "amount": 0}], {"H1": 0}),
       "no hourly rate")
expect("a run with nobody in it", O.run_errors(O.PER_METER, []), "nothing in this run")
expect("a run paying less than nothing", O.run_errors(O.PER_METER, [{"employee": "E1", "amount": -1}]),
       "less than nothing")
print("hourly: the standard hours paid at the rate, the overtime counted but not paid twice")

# ── 8. The paper ──────────────────────────────────────────────────────


def spec(name):
    folder = name.lower().replace(" ", "_")
    return json.loads(read("hrms_addon", "hrms_addon", "doctype", folder, folder + ".json"))


def fields(name):
    return {row["fieldname"]: row for row in spec(name)["fields"]}


def options(name, fieldname):
    return fields(name).get(fieldname, {}).get("options", "").split("\n")


for gone in ("production_machine", "machine_rate"):
    if os.path.exists(os.path.join(APP, "doctype", gone)):
        fail.append("the machines are gone, but %s is still in the app" % gone)
size = fields("Output Rate")
if spec("Output Rate").get("autoname") != "field:size" or not size.get("size", {}).get("unique"):
    fail.append("a size is known by its own label, the one the sheet has, and once")
if options("Output Rate", "section") != list(O.OUTPUT_SECTIONS):
    fail.append("a size is a Per Meter size or a Per Piece category")
if size.get("rates", {}).get("options") != "Output Rate Change" or not size["rates"].get("reqd"):
    fail.append("a size carries its Work Units, at least one")
if not size.get("work_unit", {}).get("read_only"):
    fail.append("the Work Unit in force today is shown on the size, not typed")
if "sheet_order" not in size or "enabled" not in size or spec("Output Rate").get("sort_field") != "sheet_order":
    fail.append("a size has its place on the sheet, lists in that order, and can be taken out of use")
change = fields("Output Rate Change")
for fieldname in ("work_unit", "valid_from"):
    if not change.get(fieldname, {}).get("reqd"):
        fail.append("a Work Unit row says its %s" % fieldname)
report = fields("Daily Production Report")
if not spec("Daily Production Report").get("is_submittable"):
    fail.append("the daily report is submitted by the shift supervisor, and only then counted")
if options("Daily Production Report", "source") != list(O.SOURCES) or not report["source"].get("read_only") \
        or report["source"].get("default") != O.ENTERED:
    fail.append("a report says where it came from: entered, unless the sheet made it")
if options("Daily Production Report", "shift") != list(O.SHIFTS):
    fail.append("a report is for the day or the night shift")
if report.get("looms_ot", {}).get("options") != "Daily Production Overtime":
    fail.append("the day's Looms OT is on the report")
if report.get("sheet_run", {}).get("options") != "Output Pay Run":
    fail.append("an uploaded report names the run whose sheet made it")
line_fields = fields("Daily Production Line")
if line_fields.get("size", {}).get("options") != "Output Rate" or not line_fields["size"].get("reqd"):
    fail.append("a daily line names its size")
for fieldname, label in (("output", "Work Done"), ("rate", "Work Unit"), ("amount", "Total")):
    if line_fields.get(fieldname, {}).get("label") != label:
        fail.append("a daily line's %s is called %s, as on the sheet" % (fieldname, label))
for fieldname in ("rate", "amount"):
    if not line_fields.get(fieldname, {}).get("read_only"):
        fail.append("the %s on a daily line is worked out, not typed" % fieldname)
extra_row = fields("Daily Production Overtime")
if "employee" not in extra_row or extra_row.get("amount", {}).get("label") != "Looms OT":
    fail.append("a Looms OT row is a person and an amount")
if not spec("Output Pay Run").get("is_submittable"):
    fail.append("the month's run is submitted, and only then paid")
if options("Output Pay Run", "section") != list(O.SECTIONS):
    fail.append("a run is for exactly the sections the rules know: %s" % (O.SECTIONS,))
run_fields = fields("Output Pay Run")
if run_fields.get("overtime", {}).get("options") != "Output Pay Overtime":
    fail.append("the month's Looms OT is on the run")
if not run_fields.get("uploaded_sheet", {}).get("read_only"):
    fail.append("the run keeps the sheet it was given")
pay_line = fields("Output Pay Line")
if pay_line.get("size", {}).get("options") != "Output Rate":
    fail.append("a run's line names its size")
if pay_line.get("size", {}).get("reqd"):
    fail.append("a line from a report made before the sizes has none, and is paid as it was recorded")
if not pay_line.get("from_reports", {}).get("hidden"):
    fail.append("a line knows whether it came from the reports or was typed in")
employee_fields = fields("Output Pay Employee")
for fieldname in ("output_amount", "looms_ot", "amount", "additional_salary"):
    if fieldname not in employee_fields:
        fail.append("each person's line on the run says %s" % fieldname)
if not spec("Output Pay Settings").get("issingle"):
    fail.append("Output Pay Settings is one page for the company")
if fields("Output Pay Settings").get("standard_hours", {}).get("default") != "10":
    fail.append("the standard day is ten hours by default, as the attendance register has it")
for name in ("Daily Production Report", "Daily Production Line", "Output Pay Run", "Output Pay Line"):
    left = sorted(set(fields(name)) & {"machine", "width_cm", "piece_category", "target", "daily_target",
                                       "achieved"})
    if left:
        fail.append("%s still has the machines' fields: %s" % (name, left))
print("the paper: sizes with Work Units, the day with its Looms OT, the month, and no machines")

# ── 9. The glue ───────────────────────────────────────────────────────
glue = read("hrms_addon", "hrms_addon", "output_pay.py")


def body(name):
    marker = "def %s(" % name
    if marker not in glue:
        fail.append("output_pay.py has no %s" % name)
        return ""
    return glue.split(marker)[1].split("\ndef ")[0]


priced = body("report_validate")
if 'rules.work_unit_on(size.get("rates"), doc.get("report_date"))' not in priced:
    fail.append("a daily line is priced at its size's Work Unit on the day it was made, so a new Work Unit "
                "reprices nothing already made")
if 'if size.get("enabled")' not in priced:
    fail.append("and a size taken out of use prices nothing new")
if "rules.report_errors(" not in priced or "rules.day_clashes(" not in priced:
    fail.append("every wrong line is refused, and a person's day taken from one place")
if '"docstatus": 1' not in body("_others") or '"report_date": doc.get("report_date")' not in body("_others"):
    fail.append("the clashes are with the other submitted reports of the same day")
if "_paid_run(" not in body("report_on_submit"):
    fail.append("a report for a month already paid is refused, or the shift is never counted")
if "output_pay_run" not in body("report_on_cancel"):
    fail.append("a paid report cannot be cancelled from under its pay")
if "attendance_rules.cycle_window(" not in body("_set_period"):
    fail.append("the payroll month is the attendance register's own 26th to 25th, not a second copy")
if "docstatus" not in body("_one_run_a_month"):
    fail.append("one plant's month is paid once")
typed = body("_price_lines")
if 'if not cint(row.get("from_reports"))' not in typed or 'rules.work_unit_on(size["rates"], doc.to_date)' not in typed:
    fail.append("a run checks and prices only the lines typed in; the reports' lines keep their own days' pay")
if 'cint(row.get("from_reports"))' not in body("_check_overtime"):
    fail.append("and only the Looms OT typed in")
if "rules.per_employee(" not in body("_employees_from_lines") or "overtime" not in body("_employees_from_lines"):
    fail.append("each person is paid their lines and their Looms OT")
paying = body("run_on_submit")
for needle, why in (
    ('"Additional Salary"', "the pay is an Additional Salary, which the payroll run already reads"),
    ('"payroll_date": doc.to_date', "on the last day of the payroll month"),
    ('"overwrite_salary_structure_amount": 0', "on top of the basic salary, not in place of it"),
    ('"ref_doctype": RUN', "and traceable to the run that wrote it"),
    ("earning.submit()", "and submitted, or the payroll never sees it"),
    ('row.get("additional_salary")', "and never written twice for the same person"),
):
    if needle not in paying:
        fail.append("run_on_submit: %s" % why)
if "earning.cancel()" not in body("run_on_cancel"):
    fail.append("cancelling a run takes its pay back")
component = body("_component")
if '"depends_on_payment_days": 0' not in component:
    fail.append("the earnings are not pro-rated by payment days: the output is already what was "
                "earned, and pro-rating it would take a missed day off twice")
if '"type": "Earning"' not in component:
    fail.append("output pay is an earning")
hours = body("_fill_hours")
if "rules.hourly_pay(" not in hours or "PRESENT" not in hours:
    fail.append("an hourly casual is paid the standard hours of the days they were present")
if "custom_hourly_rate" not in body("_hourly_rates") or "standard_hourly_rate" not in body("_hourly_rates"):
    fail.append("their own rate first, the standard one otherwise")
kept = body("_fill")
if "from_reports" not in kept or "HAND_LINE_FIELDS" not in kept or "HAND_OT_FIELDS" not in kept \
        or "REPORT_OT" not in kept:
    fail.append("getting the output again takes the reports' Looms OT too, and keeps what was typed in by hand")
out_sheet = body("download_sheet")
if "sheet.build(" not in out_sheet or 'frappe.response["type"] = "binary"' not in out_sheet \
        or "rules.PER_METER" not in out_sheet:
    fail.append("a Per Meter run downloads the month's sheet as a file")
if 'seat["enabled"] and any(unit is not None for unit in units.values())' not in out_sheet:
    fail.append("the sheet has the sizes in use that have a Work Unit in the month, and no others")
if "cint(row.sheet_order) <= 0" not in body("_sizes"):
    fail.append("a size given no place on the sheet goes after Luuka's own, not before them")
taking = body("upload_sheet")
for needle, why in (
    ("doc.docstatus != 0", "only a draft run takes a sheet"),
    ('tag.get("branch") != doc.branch', "a sheet made for another plant or month is refused"),
    ("rules.take_sheet(", "every line of it is checked"),
    ("_keep_entered_days(", "a day already on a submitted report stays there"),
    ("_replace_upload(", "the last upload is replaced, not added to"),
    ("_fill(doc)", "and the run is filled again from the reports"),
):
    if needle not in taking:
        fail.append("upload_sheet: %s" % why)
replacing = body("_replace_upload")
last = replacing.split("frappe.get_all(REPORT, filters={")[1].split("}")[0] if "frappe.get_all(REPORT, filters={" \
    in replacing else ""
if '"source": rules.UPLOADED' not in last or '"branch": doc.branch' not in last or "sheet_run" in last:
    fail.append("the last upload is found by the plant and month, not the run, so an amended run still "
                "replaces it")
if "report.cancel()" not in replacing or "frappe.delete_doc(" not in replacing or "report.docstatus = 1" not in replacing:
    fail.append("the last upload's reports are cancelled and deleted, the new ones submitted")
entered = body("_keep_entered_days")
if "row.source or rules.ENTERED" not in entered or '["!=", rules.UPLOADED]' in entered:
    fail.append("a report with no source was entered by hand: a != filter would skip it")
if "rules.keep_reported_days(" not in entered or "row.branch == doc.branch" not in entered:
    fail.append("the days on other submitted reports stay theirs; only this plant's last upload is replaced")
for name in ("output_pay.py", "output_rules.py", "output_sheet.py"):
    if re.search(r"Production Machine|Machine Rate|machine_name|width_cm", read("hrms_addon", "hrms_addon", name)):
        fail.append("%s still reaches for the machines" % name)
print("glue: priced on the day, one place for a day, paid once as Additional Salary, taken back on cancel")

# ── 10. The wiring ────────────────────────────────────────────────────
hooks = read("hrms_addon", "hooks.py")
for doctype, events in (("Output Rate", ("validate",)),
                        ("Daily Production Report", ("validate", "on_submit", "on_cancel")),
                        ("Output Pay Run", ("validate", "on_submit", "on_cancel"))):
    block = hooks.split('"%s": {' % doctype)[1].split("},")[0] if '"%s": {' % doctype in hooks else ""
    for event in events:
        if '"%s"' % event not in block:
            fail.append("%s has no %s event" % (doctype, event))
if "Production Machine" in hooks:
    fail.append("hooks.py still has the machines")
after_sync = read("hrms_addon", "patches.txt").split("[post_model_sync]")[-1]
for patch in ("seed_output_pay", "output_rates_by_size"):
    if "hrms_addon.patches.v1_0.%s" % patch not in after_sync:
        fail.append("%s runs on migrate, after the DocTypes are there" % patch)
seed = read("hrms_addon", "patches", "v1_0", "seed_output_pay.py")
if '"depends_on_payment_days": 0' not in seed:
    fail.append("the migrate makes the earnings the same way the code does")
moved = read("hrms_addon", "patches", "v1_0", "output_rates_by_size.py")
for needle, why in (
    ("rules.SHEET_SIZES", "the August sheet's six sizes are made on migrate"),
    ("rules.SHEET_FROM", "each with its Work Unit from the start of August's payroll month"),
    ('table_exists("Machine Rate")', "the Per Piece prices are read from the machines' table while it is there"),
    ("machine.section = 'Per Piece'", "only the Per Piece machines' categories are moved"),
    ("if frappe.db.exists(SIZE, size)", "a size already there is left as it is"),
):
    if needle not in moved:
        fail.append("output_rates_by_size: %s" % why)
rows = json.loads(read("hrms_addon", "fixtures", "custom_field.json"))
if not [row for row in rows if row["dt"] == "Employee" and row["fieldname"] == "custom_hourly_rate"]:
    fail.append("an hourly casual can carry their own rate")
if '"Employee-custom_hourly_rate"' not in hooks:
    fail.append("and that field is in the fixtures hooks.py syncs")
nav = read("hrms_addon", "hrms_addon", "navigation_rules.py")
for name in ("Daily Production Report", "Output Pay Run", "Output Rate", "Output Pay Settings",
             "Sizes and Work Units"):
    if nav.count('"%s"' % name) < 2:
        fail.append("%s needs a way in from the Payroll page and its sidebar" % name)
if "Production Machine" in nav:
    fail.append("the Payroll page still offers the machines")
report_js = read("hrms_addon", "hrms_addon", "doctype", "daily_production_report", "daily_production_report.js")
if 'frm.set_query("size", "lines"' not in report_js or "section: frm.doc.section, enabled: 1" not in report_js:
    fail.append("a report offers only its own section's sizes in use")
if 'frm.set_query("employee", "looms_ot"' not in report_js:
    fail.append("and only its own section's people for the Looms OT")
run_js = read("hrms_addon", "hrms_addon", "doctype", "output_pay_run", "output_pay_run.js")
for needle in ("output_pay.download_sheet", "output_pay.upload_sheet", 'allowed_file_types: [".xlsx"]'):
    if needle not in run_js:
        fail.append("a Per Meter run downloads and uploads its sheet: %s" % needle)
print("wiring: the events, the patches, the hourly rate, and the Payroll page")

if fail:
    print("\nFAILURES:")
    for message in fail:
        print("  -", message)
    sys.exit(1)
print("\nALL OUTPUT PAY CHECKS PASSED")
