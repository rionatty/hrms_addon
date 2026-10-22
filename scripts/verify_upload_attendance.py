"""Checks Upload Attendance without a bench:

    python scripts/verify_upload_attendance.py

Frappe HR v16 deleted its own Upload Attendance (Luuka, 21 Oct 2026: "i
cant find doct type upload attendace"). This app has one again, under the
same name: a template of every day still to be marked, a row per employee
per day, filled in and imported, each day a Frappe HR Attendance marked as
the person importing. For a day the template was the register instead;
Luuka, 22 Oct 2026: "you changed the import file please take it back to
this". What that day kept is kept: no number of rows short of what an Excel
sheet holds is refused, and the marking reads the file as it goes, a part
at a time, so a file of any size is never held whole.

  1. the template: who is on it, from when to when, the days left out, the
     holidays said, and the file as Excel opens it (lists to pick from); a
     company bigger than the cap that once refused 103,182 rows, built in
     seconds with its rows written as openpyxl writes them
  2. the file back: the columns found by their heading, a Data Import file
     and a .csv read too, dates written the ways people write them, read
     as it goes
  3. the checks: each problem said against its row, the rest marked as the
     site names things, a day marked already never marked twice, a day twice
     in the file caught across the parts it is read in
  4. the rows not marked: a file to put right and import again, as it is
  5. the doctype, the background job and the form, wired as they must be
"""
import ast
import datetime
import importlib.util
import io
import json
import os
import re
import sys
import time
import zipfile

import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
APP = os.path.join(REPO, "hrms_addon", "hrms_addon")
fail = []


def check(label, condition, detail=""):
    print(("ok   " if condition else "FAIL ") + label + ("" if condition else "   <- %s" % (detail,)))
    if not condition:
        fail.append(label)


def read(*parts):
    return open(os.path.join(REPO, *parts), encoding="utf-8").read()


spec = importlib.util.spec_from_file_location("attendance_upload_rules", os.path.join(APP, "attendance_upload_rules.py"))
R = importlib.util.module_from_spec(spec)
spec.loader.exec_module(R)  # proves it has no Frappe import
print("loaded attendance_upload_rules.py without Frappe")
D = datetime.date


def raised(run):
    try:
        run()
    except ValueError as error:
        return str(error)
    return None


# ── 1. The template ───────────────────────────────────────────────────
START, END, TODAY = D(2026, 9, 26), D(2026, 10, 25), D(2026, 10, 21)
PEOPLE = [
    {"employee": "LPL001", "employee_name": "Alice Akello", "department": "PE Department - LPT",
     "date_of_joining": D(2020, 1, 6), "relieving_date": None, "holiday_list": "Luuka 2026"},
    # joined in the middle of the cycle
    {"employee": "LPL002", "employee_name": "Bob Okot", "department": "PE Department - LPT",
     "date_of_joining": "2026-10-10", "relieving_date": None, "holiday_list": "Luuka 2026"},
    # left in the middle of it
    {"employee": "LPL003", "employee_name": "Carol Nabirye", "department": "PP Department - LPT",
     "date_of_joining": D(2019, 5, 1), "relieving_date": D(2026, 9, 30), "holiday_list": None},
    # every day marked already
    {"employee": "LPL004", "employee_name": "Dan Mugisha", "department": "PP Department - LPT",
     "date_of_joining": D(2019, 5, 1), "relieving_date": None, "holiday_list": None},
]
MARKED = {("LPL001", D(2026, 9, 30)), ("LPL001", D(2026, 10, 1))}
MARKED |= {("LPL004", START + datetime.timedelta(days=n)) for n in range(40)}
HOLIDAYS = {"Luuka 2026": {D(2026, 10, 9), D(2026, 10, 4)}}
made = R.template_rows(PEOPLE, START, END, TODAY, MARKED, HOLIDAYS)
check("the rows are made one at a time, never held together", hasattr(made, "__next__"))
rows = list(made)
by = {}
for row in rows:
    by.setdefault(row[0], []).append(row[3])
check("a person's rows run from the cycle's first day to today, the days marked left out (26 - 2)",
      len(by.get("LPL001", [])) == 24 and by["LPL001"][0] == START and by["LPL001"][-1] == TODAY
      and D(2026, 9, 30) not in by["LPL001"] and D(2026, 10, 1) not in by["LPL001"],
      (len(by.get("LPL001", [])), by.get("LPL001", [])[:3]))
check("a new hire's start on the day they joined", by.get("LPL002", [None])[0] == D(2026, 10, 10)
      and len(by["LPL002"]) == 12, by.get("LPL002"))
check("a leaver's stop on the day they left", by.get("LPL003", [None])[-1] == D(2026, 9, 30)
      and len(by["LPL003"]) == 5, by.get("LPL003"))
check("somebody with every day marked has no rows", "LPL004" not in by, by.get("LPL004"))
check("rows come person by person, in the order given", [row[0] for row in rows] ==
      ["LPL001"] * 24 + ["LPL002"] * 12 + ["LPL003"] * 5, [row[0] for row in rows][:30])
check("the rows are counted before they are made, as many as there are",
      R.count_rows(PEOPLE, START, END, TODAY, MARKED) == len(rows) == 41, R.count_rows(PEOPLE, START, END, TODAY, MARKED))
notes = {(row[0], row[3]): row[7] for row in rows}
check("a holiday on the person's list says so; the Note is empty otherwise",
      notes[("LPL001", D(2026, 10, 9))] == R.HOLIDAY and notes[("LPL001", D(2026, 10, 4))] == R.HOLIDAY
      and notes[("LPL001", D(2026, 10, 5))] is None and notes[("LPL003", D(2026, 9, 27))] is None)
check("nothing is filled in for them: Status, Leave Type and Shift are blank",
      all(row[4] is None and row[5] is None and row[6] is None for row in rows))
check("a day after today is never on the template", all(row[3] <= TODAY for row in rows))
check("period_days: a person who joins after the cycle has no days",
      R.period_days(START, END, TODAY, joined=D(2026, 11, 1)) == [])

LEAVES = ["Annual", "Leave Without Pay", "Sick Leave"]
SHIFTS = ["Day Shift", "Night Shift"]
content = R.build_template(iter(rows), len(rows), LEAVES, SHIFTS)
book = openpyxl.load_workbook(io.BytesIO(content))
check("the file has the sheet to fill, the lists (hidden) and how to fill it",
      book.sheetnames == [R.SHEET, R.LISTS, R.HOW] and book[R.LISTS].sheet_state == "hidden", book.sheetnames)
sheet = book[R.SHEET]
check("its header is the template's: Employee, Employee Name, Department, Attendance Date, Status, Leave Type, "
      "Shift, Note", [cell.value for cell in sheet[1]] == list(R.HEADERS) == [
          "Employee", "Employee Name", "Department", "Attendance Date", "Status", "Leave Type", "Shift", "Note"],
      [cell.value for cell in sheet[1]])
check("a row for every day, under the header", sheet.max_row == len(rows) + 1, sheet.max_row)
check("dates are dates, written 2026-10-05", isinstance(sheet["D2"].value, datetime.datetime)
      and sheet["D2"].number_format == "yyyy-mm-dd", (sheet["D2"].value, sheet["D2"].number_format))
check("the employee number is kept as text", sheet["A2"].number_format == "@", sheet["A2"].number_format)
check("the header is bold on the template's blue-grey", sheet["A1"].font.bold
      and sheet["A1"].fill.fgColor.rgb.endswith("DDE4EE"), sheet["A1"].fill.fgColor.rgb)
check("a holiday's Note is grey", sheet.cell(2 + [r[3] for r in rows].index(D(2026, 10, 4)), 8).font.color.rgb.endswith("808080"))
check("the columns are as wide as they were",
      [round(sheet.column_dimensions[c].width) for c in "ABCDEFGH"] == [14, 30, 28, 16, 18, 22, 18, 12],
      [sheet.column_dimensions[c].width for c in "ABCDEFGH"])
pick = {str(rule.sqref): (rule.formula1, rule.showErrorMessage) for rule in sheet.data_validations.dataValidation}
check("each Status is picked from Frappe HR's five", pick.get("E2:E%d" % (len(rows) + 1)) ==
      ('"%s"' % ",".join(R.STATUSES), False), pick)
check("each Leave Type and Shift from the site's own, on the hidden sheet",
      pick.get("F2:F%d" % (len(rows) + 1)) == ("'Lists'!$A$2:$A$4", False)
      and pick.get("G2:G%d" % (len(rows) + 1)) == ("'Lists'!$B$2:$B$3", False), pick)
check("the lists hold the names given", [book[R.LISTS].cell(r, 1).value for r in range(2, 5)] == LEAVES
      and [book[R.LISTS].cell(r, 2).value for r in range(2, 4)] == SHIFTS)
check("the header stays in view and filters", sheet.freeze_panes == "B2" and sheet.auto_filter.ref.startswith("A1:H"),
      (sheet.freeze_panes, sheet.auto_filter.ref))
check("how to fill it says what a blank day means and what On Leave needs",
      [row[0] for row in book[R.HOW].iter_rows(values_only=True)] == list(R.HOW_TO_FILL)
      and any("blank is not marked" in line for line in R.HOW_TO_FILL)
      and any("On Leave needs a Leave Type" in line for line in R.HOW_TO_FILL))
plain = R.build_template(iter(rows[:2]), 2)
check("with no leave types or shifts on the site, no list is offered for them",
      len(openpyxl.load_workbook(io.BytesIO(plain))[R.SHEET].data_validations.dataValidation) == 1)
check("no number of rows short of an Excel sheet's is too many", R.MAX_ROWS == 1048575 and R.MAX_DAYS == 366)
company = [{"employee": "LPL%05d" % n, "employee_name": "Person %d" % n, "department": "PE Department - LPT",
            "date_of_joining": D(2020, 1, 1), "relieving_date": None, "holiday_list": None} for n in range(3600)]
began = time.time()
total = R.count_rows(company, D(2026, 8, 26), D(2026, 9, 25), D(2026, 9, 25), set())
big = R.build_template(R.template_rows(company, D(2026, 8, 26), D(2026, 9, 25), D(2026, 9, 25), set()), total)
took = time.time() - began
with zipfile.ZipFile(io.BytesIO(big)) as packed:
    counted = packed.read("xl/worksheets/sheet1.xml").count(b"<row ")
check("a whole company of 3,600 over a cycle is 111,600 rows, past the 100,000 that once refused 103,182",
      total == 111600 and counted == 111601, (total, counted))
check("built while the person waits: %.1f seconds, where openpyxl writing each cell took 45" % took, took < 15, took)


def by_openpyxl(rows):
    """The rows as openpyxl writes them a cell at a time, as build_template
    did before 22 Oct 2026."""
    book = openpyxl.Workbook(write_only=True)
    R._add_styles(book)
    sheet = book.create_sheet(R.SHEET)
    sheet.append([R._styled(sheet, title, "ha_head") for title in R.HEADERS])
    for row in rows:
        cells = list(row)
        cells[0] = R._styled(sheet, cells[0], "ha_text")
        cells[3] = R._styled(sheet, cells[3], "ha_date")
        if cells[7]:
            cells[7] = R._styled(sheet, cells[7], "ha_note")
        sheet.append(cells)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def sheet_data(content):
    with zipfile.ZipFile(io.BytesIO(content)) as packed:
        xml = packed.read("xl/worksheets/sheet1.xml").decode("utf-8")
    return xml[xml.index("<sheetData>"):xml.index("</sheetData>")]


ODD = [{"employee": "LPL010", "employee_name": "Alice & <Akello> \"A\"", "department": "Q&A > PE",
        "date_of_joining": D(2020, 1, 6), "relieving_date": None, "holiday_list": "Luuka 2026"},
       {"employee": "0042", "employee_name": " Bob Okot ", "department": "Stores ",
        "date_of_joining": D(2026, 10, 2), "relieving_date": None, "holiday_list": None}]
odd = list(R.template_rows(ODD, START, END, TODAY, set(), HOLIDAYS))
check("the rows are written as openpyxl writes them, to the byte: a name with & < > and quotes, spaces at either "
      "end, an employee number of digits, a holiday", sheet_data(R.build_template(iter(odd), len(odd), LEAVES, SHIFTS))
      == sheet_data(by_openpyxl(iter(odd))) and "&amp; &lt;Akello&gt;" in sheet_data(by_openpyxl(iter(odd))))
odd_back = openpyxl.load_workbook(io.BytesIO(R.build_template(iter([
    ("LPL011", "Ali\x07ce", None, D(2026, 10, 5), None, None, None, None),
    ("LPL012", "", "", D(2026, 10, 5), None, None, None, None)]), 2)))[R.SHEET]
check("a character Excel cannot hold is left out, where openpyxl would refuse the whole template; a blank stays blank",
      [cell.value for cell in odd_back[2]][:4] == ["LPL011", "Alice", None, datetime.datetime(2026, 10, 5)]
      and [cell.value for cell in odd_back[3]][:3] == ["LPL012", None, None], [cell.value for cell in odd_back[2]])

# ── 2. The file back ──────────────────────────────────────────────────
check("the template's heading is read back as it went out", R.header_of(content, "template.xlsx") == R.HEADERS)
check("left as it went out, nothing is read back: a row with no Status is a day nobody marked",
      list(R.days_in(content, "template.xlsx")) == [])
filled = openpyxl.load_workbook(io.BytesIO(content))
ws = filled[R.SHEET]
ws["E2"], ws["E3"] = "Present", " absent "
ws.append(["LPL002", None, None, "05/10/2026", "On Leave", "annual", None, None])
out = io.BytesIO()
filled.save(out)
back = list(R.days_in(out.getvalue(), "filled.xlsx"))
check("what is filled in is read against its row, trimmed",
      [(r["row"], r["employee"], r["status"]) for r in back] ==
      [(2, "LPL001", "Present"), (3, "LPL001", "absent"), (len(rows) + 2, "LPL002", "On Leave")], back[:2])
check("each row keeps its values as they are, for the file of rows not marked",
      len(back[0]["values"]) == len(R.HEADERS) and back[0]["values"][1] == "Alice Akello")
check("a generator: the file is read as it goes, never held whole", hasattr(R.days_in(content, "x.xlsx"), "__next__"))
csv_file = ("﻿Title line\nID,Employee (employee),Attendance Date (attendance_date),Status (status),"
            "Leave Type (leave_type),Shift (shift)\n,LPL001,2026-10-05,Present,,\n,LPL002,,Absent,,\n"
            ",LPL003,2026-10-06,,,\n\n").encode("utf-8")
taken = list(R.days_in(csv_file, "import.CSV"))
check("a Data Import .csv reads too: its headings found under a title line, a row with no Status no day",
      [(row["row"], row["employee"], row["attendance_date"], row["status"]) for row in taken] ==
      [(3, "LPL001", "2026-10-05", "Present"), (4, "LPL002", "", "Absent")], taken)
check("a heading is compared without case, (fieldname) or underscores",
      R.heading("Attendance Date (attendance_date)") == R.heading("attendance_date") == "attendance date"
      and R.heading("  STATUS ") == "status")
check("the columns found are the five read back",
      R.header_columns(["Employee", "Employee Name", "Date", "Status", "Leave Type", "Shift", "Note"]) ==
      {"employee": 0, "attendance_date": 2, "status": 3, "leave_type": 4, "shift": 5})
check("Employee Name is not taken for the Employee", R.header_columns(["Employee Name", "Date", "Status"]) is None)
check("a file with no Employee, Attendance Date and Status columns is refused, saying so",
      raised(lambda: R.header_of(b"Name,When\nAlice,2026-10-05\n", "x.csv")) == R.NO_HEADER
      and raised(lambda: list(R.days_in(b"Name,When\nAlice,2026-10-05\n", "x.csv"))) == R.NO_HEADER)
check("a file that is not a workbook is refused, saying what to save it as",
      raised(lambda: R.header_of(b"\xd0\xcf\x11\xe0 an old .xls", "x.xls")) == R.NOT_A_SHEET)
check("the heading is looked for near the top only",
      raised(lambda: R.header_of(("\n" * 12 + "Employee,Attendance Date,Status\nLPL001,2026-10-05,Present\n").encode(),
                                 "x.csv")) == R.NO_HEADER)
check("a register of letters, a column per day, is no template now: it is said so",
      raised(lambda: R.header_of("Employee,Employee Name,26/09/2026\nLPL001,Alice,M\n".encode(), "x.csv")) == R.NO_HEADER)
for given, wanted in (("2026-10-05", D(2026, 10, 5)), ("05/10/2026", D(2026, 10, 5)), ("05-10-2026", D(2026, 10, 5)),
                      ("5.10.2026", D(2026, 10, 5)), ("2026-10-05 00:00:00", D(2026, 10, 5)),
                      (datetime.datetime(2026, 10, 5, 0, 0), D(2026, 10, 5)), (D(2026, 10, 5), D(2026, 10, 5)),
                      (46300, D(2026, 10, 5)), ("31/02/2026", None), ("soon", None), ("", None), (None, None),
                      (12, None), (True, None)):
    check("the date %r reads as %s" % (given, wanted), R.parse_date(given) == wanted, R.parse_date(given))

# ── 3. The checks ─────────────────────────────────────────────────────
EMPLOYEES = {
    "LPL001": {"date_of_joining": D(2020, 1, 6), "relieving_date": None, "status": "Active"},
    "LPL002": {"date_of_joining": D(2026, 10, 10), "relieving_date": None, "status": "Active"},
    "LPL003": {"date_of_joining": D(2019, 5, 1), "relieving_date": D(2026, 9, 30), "status": "Left"},
    "LPL004": {"date_of_joining": D(2019, 5, 1), "relieving_date": None, "status": "Inactive"},
    "LPL005": {"date_of_joining": D(2019, 5, 1), "relieving_date": None, "status": "Suspended"},
}
NAMES = {"leave_types": LEAVES, "shifts": SHIFTS}
ON_SITE = {("LPL001", D(2026, 9, 30)): ("Present", "HR-ATT-2026-00001"),
           ("LPL001", D(2026, 10, 1)): ("Absent", "HR-ATT-2026-00002")}
CASES = [
    # (employee, date, status, leave type, shift, what it comes to)
    ("LPL001", D(2026, 10, 2), "Present", None, None, "mark"),
    ("lpl001", "2026-10-03", " absent ", None, None, "mark"),
    ("LPL001", "05/10/2026", "on leave", "annual", None, "mark"),
    ("LPL001", D(2026, 10, 6), "Half Day", "Sick Leave", None, "mark"),
    ("LPL001", D(2026, 10, 7), "Present", "Annual", "night shift", "mark"),
    ("LPL005", D(2026, 10, 7), "Work From Home", None, None, "mark"),
    (None, D(2026, 10, 8), "Present", None, None, "No employee on this row."),
    ("LPL999", D(2026, 10, 8), "Present", None, None, "No employee LPL999."),
    ("LPL001", None, "Present", None, None, "No Attendance Date."),
    ("LPL001", "31/02/2026", "Present", None, None, "31/02/2026 is not a date. Write it as 2026-10-05."),
    ("LPL001", D(2026, 10, 8), "Here", None, None,
     "Here is not a Status. Use Present, Absent, On Leave, Half Day or Work From Home."),
    ("LPL001", D(2026, 10, 8), "M", None, None,
     "M is not a Status. Use Present, Absent, On Leave, Half Day or Work From Home."),
    ("LPL001", D(2026, 10, 22), "Present", None, None, "2026-10-22 is after today."),
    ("LPL002", D(2026, 10, 9), "Present", None, None, "LPL002 joined on 2026-10-10."),
    ("LPL003", D(2026, 10, 1), "Present", None, None, "LPL003 left on 2026-09-30."),
    ("LPL004", D(2026, 10, 8), "Present", None, None,
     "LPL004 is Inactive, and attendance is not marked for an inactive employee."),
    ("LPL001", D(2026, 10, 9), "On Leave", None, None, "On Leave needs a Leave Type."),
    ("LPL001", D(2026, 10, 9), "On Leave", "Maternity", None, "No leave type Maternity."),
    ("LPL001", D(2026, 10, 9), "Present", None, "Swing", "No shift Swing."),
    ("LPL001", D(2026, 10, 2), "Absent", None, None, "LPL001 on 2026-10-02 is on row 2 too."),
    ("LPL001", D(2026, 9, 30), "present", None, None, "already"),
    ("LPL001", D(2026, 10, 1), "Present", None, None,
     "Already marked Absent on HR-ATT-2026-00002. Cancel that one to mark the day again."),
    ("LPL003", D(2026, 9, 30), "Present", None, None, "mark"),
]
ROWS = [{"row": number, "employee": who, "attendance_date": day, "status": status, "leave_type": leave, "shift": shift,
         "values": (who, None, None, day, status, leave, shift, None)}
        for number, (who, day, status, leave, shift, _what) in enumerate(CASES, 2)]
seen = {}
result = R.check_days(ROWS[:12], EMPLOYEES, NAMES, ON_SITE, TODAY, seen=seen)
rest = R.check_days(ROWS[12:], EMPLOYEES, NAMES, ON_SITE, TODAY, seen=seen)
said = {row["row"]: why for row, why in result["problems"] + rest["problems"]}
marking = {row["row"]: row for row in result["to_mark"] + rest["to_mark"]}
for number, (who, day, status, leave, shift, what) in enumerate(CASES, 2):
    if what == "mark":
        check("row %d (%s, %s) is marked" % (number, who, status), number in marking and number not in said,
              said.get(number))
    elif what == "already":
        check("row %d is marked already, not a problem" % number, number not in said and number not in marking,
              said.get(number))
    else:
        check("row %d: %s" % (number, what), said.get(number) == what, said.get(number))
check("the day marked already with the same Status is counted as marked already",
      result["already"] + rest["already"] == 1, result["already"] + rest["already"])
check("a day twice in the file is caught though the two are read in different parts",
      21 in said and 21 not in marking and "on row 2 too" in said[21])
check("every problem names its row", all(row.get("row") for row, _why in result["problems"] + rest["problems"]))
two, three, five, six, seven = (marking[n] for n in (2, 3, 4, 5, 6))
check("the names are the site's: employee, Status, Leave Type and Shift as the site spells them",
      (three["employee"], three["status"], five["status"], five["leave_type"], seven["shift"]) ==
      ("LPL001", "Absent", "On Leave", "Annual", "Night Shift"), (three, five, seven))
check("a Half Day keeps its Leave Type; a Present day drops one, as Frappe HR would",
      six["leave_type"] == "Sick Leave" and seven["leave_type"] is None, (six, seven))
check("a date written day first is the right day", five["attendance_date"] == D(2026, 10, 5),
      five["attendance_date"])
check("a suspended employee is marked: Frappe HR refuses only the inactive",
      marking[7]["employee"] == "LPL005", marking.get(7))
check("a leaver is marked up to the day they left", marking[24]["attendance_date"] == D(2026, 9, 30),
      marking.get(24))
check("each day to mark keeps the row it came from, for the file of rows not marked",
      two["source"] is ROWS[0] and two["row"] == 2)

# ── 4. The rows not marked ────────────────────────────────────────────
header = R.HEADERS
problems = result["problems"] + rest["problems"]
again = R.problems_file(header, list(reversed(problems)))
sheet = openpyxl.load_workbook(io.BytesIO(again))[R.SHEET]
check("the rows not marked come with why, in a last column", [c.value for c in sheet[1]] == list(header) + [R.WHY],
      [c.value for c in sheet[1]])
check("in the order of the file, whatever order they were found in",
      [sheet.cell(r, 9).value for r in range(2, sheet.max_row + 1)] ==
      [why for _row, why in sorted(problems, key=lambda p: p[0]["row"])])
reread = list(R.days_in(again, "rows not marked.xlsx"))
check("the file goes back in as it is: each row with its Status read again",
      len(reread) == len(problems) and R.header_of(again, "x.xlsx")[-1] == R.WHY, len(reread))
twice = openpyxl.load_workbook(io.BytesIO(R.problems_file(R.header_of(again, "x.xlsx"),
                                                          [(row, "again") for row in reread])))
check("put right and not marked again, it carries one Why column, not two",
      [c.value for c in twice[R.SHEET][1]].count(R.WHY) == 1, [c.value for c in twice[R.SHEET][1]])
dated = openpyxl.load_workbook(io.BytesIO(R.problems_file(header, [(ROWS[0], "x")])))[R.SHEET]
check("a date in a row not marked is still a date", isinstance(dated["D2"].value, datetime.datetime)
      and dated["D2"].number_format == "yyyy-mm-dd", (dated["D2"].value, dated["D2"].number_format))
lines = R.listed(list(reversed(problems)), 3).split("\n")
check("the upload lists the first problems by row, whatever order they were found in, and says how many more are "
      "in the file", lines[:2] == ["Row 8: No employee on this row.", "Row 9: No employee LPL999."]
      and lines[-1] == "And %d more, in the file." % (len(problems) - 3) and len(lines) == 4, lines)
check("all of them, when there are few", R.listed(problems[:2], 50).count("\n") == 1)

# ── 5. The doctype, the job and the form ──────────────────────────────
doc = json.loads(read("hrms_addon", "hrms_addon", "doctype", "upload_attendance", "upload_attendance.json"))
fields = {f["fieldname"]: f for f in doc["fields"]}
glue = read("hrms_addon", "hrms_addon", "attendance_upload.py")
tree = ast.parse(glue)
consts = {}
for node in tree.body:
    if isinstance(node, ast.Assign):
        for target in node.targets:
            if isinstance(target, ast.Name):
                try:
                    consts[target.id] = ast.literal_eval(node.value)
                except ValueError:
                    # a tuple of the names above, as RUNNING is
                    if isinstance(node.value, ast.Tuple) and all(isinstance(elt, ast.Name) for elt in node.value.elts):
                        consts[target.id] = tuple(consts.get(elt.id) for elt in node.value.elts)
            elif isinstance(target, ast.Tuple) and isinstance(node.value, ast.Tuple):
                for name, value in zip(target.elts, node.value.elts):
                    consts[name.id] = ast.literal_eval(value)
check("the doctype is Upload Attendance, the name people search for", doc["name"] == "Upload Attendance"
      and consts.get("DOCTYPE") == "Upload Attendance")
check("one record per upload, numbered, kept in this app", not doc.get("issingle") and not doc.get("is_submittable")
      and doc.get("autoname", "").startswith("HR-ATU-") and doc.get("module") == "HRMS Addon")
check("its statuses are the job's", fields["status"]["options"].split("\n") == [
    consts.get(name) for name in ("DRAFT", "QUEUED", "IMPORTING", "DONE", "DONE_WITH_PROBLEMS", "FAILED")],
    fields["status"]["options"])
check("what the job writes is read-only on the form", all(fields[name].get("read_only") for name in
                                                         ("status", "imported_by", "imported_on", "days_in_file",
                                                          "marked", "already_marked", "not_marked", "problems_file",
                                                          "problems")))
check("the counts the job writes are fields", all(name in fields for name in consts.get("COUNTS", ())),
      consts.get("COUNTS"))
check("the company and the period are asked for", all(fields[name].get("reqd") for name in
                                                      ("company", "from_date", "to_date")))
check("the plant is called a plant, as on Luuka's other forms", fields["branch"]["label"] == "Plant")
check("the form says a Status on each day, as the template asks",
      "a Status on each day" in fields["attendance_file"].get("description", "")
      and "a Status on each day" in doc.get("description", "")
      and not re.search(r"letter|register", json.dumps(doc)), fields["attendance_file"].get("description"))
check("the doctype is newer than the one that said the register, so the migrate takes it",
      doc.get("modified", "") >= "2026-10-22", doc.get("modified"))
roles = {p["role"] for p in doc["permissions"]}
check("HR only: the HR User, the HR Manager and the System Manager", roles == {"HR User", "HR Manager",
                                                                                 "System Manager"}, roles)
check("Frappe HR's Inactive is not among those put on the template", "Inactive" not in consts.get("ON_STAFF", ())
      and "Left" in consts.get("ON_STAFF", ()), consts.get("ON_STAFF"))


def body(name):
    found = re.search(r"^def %s\(.*?(?=^def |^# ──|\Z)" % name, glue, re.S | re.M)
    return found.group(0) if found else ""


for name in ("this_cycle", "download_template"):
    check("%s can be called from the form" % name,
          re.search(r"@frappe\.whitelist\(\)\ndef %s\(" % name, glue) is not None)
template = body("download_template")
check("the template is for HR, and for those who may make an upload",
      "attendance.only_hr()" in template and 'frappe.has_permission(DOCTYPE, "create", throw=True)' in template)
check("the period is held to a year, and the rows to what one Excel sheet holds, counted before the file is made",
      "(end - start).days + 1 > rules.MAX_DAYS" in template and "total > rules.MAX_ROWS" in template
      and template.index("rules.count_rows(") < template.index("rules.build_template(")
      and "100000" not in glue)
check("a period with no day left to mark is said, rather than an empty file",
      "if not total:" in template and template.index("if not total:") < template.index("rules.build_template("))
check("the rows go into the file as they are made, the days marked already left out, in draft too",
      "rules.build_template(rules.template_rows(" in template and '"docstatus": ["<", 2]' in body("_marked")
      and "register_code" not in glue)
check("a department takes in the departments under it", '"descendants of (inclusive)"' in body("_people"))
check("the template comes back in the answer with how many rows it has, so a problem is a message and not an "
      "error page", "base64.b64encode(content)" in template and '"rows": total' in template
      and "frappe.response" not in template)
importing = body("start_import")
check("Import is HR's, on an upload they may change", "attendance.only_hr()" in importing
      and 'doc.check_permission("write")' in importing)
check("Import reads only the file's heading while the person waits, then queues the rest",
      importing.index("rules.header_of(") < importing.index("frappe.enqueue(") and "rules.days_in(" not in importing)
check("an upload already being marked is not queued twice", "if running(doc):" in importing
      and 'job_id="upload_attendance::%s" % doc.name' in importing and "deduplicate=True" in importing)
check("the job starts after the upload is saved as Queued", "enqueue_after_commit=True" in importing
      and importing.index("doc.db_set(") < importing.index("frappe.enqueue("))
check("the job runs where long jobs do, given a day", 'queue="long"' in importing
      and "timeout=JOB_HOURS * 3600" in importing and consts.get("JOB_HOURS") == 24)
marking = body("_mark_file")
check("the rows are counted as they are read, for the bar, and a file with none is said",
      "sum(1 for _record in rules.days_in(" in marking and "rules.NOTHING_FILLED" in marking)
check("then read again a part at a time, never held whole",
      "for part in _parts(rules.days_in(content, doc.attendance_file), PART):" in marking
      and consts.get("PART") == 2000)
check("each part checked before a day of it is written, the days met so far carried from part to part",
      marking.index("rules.check_days(") < marking.index("_mark_day(") and "seen=seen" in marking)
check("each part's employees read once, by an index, and a name the site lacks not asked twice",
      "index[name.lower()] = None" in body("_learn") and "index[row.name.lower()] = row.name" in body("_learn")
      and "_learn(people, index, part)" in marking)
check("the counts saved as it goes, which keeps the upload moving, and the bar moved by the percent",
      "doc.db_set(counts, commit=True)" in marking and "if done % BATCH == 0:" in marking
      and "frappe.publish_progress(" in marking and "if percent != shown:" in marking)
check("the rows not marked go back as a file on the upload, under the file's own heading",
      "_problems_file(doc, rules.header_of(content, doc.attendance_file), problems)" in marking
      and '"attached_to_field": "problems_file"' in body("_problems_file")
      and "rules.problems_file(header, problems)" in body("_problems_file"))
day = body("_mark_day")
check("each day is marked and submitted on its own, rolled back alone when it fails",
      "frappe.db.savepoint(SAVEPOINT)" in day and "frappe.db.rollback(save_point=SAVEPOINT)" in day
      and day.index("day.insert()") < day.index("day.submit()"))
check("and as the person who pressed Import: no permission is waived", "ignore_permissions" not in day)
check("Frappe HR's notes on each day are not kept in memory", "frappe.local.message_log = []" in day)
job = body("mark_all")
check("a job that stops says so on the upload, as Failed, and in the Error Log",
      "frappe.db.rollback()" in job and "frappe.log_error(" in job and '"status": FAILED' in job)
check("an upload that has not moved for half an hour has lost its job, and is free to go on",
      "add_to_date(now_datetime(), minutes=-STOPPED_AFTER)" in body("running") and consts.get("STOPPED_AFTER") == 30
      and 'doc.get("modified")' in body("running"))
check("and the form is told so", 'doc.set_onload("stopped", doc.get("status") in RUNNING and not running(doc))'
      in body("upload_onload"))
check("a file is read as bytes, a .csv too", 'open(path, "rb")' in body("_file_bytes")
      and "get_content()" not in body("_file_bytes"))
check("a row's Leave Type and Shift are checked against the site's own",
      '{"leave_types": _names("Leave Type"), "shifts": _names("Shift Type")}' in body("_site_names"))

hooks = read("hrms_addon", "hooks.py")
after = hooks[hooks.index("after_migrate = ["):]
after = after[:after.index("]")]
check("after every migrate the doctype is put back if Frappe HR's patch took it, before the navigation",
      '"hrms_addon.hrms_addon.attendance_upload.keep_doctype"' in after
      and after.index("attendance_upload.keep_doctype") < after.index("navigation.setup_on_migrate"))
check("which reloads ours only when it is gone", 'if not frappe.db.exists("DocType", DOCTYPE):' in body("keep_doctype")
      and 'frappe.reload_doc("hrms_addon", "doctype", "upload_attendance", force=True)' in body("keep_doctype"))
controller = read("hrms_addon", "hrms_addon", "doctype", "upload_attendance", "upload_attendance.py")
check("the form's Import is a method of the upload", re.search(
    r"@frappe\.whitelist\(\)\n    def start_import\(self\):\n        return attendance_upload\.start_import\(self\)",
    controller) is not None)
check("the upload's own checks run on load, save and delete", "attendance_upload.upload_validate(self)" in controller
      and "attendance_upload.upload_on_trash(self)" in controller and "attendance_upload.upload_onload(self)" in controller)
form = read("hrms_addon", "hrms_addon", "doctype", "upload_attendance", "upload_attendance.js")
check("the form calls the template and the cycle where they are",
      '"hrms_addon.hrms_addon.attendance_upload."' in form and 'HA_UPLOAD_ATTENDANCE + "download_template"' in form
      and 'HA_UPLOAD_ATTENDANCE + "this_cycle"' in form)
check("it says to put a Status on each day, and how many rows the template has",
      "put a Status on each day to mark" in form and '__("{0} row(s) to fill.", [file.rows])' in form
      and "letter" not in form and "register" not in form)
check("Import calls the upload's own method with the upload", 'method: "start_import"' in form
      and "doc: frm.doc" in form)
check("an upload is saved before it is imported", "frm.is_new() || frm.is_dirty() ? frm.save()" in form)
check("the form knows which statuses are running", "HA_UPLOAD_RUNNING = [%s]" % ", ".join(
    '"%s"' % status for status in consts.get("RUNNING", ())) in form)
check("a stopped marking offers Import again, to go on", "(frm.doc.__onload || {}).stopped" in form
      and "HA_UPLOAD_RUNNING.includes(frm.doc.status) && !stopped" in form
      and "if (frm.doc.attendance_file && !running)" in form)
check("one hint at a time: the form clears its own before it draws one (two stood on the page, 21 Oct 2026)",
      'frm.layout.message.find(".ha-upload-hint").remove();' in form
      and 'frm.$intro_message.addClass("ha-upload-hint")' in form and form.count("frm.set_intro(") == 1)
listing = read("hrms_addon", "hrms_addon", "doctype", "upload_attendance", "upload_attendance_list.js")
check("the list colours every status", all('%s:' % s in listing or '"%s":' % s in listing
                                           for s in fields["status"]["options"].split("\n")))

print()
if fail:
    print("%d FAILED" % len(fail))
    sys.exit(1)
print("ALL UPLOAD ATTENDANCE CHECKS PASSED")
